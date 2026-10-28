from playwright.sync_api import (
    sync_playwright,
    TimeoutError as PlaywrightTimeoutError,
)

from openpyxl import load_workbook

import csv
import os
import re
import zipfile
import io
import json
import unicodedata
import shutil
import traceback
import time
import threading
import tempfile
import html as html_lib
from urllib.parse import quote
from urllib.request import Request, urlopen
from urllib.error import HTTPError, URLError


# ============================================================
# CONFIGURACIÓN
# ============================================================

# URL_PRIZMA, CHROMIUM_ARGS y VIEWPORT_NAVEGADOR viven en auto_prizma.core.config.
from auto_prizma.core.config import (  # noqa: E402
    CHROMIUM_ARGS,
    EXTENSIONES_HTML,
    PRIZMA_ESPERAS_REINTENTO_SESION,
    PRIZMA_VACIAR_HTML_CON_H5P,
    PRIZMA_VACIAR_RECURSO_CON_HTML,
    FORMATO_RECURSO_POR_DEFECTO,
    URL_PRIZMA,
    VIEWPORT_NAVEGADOR,
)
from auto_prizma.automation.prizma.html_field import (  # noqa: E402
    NOMBRE_ARCHIVO_DESCONOCIDO,
    XPATH_TARJETA_ARCHIVO_RECURSO,
    XPATH_SECCION_RECURSO,
    archivo_en_recurso,
    es_campo_de_texto as _es_campo_de_texto,
    escribir_html_en_editor,
    guardar_estructura,
    leer_archivo_html,
    normalizar_saltos,
    quitar_archivo_de_recurso,
)
from auto_prizma.services.course_state import (  # noqa: E402
    registrar_resultado_actividad,
)
from auto_prizma.services.error_messages import es_codigo_error, mensaje_error  # noqa: E402

VERSION_SCRIPT = "PRUEBA_H5P_NUEVO_GESTION_FINANCIERA_V2"

BUILD_INTERNO = "CONFIRMACION_PATCH_DIRECTA_V11"


TEXTOS_DESCRIPCION_A_BORRAR = [
    "NO_DISPONIBLE",
    "NO_DISPOINBLE",
]


# ============================================================
# ESTADO DEL TRABAJO
# ============================================================

def actualizar_estado(
    estado,
    etapa=None,
    mensaje=None,
    total=None,
    procesadas=None,
    exitosas=None,
    errores=None,
    terminado=None,
):

    if etapa is not None:
        estado["etapa"] = etapa

    if mensaje is not None:
        estado["mensaje"] = mensaje

    if total is not None:
        estado["total"] = total

    if procesadas is not None:
        estado["procesadas"] = procesadas

    if exitosas is not None:
        estado["exitosas"] = exitosas

    if errores is not None:
        estado["errores"] = errores

    if terminado is not None:
        estado["terminado"] = terminado


# ============================================================
# DETALLE DE ACTIVIDADES PARA LA INTERFAZ
# ============================================================

def preparar_detalle_actividades(
    estado,
    actividades,
):

    estado["detalle_actividades"] = [
        {
            "numero": numero,
            "nombre": actividad["nombre"],
            "estado": "pendiente",
            "error": "",
        }
        for numero, actividad in enumerate(
            actividades,
            start=1,
        )
    ]


def actualizar_detalle_actividad(
    estado,
    numero,
    nuevo_estado,
    error="",
):

    detalle = estado.get(
        "detalle_actividades",
        [],
    )

    indice = numero - 1

    if not (
        0 <= indice < len(detalle)
    ):
        return

    detalle[indice][
        "estado"
    ] = nuevo_estado

    detalle[indice][
        "error"
    ] = error or ""


# ============================================================
# NORMALIZACIÓN
# ============================================================

def normalizar_texto(texto):

    if texto is None:
        return ""

    texto = str(texto)

    texto = unicodedata.normalize(
        "NFD",
        texto,
    )

    texto = "".join(
        caracter
        for caracter in texto
        if unicodedata.category(caracter) != "Mn"
    )

    texto = texto.lower()

    texto = re.sub(
        r"[^a-z0-9]+",
        " ",
        texto,
    )

    return " ".join(
        texto.split()
    )


def normalizar_categoria(categoria):

    valor = normalizar_texto(
        categoria
    )

    if valor == "ovi":
        return "OVI"

    if valor == "ova":
        return "OVA"

    if valor == "challenge":
        return "CHALLENGE"

    return None


# ============================================================
# VARIANTES SEGURAS DE PROGRAMA
# ============================================================

# Equivalencias academicas explicitas observadas en PRIZMA.
# Cada grupo representa nombres distintos que pertenecen a la misma
# linea/programa para efectos de localizar una actividad.
#
# IMPORTANTE: evitar fuzzy matching general. Para agregar un caso nuevo,
# agregarlo explicitamente al grupo correcto o crear uno nuevo.
EQUIVALENCIAS_PROGRAMAS = (
    {
        "administracion de empresas",
        "tecnologia en gestion empresarial",
        "tecnica profesional en procesos empresariales",
    },
    {
        "administracion turistica y hotelera",
        "tecnica profesional en procesos turisticos y hoteleros",
        "tecnologia en gestion de turismo y hoteles",
        # Variante historica ya soportada; se conserva por compatibilidad.
        "tecnologia en gestion turistica y hoteles",
    },
    {
        "contaduria publica",
        "tecnologia virtual en gestion contable",
        "tecnica profesional en procesos contables",
    },
    {
        "negocios internacionales",
        "tecnica virtual profesional en procesos logisticos y de comercio exterior",
        "tecnica vistural profesional en procesos logisticos y de comercio exterior",
    },
)


def obtener_variantes_programa(
    programa,
):

    programa_n = normalizar_texto(
        programa
    )

    if not programa_n:
        return set()

    variantes = {
        programa_n
    }

    # Si el nombre pertenece a un grupo de equivalencias conocido,
    # aceptamos todos los nombres de ese mismo grupo. La relacion es
    # deliberadamente explicita y bidireccional.
    for grupo in EQUIVALENCIAS_PROGRAMAS:
        if programa_n in grupo:
            variantes.update(grupo)
            break

    partes = programa_n.split()

    # PRIZMA y las matrices no siempre usan exactamente el mismo
    # prefijo para las especializaciones. Ejemplos reales:
    #   Matriz:  "Esp en Crisis, Conflicto..."
    #   PRIZMA:  "Especialización Crisis, Conflicto... - Virtual"
    #
    # Generamos únicamente variantes equivalentes del prefijo, sin
    # relajar la comparación del nombre propio del programa.
    if partes and partes[0] in {"esp", "especializacion"}:

        resto_partes = partes[1:]

        if resto_partes:
            resto = " ".join(resto_partes)
            variantes.add(
                "especializacion " + resto
            )

            if resto_partes[0] == "en":
                resto_sin_en = " ".join(
                    resto_partes[1:]
                )

                if resto_sin_en:
                    variantes.add(
                        "especializacion "
                        + resto_sin_en
                    )
                    variantes.add(
                        "esp " + resto_sin_en
                    )
            else:
                variantes.add(
                    "especializacion en " + resto
                )
                variantes.add(
                    "esp en " + resto
                )

    return variantes


# ============================================================
# NOMBRES DE RECURSOS CORTOS / TRUNCADOS
# ============================================================

def limpiar_prefijo_tecnico_recurso(
    nombre_archivo,
):

    nombre_sin_extension = os.path.splitext(
        os.path.basename(
            nombre_archivo
        )
    )[0]

    # Copias de Google Drive: "archivo (1).h5p", "archivo (2).pdf".
    # Sin quitar esto, el " (1)" deja el nombre mas largo que la actividad
    # y el prefijo deja de coincidir (fallo real: filas 24 y 46).
    nombre_sin_extension = re.sub(
        r"\s*\(\d+\)$",
        "",
        nombre_sin_extension,
    ).strip()

    nombre_n = normalizar_texto(
        nombre_sin_extension
    )

    # Prefijos técnicos observados en paquetes de recursos.
    # Ejemplos:
    # ID003-Tema_4-Ruta_para_analizar.h5p
    # U1-T1-Principios jurídicos.pdf
    # U2-T4-Caso aplicado.pdf
    #
    # normalizar_texto() convierte guiones y guiones bajos en espacios,
    # por eso los patrones se aplican sobre la versión normalizada.
    nombre_n = re.sub(
        r"^id\d+\s+tema\s+\d+\s+",
        "",
        nombre_n,
    )

    nombre_n = re.sub(
        r"^u\d+\s*t\d+\s+",
        "",
        nombre_n,
    )

    return nombre_n


def recurso_es_prefijo_de_actividad(
    nombre_archivo,
    actividad,
):

    archivo_n = limpiar_prefijo_tecnico_recurso(
        nombre_archivo
    )

    actividad_n = normalizar_texto(
        actividad["nombre"]
    )

    if not archivo_n or not actividad_n:
        return False

    palabras_archivo = archivo_n.split()
    palabras_actividad = actividad_n.split()

    if not palabras_archivo:
        return False

    if len(palabras_archivo) > len(palabras_actividad):
        return False

    # Permitimos incluso una sola palabra cuando es suficientemente
    # descriptiva. La seguridad real se valida después comprobando
    # que ese prefijo identifique una única actividad del curso.
    if len(palabras_archivo[0]) < 4:
        return False

    for indice, palabra_archivo in enumerate(
        palabras_archivo
    ):

        palabra_actividad = palabras_actividad[indice]

        # Todas las palabras anteriores a la última deben coincidir.
        if indice < len(palabras_archivo) - 1:

            if palabra_archivo != palabra_actividad:
                return False

            continue

        # La última palabra puede estar completa o truncada.
        if palabra_archivo == palabra_actividad:
            continue

        if len(palabra_archivo) < 4:
            return False

        if not palabra_actividad.startswith(
            palabra_archivo
        ):
            return False

    return True


def recurso_corto_es_unico_para_actividad(
    nombre_archivo,
    actividad,
    actividades_curso,
):

    if not actividades_curso:
        return False

    tipo_objetivo = actividad[
        "tipo_archivo"
    ]

    compatibles = []

    for candidata in actividades_curso:

        if candidata[
            "tipo_archivo"
        ] != tipo_objetivo:
            continue

        if recurso_es_prefijo_de_actividad(
            nombre_archivo,
            candidata,
        ):

            compatibles.append(
                candidata
            )

    if len(compatibles) != 1:
        return False

    unica = compatibles[0]

    return (
        unica["fila_excel"]
        == actividad["fila_excel"]
    )


def obtener_titulo_interno_h5p(
    recurso,
):
    """
    Lee únicamente h5p.json de un H5P concreto.

    IMPORTANTE:
    - No se usa durante el indexado general.
    - Solo se invoca como fallback cuando el nombre externo quedó
      ambiguo y el flujo normal no pudo resolverlo de forma segura.
    - No modifica la regla de ERROR_RECURSO_DUPLICADO.
    """

    if recurso.get("extension") != ".h5p":
        return ""

    try:
        with zipfile.ZipFile(
            recurso["zip"],
            "r",
        ) as zip_ref:

            datos_h5p = zip_ref.read(
                recurso["miembro"]
            )

        with zipfile.ZipFile(
            io.BytesIO(datos_h5p),
            "r",
        ) as h5p_ref:

            if "h5p.json" not in h5p_ref.namelist():
                return ""

            datos_json = h5p_ref.read(
                "h5p.json"
            )

        manifiesto = json.loads(
            datos_json.decode(
                "utf-8-sig"
            )
        )

        titulo = manifiesto.get(
            "title",
            "",
        )

        return str(titulo).strip()

    except Exception as error:
        print(
            "No se pudo leer h5p.json de",
            recurso.get("nombre", ""),
            "|",
            type(error).__name__,
            str(error),
        )
        return ""


def resolver_h5p_ambiguo_por_titulo_interno(
    actividad,
    recursos_ambiguos,
):
    """
    Intenta resolver un ERROR_RECURSO_AMBIGUO leyendo el título interno
    de los pocos H5P implicados. Exige coincidencia EXACTA normalizada.

    Si encuentra 0 o más de 1 coincidencias, no adivina.
    """

    if not recursos_ambiguos:
        return None

    actividad_n = normalizar_texto(
        actividad["nombre"]
    )

    coincidencias = []

    # Evitar abrir dos veces el mismo miembro si llegó repetido a la lista.
    vistos = set()

    for recurso in recursos_ambiguos:
        clave = (
            recurso.get("zip"),
            recurso.get("miembro"),
        )

        if clave in vistos:
            continue

        vistos.add(clave)

        titulo = obtener_titulo_interno_h5p(
            recurso
        )

        if not titulo:
            continue

        titulo_n = normalizar_texto(
            titulo
        )

        print(
            "Título H5P interno:",
            recurso.get("nombre", ""),
            "->",
            titulo,
        )

        if titulo_n == actividad_n:
            coincidencias.append(
                recurso.copy()
            )

    if len(coincidencias) != 1:
        return None

    elegido = coincidencias[0]
    elegido["puntuacion"] = 1000
    elegido["metodo"] = "TITULO_INTERNO_H5P_EXACTO"

    return elegido


# ============================================================
# TIPO DE ARCHIVO
# ============================================================

def determinar_tipo_archivo(
    categoria,
    tipo_recurso,
    enlace,
):

    categoria_n = normalizar_texto(
        categoria
    )

    tipo_n = normalizar_texto(
        tipo_recurso
    )

    enlace_n = normalizar_texto(
        enlace
    )

    # OVI / OVA / RETO EVALUATIVO
    if categoria_n not in [
        "ovi",
        "ova",
        "challenge",
    ]:
        return None

    # RETO EVALUATIVO SIEMPRE PDF
    if (
        categoria_n == "challenge"
        and tipo_n == "reto evaluativo"
    ):
        return "PDF"

    # VIDEOS EXCLUIDOS
    if tipo_n in [
        "video intro",
        "video cierre",
        "video a camara",
    ]:
        return None

    # OVA SIEMPRE H5P
    if categoria_n == "ova":
        return "H5P"

    # OVI PDF
    if "pdf" in tipo_n:
        return "PDF"

    if "pdf" in enlace_n:
        return "PDF"

    # OVI H5P
    if "h5p" in tipo_n:
        return "H5P"

    if "h5p" in enlace_n:
        return "H5P"

    # INFOGRAFÍA OVI
    if "infografia" in tipo_n:
        return "H5P"

    return None


def aplicar_formato_recurso(actividades, formato_recurso):
    """Aplica el formato de recurso interactivo (H5P/HTML) elegido para el cargue.

    En modo "HTML" las actividades de recurso interactivo (tipo H5P) pasan a
    tipo "HTML": su codigo se pega en el campo Html. Los PDF no cambian.

    Args:
        actividades: lista devuelta por leer_actividades_excel (se modifica).
        formato_recurso: "H5P" o "HTML" (config.FORMATOS_RECURSO_INTERACTIVO).

    Returns:
        La misma lista de actividades.
    """
    if str(formato_recurso or "").strip().upper() != "HTML":
        return actividades

    for actividad in actividades:
        if actividad.get("tipo_archivo") == "H5P":
            actividad["tipo_archivo"] = "HTML"

    return actividades


# ============================================================
# VALIDACIÓN DE CARPETAS DE RECURSOS EN GOOGLE DRIVE
# ============================================================

_CACHE_CARPETAS_DRIVE = {}
_CACHE_CARPETAS_DRIVE_LOCK = threading.Lock()


def extraer_url_celda_recurso(celda_valor, celda_formula):
    """Obtiene el hipervínculo real conservado en el XLSX exportado por Sheets."""
    for celda in (celda_formula, celda_valor):
        try:
            if celda is not None and celda.hyperlink and celda.hyperlink.target:
                return str(celda.hyperlink.target).strip()
        except Exception:
            pass

    valor_formula = ""
    try:
        valor_formula = str(celda_formula.value or "").strip()
    except Exception:
        pass

    coincidencia = re.search(
        r'^\s*=\s*HYPERLINK\(\s*["\']([^"\']+)["\']',
        valor_formula,
        flags=re.IGNORECASE,
    )
    if coincidencia:
        return coincidencia.group(1).strip()

    for celda in (celda_valor, celda_formula):
        try:
            valor = str(celda.value or "").strip()
        except Exception:
            valor = ""
        if valor.lower().startswith(("https://", "http://")):
            return valor

    return ""


def extraer_id_carpeta_drive(url):
    url = str(url or "").strip()
    coincidencia = re.search(
        r"drive\.google\.com/(?:drive/(?:u/\d+/)?folders|folders)/([A-Za-z0-9_-]+)",
        url,
        flags=re.IGNORECASE,
    )
    if not coincidencia:
        return None
    return coincidencia.group(1)


def _limpiar_html_drive(texto):
    texto = html_lib.unescape(str(texto or ""))
    reemplazos = {
        r"\u0026": "&",
        r"\u003d": "=",
        r"\u003c": "<",
        r"\u003e": ">",
        r"\u0027": "'",
        r"\u0022": '"',
    }
    for origen, destino in reemplazos.items():
        texto = texto.replace(origen, destino)
    return texto


def _extraer_nombres_archivos_drive(html_drive):
    texto = _limpiar_html_drive(html_drive)
    nombres = []
    vistos = set()

    # Vista embebida: los nombres suelen aparecer como texto de enlaces.
    for contenido in re.findall(
        r"<a[^>]*>(.*?)</a>",
        texto,
        flags=re.IGNORECASE | re.DOTALL,
    ):
        limpio = re.sub(r"<[^>]+>", " ", contenido)
        limpio = html_lib.unescape(limpio)
        limpio = " ".join(limpio.split()).strip()
        if not limpio or "." not in limpio:
            continue
        if len(limpio) > 260:
            continue
        clave = limpio.casefold()
        if clave not in vistos:
            vistos.add(clave)
            nombres.append(limpio)

    # Fallback para nombres serializados en JSON/HTML.
    for patron in [
        r'["\']([^"\'<>\\]{1,250}\.(?:h5p|pdf|zip|docx?|xlsx?|pptx?|webm|mp4))["\']',
        r'>([^<>]{1,250}\.(?:h5p|pdf|zip|docx?|xlsx?|pptx?|webm|mp4))<',
    ]:
        for nombre in re.findall(patron, texto, flags=re.IGNORECASE):
            nombre = html_lib.unescape(" ".join(nombre.split()).strip())
            clave = nombre.casefold()
            if nombre and clave not in vistos:
                vistos.add(clave)
                nombres.append(nombre)

    return nombres


def listar_archivos_carpeta_drive(url_carpeta):
    carpeta_id = extraer_id_carpeta_drive(url_carpeta)
    if not carpeta_id:
        return {
            "ok": False,
            "error": "ERROR_CARPETA_RECURSO_NO_ACCESIBLE",
            "archivos": [],
        }

    with _CACHE_CARPETAS_DRIVE_LOCK:
        cache = _CACHE_CARPETAS_DRIVE.get(carpeta_id)
    if cache is not None:
        return cache

    urls = [
        "https://drive.google.com/embeddedfolderview?id="
        + quote(carpeta_id)
        + "#list",
        "https://drive.google.com/drive/folders/"
        + quote(carpeta_id),
    ]

    ultimo_error = None
    for url in urls:
        solicitud = Request(
            url,
            headers={
                "User-Agent": (
                    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
                    "AppleWebKit/537.36 Chrome/124 Safari/537.36"
                )
            },
        )
        try:
            with urlopen(solicitud, timeout=30) as respuesta:
                url_final = str(respuesta.geturl() or "")
                datos = respuesta.read(8 * 1024 * 1024)
        except (HTTPError, URLError, TimeoutError, OSError) as error:
            ultimo_error = error
            continue

        if "accounts.google.com" in url_final.lower():
            ultimo_error = RuntimeError("La carpeta requiere autenticación de Google")
            continue

        texto = datos.decode("utf-8", errors="ignore")
        archivos = _extraer_nombres_archivos_drive(texto)

        resultado = {
            "ok": True,
            "error": None,
            "archivos": archivos,
        }
        with _CACHE_CARPETAS_DRIVE_LOCK:
            if len(_CACHE_CARPETAS_DRIVE) > 500:
                _CACHE_CARPETAS_DRIVE.clear()
            _CACHE_CARPETAS_DRIVE[carpeta_id] = resultado
        return resultado

    print(
        "No se pudo consultar carpeta Drive:",
        carpeta_id,
        type(ultimo_error).__name__ if ultimo_error else "Error",
        str(ultimo_error or ""),
    )
    return {
        "ok": False,
        "error": "ERROR_CARPETA_RECURSO_NO_ACCESIBLE",
        "archivos": [],
    }


def validar_carpeta_drive_para_actividad(actividad):
    """
    Si la matriz contiene un hipervínculo a una carpeta Drive, comprueba:
    - que sea accesible;
    - que no esté vacía;
    - que contenga un H5P/PDF del tipo esperado que corresponda a la actividad.

    Si la celda no contiene una carpeta Drive, no altera el flujo histórico del ZIP.
    """
    url_recurso = str(actividad.get("url_recurso") or "").strip()
    if not url_recurso:
        return None

    if not extraer_id_carpeta_drive(url_recurso):
        return None

    resultado = listar_archivos_carpeta_drive(url_recurso)
    if not resultado.get("ok"):
        return resultado.get("error") or "ERROR_CARPETA_RECURSO_NO_ACCESIBLE"

    archivos = resultado.get("archivos") or []
    if not archivos:
        return "ERROR_CARPETA_RECURSO_VACIA"

    extension_objetivo = (
        ".h5p"
        if actividad.get("tipo_archivo") == "H5P"
        else ".pdf"
        if actividad.get("tipo_archivo") == "PDF"
        else ""
    )

    compatibles = [
        nombre
        for nombre in archivos
        if os.path.splitext(nombre)[1].lower() == extension_objetivo
    ]

    if not compatibles:
        return "ERROR_RECURSO_NO_COINCIDE_CON_ACTIVIDAD"

    coincidencias = []
    for nombre in compatibles:
        puntuacion = puntuar_recurso(nombre, actividad)
        if puntuacion >= 100 or recurso_es_prefijo_de_actividad(nombre, actividad):
            coincidencias.append(nombre)

    if not coincidencias:
        return "ERROR_RECURSO_NO_COINCIDE_CON_ACTIVIDAD"

    if len(coincidencias) != 1:
        return "ERROR_RECURSO_AMBIGUO_EN_CARPETA_DRIVE"

    nombre_drive = coincidencias[0]
    actividad["nombre_recurso_drive"] = nombre_drive

    print(
        "✅ Carpeta Drive validada:",
        actividad["nombre"],
        "->",
        nombre_drive,
    )
    return None


# ============================================================
# LEER ACTIVIDADES DEL EXCEL
# ============================================================

def leer_actividades_excel(
    ruta_excel,
    procesar_ovi=True,
    procesar_ova=True,
    procesar_retos=True,
):

    libro = load_workbook(
        ruta_excel,
        data_only=True,
    )
    libro_formulas = load_workbook(
        ruta_excel,
        data_only=False,
    )

    actividades = []

    for nombre_hoja in libro.sheetnames:

        hoja = libro[
            nombre_hoja
        ]
        hoja_formulas = libro_formulas[
            nombre_hoja
        ]

        fila_cabecera = None

        # ----------------------------------------------------
        # BUSCAR CABECERA
        # ----------------------------------------------------

        for fila in range(
            1,
            min(
                hoja.max_row,
                30,
            ) + 1,
        ):

            valor = hoja.cell(
                row=fila,
                column=1,
            ).value

            if normalizar_texto(
                valor
            ) == normalizar_texto(
                "Semana correspondiente"
            ):

                fila_cabecera = fila
                break

        if fila_cabecera is None:
            continue

        programa = hoja.cell(
            row=1,
            column=1,
        ).value

        curso = hoja.cell(
            row=2,
            column=1,
        ).value

        programa = (
            str(programa).strip()
            if programa
            else ""
        )

        curso = (
            str(curso).strip()
            if curso
            else nombre_hoja
        )

        # ----------------------------------------------------
        # ACTIVIDADES
        # ----------------------------------------------------

        for fila in range(
            fila_cabecera + 1,
            hoja.max_row + 1,
        ):

            semana = hoja.cell(
                row=fila,
                column=1,
            ).value

            unidad = hoja.cell(
                row=fila,
                column=2,
            ).value

            nombre = hoja.cell(
                row=fila,
                column=4,
            ).value

            categoria = hoja.cell(
                row=fila,
                column=5,
            ).value

            tipo_recurso = hoja.cell(
                row=fila,
                column=6,
            ).value

            celda_enlace = hoja.cell(
                row=fila,
                column=7,
            )
            celda_enlace_formula = hoja_formulas.cell(
                row=fila,
                column=7,
            )
            enlace = celda_enlace.value
            url_recurso = extraer_url_celda_recurso(
                celda_enlace,
                celda_enlace_formula,
            )

            if not nombre:
                continue

            if not categoria:
                continue

            categoria_prizma = normalizar_categoria(
                categoria
            )

            if categoria_prizma not in [
                "OVI",
                "OVA",
                "CHALLENGE",
            ]:
                continue

            if (
                categoria_prizma == "OVI"
                and not procesar_ovi
            ):
                continue

            if (
                categoria_prizma == "OVA"
                and not procesar_ova
            ):
                continue

            if (
                categoria_prizma == "CHALLENGE"
                and not procesar_retos
            ):
                continue

            tipo_archivo = determinar_tipo_archivo(
                categoria,
                tipo_recurso,
                enlace,
            )

            if tipo_archivo is None:
                continue

            actividades.append(
                {
                    "fila_excel": fila,
                    "hoja": nombre_hoja,
                    "programa": programa,
                    "curso": curso,

                    "semana": (
                        str(semana).strip()
                        if semana
                        else ""
                    ),

                    "unidad": (
                        str(unidad).strip()
                        if unidad
                        else ""
                    ),

                    "nombre": str(
                        nombre
                    ).strip(),

                    "categoria_prizma":
                        categoria_prizma,

                    "tipo_recurso": (
                        str(tipo_recurso).strip()
                        if tipo_recurso
                        else ""
                    ),

                    "tipo_archivo":
                        tipo_archivo,

                    "enlace": (
                        str(enlace).strip()
                        if enlace
                        else ""
                    ),

                    "url_recurso": url_recurso,
                }
            )

    return actividades


# ============================================================
# ÍNDICE DE RECURSOS
# ============================================================

def crear_indice_recursos(
    ruta_zip,
):

    indice = []

    with zipfile.ZipFile(
        ruta_zip,
        "r",
    ) as zip_ref:

        for miembro in zip_ref.namelist():

            if miembro.endswith("/"):
                continue

            nombre = os.path.basename(
                miembro
            )

            extension = os.path.splitext(
                nombre
            )[1].lower()

            if extension not in [
                ".h5p",
                ".pdf",
            ] + list(EXTENSIONES_HTML):
                continue

            # .htm se trata igual que .html (mismo tipo de recurso).
            if extension in EXTENSIONES_HTML:
                extension = ".html"

            indice.append(
                {
                    "nombre": nombre,
                    "extension": extension,
                    "zip": ruta_zip,
                    "miembro": miembro,
                }
            )

    return indice


# ============================================================
# PUNTUAR RECURSO
# ============================================================

def puntuar_recurso(
    nombre_archivo,
    actividad,
):

    nombre_sin_extension = os.path.splitext(
        os.path.basename(
            nombre_archivo
        )
    )[0]

    # Copias de Google Drive: "archivo (1).h5p".
    nombre_sin_extension = re.sub(
        r"\s*\(\d+\)$",
        "",
        nombre_sin_extension,
    ).strip()

    archivo_n = normalizar_texto(
        nombre_sin_extension
    )

    actividad_n = normalizar_texto(
        actividad["nombre"]
    )

    referencia_n = normalizar_texto(
        actividad["enlace"]
    )

    puntuacion = 0

    if archivo_n == actividad_n:
        puntuacion += 300

    if (
        actividad_n
        and actividad_n in archivo_n
    ):
        puntuacion += 180

    if (
        archivo_n
        and archivo_n in referencia_n
    ):
        puntuacion += 80

    if (
        referencia_n
        and referencia_n in archivo_n
    ):
        puntuacion += 120

    palabras_actividad = {
        palabra
        for palabra in actividad_n.split()
        if len(palabra) >= 4
    }

    palabras_archivo = set(
        archivo_n.split()
    )

    if palabras_actividad:

        # Tambien cuenta palabras cortadas por el paquete o por Drive:
        # "investigac" reconoce "investigacion" y "problemas" a "problema".
        comunes = sum(
            1
            for palabra in palabras_actividad
            if any(
                palabra == otra
                or (len(otra) >= 4 and palabra.startswith(otra))
                or otra.startswith(palabra)
                for otra in palabras_archivo
            )
        )

        porcentaje = (
            comunes
            / len(palabras_actividad)
        )

        puntuacion += int(
            porcentaje * 100
        )

    return puntuacion


# ============================================================
# GOOGLE DRIVE -> ZIP: VALIDACIÓN EXACTA ANTES DE PRIZMA
# ============================================================

def clave_nombre_archivo_exacto(nombre):
    return unicodedata.normalize(
        "NFC",
        os.path.basename(str(nombre or "")).strip(),
    ).casefold()


def prevalidar_recursos_drive_zip(
    ruta_excel,
    ruta_zip,
    procesar_ovi=True,
    procesar_ova=True,
    procesar_retos=True,
):
    """Valida todos los enlaces Drive y exige el mismo archivo dentro del ZIP."""

    actividades = leer_actividades_excel(
        ruta_excel,
        procesar_ovi,
        procesar_ova,
        procesar_retos,
    )
    indice_recursos = crear_indice_recursos(ruta_zip)

    errores = []
    recursos = []

    if not actividades:
        return {
            "ok": False,
            "total_actividades": 0,
            "validadas_drive": 0,
            "errores": ["No se encontraron actividades compatibles en la matriz."],
            "recursos": [],
        }

    if not indice_recursos:
        return {
            "ok": False,
            "total_actividades": len(actividades),
            "validadas_drive": 0,
            "errores": ["El ZIP no contiene archivos H5P o PDF."],
            "recursos": [],
        }

    for actividad in actividades:
        descripcion = (
            f'Hoja {actividad.get("hoja", "")} · '
            f'fila {actividad.get("fila_excel", "")} · '
            f'{actividad.get("nombre", "")}'
        )

        url_recurso = str(actividad.get("url_recurso") or "").strip()

        if not url_recurso:
            errores.append(descripcion + ": no tiene enlace de Google Drive.")
            continue

        if not extraer_id_carpeta_drive(url_recurso):
            errores.append(
                descripcion
                + ": el enlace no corresponde a una carpeta válida de Google Drive."
            )
            continue

        error_drive = validar_carpeta_drive_para_actividad(actividad)
        if error_drive:
            errores.append(descripcion + ": " + str(error_drive))
            continue

        nombre_drive = str(actividad.get("nombre_recurso_drive") or "").strip()
        if not nombre_drive:
            errores.append(
                descripcion + ": Google Drive no permitió determinar un archivo único."
            )
            continue

        clave_drive = clave_nombre_archivo_exacto(nombre_drive)
        exactos = [
            recurso
            for recurso in indice_recursos
            if clave_nombre_archivo_exacto(recurso.get("nombre")) == clave_drive
        ]

        if len(exactos) == 0:
            errores.append(
                descripcion
                + f': Drive indica "{nombre_drive}", pero ese archivo no está en el ZIP.'
            )
            continue

        if len(exactos) != 1:
            errores.append(
                descripcion
                + f': el archivo "{nombre_drive}" aparece más de una vez en el ZIP.'
            )
            continue

        recursos.append({
            "hoja": actividad.get("hoja", ""),
            "fila": actividad.get("fila_excel", 0),
            "actividad": actividad.get("nombre", ""),
            "archivo": nombre_drive,
        })

    return {
        "ok": len(errores) == 0 and len(recursos) == len(actividades),
        "total_actividades": len(actividades),
        "validadas_drive": len(recursos),
        "errores": errores,
        "recursos": recursos,
    }


# ============================================================
# RESOLVER RECURSO
# ============================================================

def archivo_parecido_a_actividad(
    actividad,
    indice_recursos,
    umbral=50,
):
    """Mejor archivo del ZIP parecido a la actividad aunque no llegue al
    umbral de seleccion.

    Sirve unicamente para explicar ERROR_RECURSO_NO_ENCONTRADO: decirle al
    usuario que archivo si hay en el ZIP en vez de solo "no se encontro".
    Devuelve "" si no hay nada medianamente parecido.
    """
    tipo = str(actividad.get("tipo_archivo") or "").upper()
    extension = {"H5P": ".h5p", "PDF": ".pdf", "HTML": ".html"}.get(tipo)
    if not extension:
        return ""

    mejor = ""
    mejor_puntaje = umbral
    for recurso in indice_recursos:
        if recurso.get("extension") != extension:
            continue
        puntaje = puntuar_recurso(recurso["nombre"], actividad)
        if recurso_es_prefijo_de_actividad(recurso["nombre"], actividad):
            puntaje = max(puntaje, 99)
        if puntaje >= mejor_puntaje:
            mejor_puntaje = puntaje
            mejor = recurso["nombre"]
    return mejor


def resolver_recurso(
    actividad,
    indice_recursos,
    carpeta_temp,
    actividades_curso,
):

    print()
    print("--------------------------------------")
    print("BUSCANDO RECURSO")
    print("--------------------------------------")

    print(
        "Actividad:",
        actividad["nombre"],
    )

    print(
        "Tipo esperado:",
        actividad["tipo_archivo"],
    )

    print(
        "Referencia:",
        actividad["enlace"],
    )

    if actividad[
        "tipo_archivo"
    ] == "H5P":

        extension = ".h5p"

    elif actividad[
        "tipo_archivo"
    ] == "PDF":

        extension = ".pdf"

    elif actividad[
        "tipo_archivo"
    ] == "HTML":

        extension = ".html"

    else:

        return (
            None,
            "ERROR_TIPO_RECURSO_NO_SOPORTADO",
        )

    candidatos = []

    nombre_drive = str(
        actividad.get("nombre_recurso_drive")
        or ""
    ).strip()

    if nombre_drive:
        clave_drive = clave_nombre_archivo_exacto(nombre_drive)

        for recurso in indice_recursos:
            if recurso["extension"] != extension:
                continue

            if clave_nombre_archivo_exacto(recurso["nombre"]) != clave_drive:
                continue

            candidato = recurso.copy()
            candidato["puntuacion"] = 10000
            candidato["metodo"] = "DRIVE_EXACTO"
            candidatos.append(candidato)

        if len(candidatos) == 0:
            return (
                None,
                "ERROR_RECURSO_DRIVE_NO_ESTA_EN_ZIP",
            )

        if len(candidatos) != 1:
            return (
                None,
                "ERROR_RECURSO_DUPLICADO",
            )

        print(
            "Coincidencia exacta Drive -> ZIP:",
            nombre_drive,
        )

    else:
        for recurso in indice_recursos:

            if recurso[
                "extension"
            ] != extension:
                continue

            puntuacion = puntuar_recurso(
                recurso["nombre"],
                actividad,
            )

            if puntuacion < 100:
                continue

            candidato = recurso.copy()

            candidato[
                "puntuacion"
            ] = puntuacion

            candidatos.append(
                candidato
            )

    print(
        "Candidatos:",
        len(candidatos),
    )

    # --------------------------------------------------------
    # SEGUNDO MÉTODO SEGURO PARA NOMBRES CORTOS / TRUNCADOS
    # --------------------------------------------------------
    # Se usa únicamente cuando el método normal no encontró
    # ningún candidato con puntuación >= 100.
    #
    # Ejemplos admitidos:
    # "Dimensiones.h5p"
    # "Dimensiones e indicadores.h5p"
    # "... principios de legalid.h5p"
    #
    # No adivina: el prefijo del archivo debe identificar una
    # única actividad del curso para ese tipo de archivo.
    # --------------------------------------------------------

    if not candidatos:

        candidatos_cortos = []
        hubo_ambiguos = False
        recursos_ambiguos = []

        for recurso in indice_recursos:

            if recurso[
                "extension"
            ] != extension:
                continue

            if not recurso_es_prefijo_de_actividad(
                recurso["nombre"],
                actividad,
            ):
                continue

            if not recurso_corto_es_unico_para_actividad(
                recurso["nombre"],
                actividad,
                actividades_curso,
            ):

                hubo_ambiguos = True
                recursos_ambiguos.append(
                    recurso
                )
                continue

            candidato = recurso.copy()

            candidato[
                "puntuacion"
            ] = puntuar_recurso(
                recurso["nombre"],
                actividad,
            )

            candidato[
                "metodo"
            ] = "PREFIJO_UNICO"

            candidatos_cortos.append(
                candidato
            )

        print(
            "Candidatos por prefijo único:",
            len(candidatos_cortos),
        )

        for candidato in candidatos_cortos:

            print(
                "  →",
                candidato["nombre"],
                "| PREFIJO_UNICO",
            )

        if len(candidatos_cortos) == 0:

            if hubo_ambiguos:

                # Fallback excepcional SOLO para H5P ambiguos.
                # No se ejecuta para duplicados y no recorre todos los H5P
                # del curso: únicamente abre los recursos que ya quedaron
                # implicados en la ambigüedad del nombre externo.
                if extension == ".h5p":

                    candidato_interno = (
                        resolver_h5p_ambiguo_por_titulo_interno(
                            actividad,
                            recursos_ambiguos,
                        )
                    )

                    if candidato_interno is not None:
                        candidatos = [
                            candidato_interno
                        ]
                    else:
                        return (
                            None,
                            "ERROR_RECURSO_AMBIGUO",
                        )

                else:
                    return (
                        None,
                        "ERROR_RECURSO_AMBIGUO",
                    )

            else:
                referencia = str(actividad.get("enlace") or "").strip()
                parecido = archivo_parecido_a_actividad(
                    actividad,
                    indice_recursos,
                )
                partes = []
                if referencia:
                    partes.append(
                        f"el Excel pide '{referencia}' pero no está en el ZIP"
                    )
                if parecido:
                    partes.append(
                        f"hay un archivo parecido a la actividad: '{parecido}'"
                    )
                if partes:
                    return (
                        None,
                        "ERROR_RECURSO_NO_ENCONTRADO: "
                        + "; ".join(partes),
                    )
                return (
                    None,
                    "ERROR_RECURSO_NO_ENCONTRADO",
                )

        if candidatos:
            pass
        elif len(candidatos_cortos) > 1:

            return (
                None,
                "ERROR_RECURSO_DUPLICADO",
            )
        else:
            candidatos = candidatos_cortos

        # Si el fallback interno ya resolvió la ambigüedad, no reemplazarlo.
        if candidatos and candidatos[0].get("metodo") == "TITULO_INTERNO_H5P_EXACTO":
            pass
        elif len(candidatos_cortos) > 1:
            return (
                None,
                "ERROR_RECURSO_DUPLICADO",
            )
        elif candidatos_cortos:
            candidatos = candidatos_cortos

    mejor_puntuacion = max(
        candidato["puntuacion"]
        for candidato in candidatos
    )

    mejores = [
        candidato
        for candidato in candidatos
        if candidato["puntuacion"]
        == mejor_puntuacion
    ]

    print(
        "Mejor puntuación:",
        mejor_puntuacion,
    )

    for candidato in mejores:

        print(
            "  →",
            candidato["nombre"],
            "|",
            candidato["puntuacion"],
        )

    # SEGURIDAD: NO ADIVINAR
    if len(mejores) != 1:

        return (
            None,
            "ERROR_RECURSO_DUPLICADO",
        )

    elegido = mejores[0]

    carpeta_fila = os.path.join(
        carpeta_temp,
        "fila_" + str(
            actividad["fila_excel"]
        ),
    )

    os.makedirs(
        carpeta_fila,
        exist_ok=True,
    )

    ruta_destino = os.path.join(
        carpeta_fila,
        elegido["nombre"],
    )

    try:

        with zipfile.ZipFile(
            elegido["zip"],
            "r",
        ) as zip_ref:

            with zip_ref.open(
                elegido["miembro"]
            ) as origen:

                with open(
                    ruta_destino,
                    "wb",
                ) as destino:

                    shutil.copyfileobj(
                        origen,
                        destino,
                    )

    except Exception:

        return (
            None,
            "ERROR_EXTRAYENDO_RECURSO",
        )

    if not os.path.isfile(
        ruta_destino
    ):

        return (
            None,
            "ERROR_EXTRAYENDO_RECURSO",
        )

    print()
    print(
        "✅ Recurso seleccionado:",
        elegido["nombre"],
    )

    return (
        {
            "ruta":
                ruta_destino,

            "nombre_original":
                elegido["nombre"],

            "puntuacion":
                mejor_puntuacion,
        },
        None,
    )


# ============================================================
# REPORTE
# ============================================================

def guardar_resultado(
    ruta_reporte,
    actividad,
    resultado,
    observacion,
    recurso="",
):

    # Acumula el estado del curso (data/estado_cursos.json): que actividades
    # ya quedaron en PRIZMA y cual es la causa de las que faltan.
    registrar_resultado_actividad(actividad, resultado, observacion, recurso)

    # Motivo en lenguaje humano solo cuando la observacion es un codigo
    # ERROR_* (si ya es texto libre, se entiende sola). El codigo original se
    # conserva en la columna Observación (diagnostico y compatibilidad).
    motivo = ""
    if str(resultado).strip().upper() != "OK" and es_codigo_error(observacion):
        motivo = mensaje_error(observacion)

    existe = os.path.isfile(
        ruta_reporte
    )

    with open(
        ruta_reporte,
        "a",
        newline="",
        encoding="utf-8-sig",
    ) as archivo:

        escritor = csv.writer(
            archivo
        )

        if not existe:

            escritor.writerow(
                [
                    "Fila Excel",
                    "Hoja",
                    "Programa",
                    "Curso",
                    "Semana",
                    "Unidad",
                    "Actividad",
                    "Categoría",
                    "Tipo",
                    "Recurso",
                    "Resultado",
                    "Observación",
                    "Motivo",
                ]
            )

        escritor.writerow(
            [
                actividad["fila_excel"],
                actividad.get("hoja", ""),
                actividad["programa"],
                actividad["curso"],
                actividad["semana"],
                actividad["unidad"],
                actividad["nombre"],
                actividad["categoria_prizma"],
                actividad["tipo_archivo"],
                recurso,
                resultado,
                observacion,
                motivo,
            ]
        )


# ============================================================
# OVERLAY POST-GUARDADO
# ============================================================

def desbloquear_interfaz_post_guardado(
    pagina,
):
    """Cierra el overlay solo si realmente existe y esta visible.

    Antes esta funcion agregaba 1.1 s de espera y un clic ciego incluso
    cuando no habia overlay. En el camino normal ahora retorna de inmediato.
    """
    overlay = pagina.locator(
        "div.MuiBox-root.css-15m6u24"
    )

    try:
        for i in range(overlay.count()):
            elemento = overlay.nth(i)
            if not elemento.is_visible():
                continue

            print("✅ Overlay detectado.")

            try:
                elemento.click(
                    position={"x": 5, "y": 5},
                    force=True,
                )
            except Exception:
                pagina.mouse.click(720, 450)

            # Pequena pausa solo despues de haber cerrado un overlay real.
            pagina.wait_for_timeout(120)
            return True
    except Exception:
        pass

    return True


# ============================================================
# CANCELAR EDICIÓN
# ============================================================

def cancelar_edicion_segura(
    pagina,
):

    try:

        boton = pagina.get_by_role(
            "button",
            name="Cancelar",
            exact=True,
        )

        if boton.count() == 0:
            return False

        if not boton.first.is_visible():
            return False

        boton.first.click(
            timeout=10000
        )

        pagina.wait_for_timeout(
            900
        )

        desbloquear_interfaz_post_guardado(
            pagina
        )

        return True

    except Exception:

        return False


# ============================================================
# ASEGURAR LISTADO
# ============================================================

def asegurar_listado(
    pagina,
):

    desbloquear_interfaz_post_guardado(
        pagina
    )

    buscador = pagina.locator(
        'input[placeholder="Buscar..."]'
    )

    try:

        if (
            buscador.count() > 0
            and buscador.first.is_visible()
        ):
            return True

    except Exception:
        pass

    try:

        cancelar = pagina.get_by_role(
            "button",
            name="Cancelar",
            exact=True,
        )

        if (
            cancelar.count() > 0
            and cancelar.first.is_visible()
        ):

            cancelar.first.click(
                timeout=10000
            )

            pagina.wait_for_timeout(
                1000
            )

            desbloquear_interfaz_post_guardado(
                pagina
            )

    except Exception:
        pass

    try:

        buscador.wait_for(
            state="visible",
            timeout=10000,
        )

        return True

    except Exception:

        return False


# ============================================================
# ENTRAR A CATEGORÍA
# ============================================================

def entrar_categoria(
    pagina,
    categoria,
    estado_navegacion=None,
):

    nombres_pestana = {
        "OVI": "OVI",
        "OVA": "OVA",
        "CHALLENGE": "Retos Evaluativos",
    }

    texto_pestana = nombres_pestana.get(
        categoria
    )

    if not texto_pestana:
        return False

    if (
        isinstance(estado_navegacion, dict)
        and estado_navegacion.get("categoria_actual") == categoria
    ):
        return True

    try:

        desbloquear_interfaz_post_guardado(
            pagina
        )

        pestana = pagina.locator(
            "label.tab",
            has_text=texto_pestana,
        ).first

        pestana.wait_for(
            state="visible",
            timeout=30000,
        )

        pestana.click(
            timeout=10000
        )

        pagina.wait_for_timeout(
            1200
        )

        if isinstance(estado_navegacion, dict):
            estado_navegacion["categoria_actual"] = categoria

        return True

    except Exception:

        return False


# ============================================================
# ENCONTRAR FILA
# ============================================================

def encontrar_fila_desde_elemento(
    elemento,
):

    for nivel in range(
        1,
        10,
    ):

        try:

            padre = elemento.locator(
                "xpath=" + "/.." * nivel
            )

            if padre.locator(
                "button"
            ).count() == 3:

                return padre

        except Exception:
            pass

    return None


def obtener_id_temporal_fila(
    fila,
):

    try:

        return fila.evaluate(
            """
            (element) => {

                if (!element.dataset.prizmaScanId) {

                    element.dataset.prizmaScanId =
                        "scan-" +
                        Math.random()
                        .toString(36)
                        .slice(2);
                }

                return element.dataset.prizmaScanId;
            }
            """
        )

    except Exception:

        return None


def obtener_nombre_fila(
    texto,
):

    lineas = [
        linea.strip()
        for linea in texto.splitlines()
        if linea.strip()
    ]

    if len(lineas) < 2:
        return ""

    return lineas[1]


# ============================================================
# ANALIZAR RESULTADOS
# ============================================================

def analizar_resultados_pagina(
    pagina,
    actividad,
):

    elementos = pagina.get_by_text(
        actividad["nombre"],
        exact=False,
    )

    filas = {}

    for i in range(
        elementos.count()
    ):

        try:

            elemento = elementos.nth(i)

            if not elemento.is_visible():
                continue

            fila = encontrar_fila_desde_elemento(
                elemento
            )

            if fila is None:
                continue

            identificador = obtener_id_temporal_fila(
                fila
            )

            if identificador:

                filas[
                    identificador
                ] = fila

        except Exception:
            pass

    nombre_objetivo = normalizar_texto(
        actividad["nombre"]
    )

    semana_objetivo = normalizar_texto(
        actividad["semana"]
    )

    unidad_objetivo = normalizar_texto(
        actividad["unidad"]
    )

    variantes_programa = obtener_variantes_programa(
        actividad["programa"]
    )

    resultados = []

    for fila in filas.values():

        try:

            texto = fila.inner_text().strip()

        except Exception:
            continue

        texto_n = normalizar_texto(
            texto
        )

        nombre_fila = normalizar_texto(
            obtener_nombre_fila(
                texto
            )
        )

        cumple_nombre = (
            nombre_fila
            == nombre_objetivo
        )

        cumple_semana = (
            semana_objetivo
            in texto_n
        )

        cumple_unidad = (
            unidad_objetivo
            in texto_n
        )

        cumple_programa = any(
            variante
            and variante in texto_n
            for variante in variantes_programa
        )

        if actividad[
            "categoria_prizma"
        ] == "OVI":

            cumple_categoria = (
                "ovi" in texto_n
            )

        elif actividad[
            "categoria_prizma"
        ] == "OVA":

            # En OVA la fila puede no mostrar literalmente la categoría
            # dentro de su texto. La categoría ya quedó validada al
            # entrar explícitamente a la pestaña OVA.
            cumple_categoria = True

        elif actividad[
            "categoria_prizma"
        ] == "CHALLENGE":

            # En Retos Evaluativos la fila no muestra la categoría
            # dentro de su texto. La categoría ya quedó validada al
            # entrar explícitamente a la pestaña Retos Evaluativos.
            cumple_categoria = True

        else:

            cumple_categoria = False

        coincide = (
            cumple_nombre
            and cumple_semana
            and cumple_unidad
            and cumple_programa
            and cumple_categoria
        )

        resultados.append(
            {
                "fila": fila,
                "texto": texto,
                "coincide": coincide,
            }
        )

    return resultados


def obtener_firma_pagina(
    pagina,
    actividad,
):

    resultados = analizar_resultados_pagina(
        pagina,
        actividad,
    )

    return "||".join(
        sorted(
            normalizar_texto(
                resultado["texto"]
            )
            for resultado in resultados
        )
    )


# ============================================================
# PAGINACIÓN
# ============================================================

def obtener_paginas_numericas(
    pagina,
):

    numeros = set()

    botones = pagina.locator(
        "button"
    )

    for i in range(
        botones.count()
    ):

        try:

            boton = botones.nth(i)

            if not boton.is_visible():
                continue

            texto = boton.inner_text().strip()

            if texto.isdigit():

                numero = int(
                    texto
                )

                if 1 <= numero <= 200:

                    numeros.add(
                        numero
                    )

            aria = (
                boton.get_attribute(
                    "aria-label"
                )
                or ""
            )

            coincidencia = re.search(
                r"(?:page|página|pagina)\s*(\d+)",
                aria,
                flags=re.IGNORECASE,
            )

            if coincidencia:

                numero = int(
                    coincidencia.group(1)
                )

                if 1 <= numero <= 200:

                    numeros.add(
                        numero
                    )

        except Exception:
            pass

    return sorted(
        numeros
    )


def ir_a_pagina_numero(
    pagina,
    numero,
):

    botones = pagina.locator(
        "button"
    )

    # POR TEXTO
    for i in range(
        botones.count()
    ):

        try:

            boton = botones.nth(i)

            if not boton.is_visible():
                continue

            texto = boton.inner_text().strip()

            if texto != str(
                numero
            ):
                continue

            boton.click(
                timeout=5000
            )

            pagina.wait_for_timeout(
                1200
            )

            return True

        except Exception:
            pass

    # POR ARIA
    for i in range(
        botones.count()
    ):

        try:

            boton = botones.nth(i)

            if not boton.is_visible():
                continue

            aria = normalizar_texto(
                boton.get_attribute(
                    "aria-label"
                )
            )

            if (
                f"page {numero}" not in aria
                and
                f"pagina {numero}" not in aria
            ):
                continue

            boton.click(
                timeout=5000
            )

            pagina.wait_for_timeout(
                1200
            )

            return True

        except Exception:
            pass

    return False


def encontrar_boton_siguiente(
    pagina,
):

    botones = pagina.locator(
        "button"
    )

    for i in range(
        botones.count()
    ):

        try:

            boton = botones.nth(i)

            if not boton.is_visible():
                continue

            if not boton.is_enabled():
                continue

            texto = normalizar_texto(
                boton.inner_text()
            )

            aria = normalizar_texto(
                boton.get_attribute(
                    "aria-label"
                )
            )

            title = normalizar_texto(
                boton.get_attribute(
                    "title"
                )
            )

            combinado = (
                texto
                + " "
                + aria
                + " "
                + title
            )

            if any(
                valor in combinado
                for valor in [
                    "next page",
                    "pagina siguiente",
                    "siguiente",
                    "go to next page",
                ]
            ):

                return boton

        except Exception:
            pass

    return None


def asegurar_pagina_1(
    pagina,
):

    paginas = obtener_paginas_numericas(
        pagina
    )

    if 1 not in paginas:
        return True

    return ir_a_pagina_numero(
        pagina,
        1,
    )


# ============================================================
# TÉRMINOS DE BÚSQUEDA
# ============================================================

def obtener_terminos_busqueda_general(
    nombre,
):

    nombre_original = str(
        nombre or ""
    ).strip()

    if not nombre_original:
        return []

    terminos = [
        nombre_original
    ]

    nombre_n = normalizar_texto(
        nombre_original
    )

    palabras = [
        palabra
        for palabra in nombre_n.split()
        if len(palabra) >= 3
        and palabra not in {
            "del", "las", "los", "una", "uno",
            "para", "por", "con", "sin", "desde",
            "entre", "sobre", "sus", "que", "como",
        }
    ]

    # PRIZMA a veces no devuelve registros con el título completo,
    # pero sí con una frase corta del mismo nombre. La fila SOLO se
    # acepta después si nombre + semana + unidad + programa + categoría
    # coinciden exactamente, así que estos términos amplían la búsqueda
    # sin volverla ambigua.
    candidatos = []

    if len(palabras) >= 2:
        candidatos.append(
            " ".join(palabras[-2:])
        )

    if len(palabras) >= 4:
        medio = max(0, len(palabras) // 2 - 1)
        candidatos.append(
            " ".join(palabras[medio:medio + 2])
        )

    if len(palabras) >= 2:
        candidatos.append(
            " ".join(palabras[:2])
        )

    for candidato in candidatos:
        if (
            candidato
            and normalizar_texto(candidato) != nombre_n
            and candidato not in terminos
        ):
            terminos.append(candidato)

    return terminos


def obtener_terminos_busqueda_reto(
    nombre,
):

    nombre_original = str(
        nombre or ""
    ).strip()

    if not nombre_original:
        return []

    terminos = [
        nombre_original
    ]

    # Fallback específico para Retos Evaluativos. PRIZMA puede mostrar
    # la fila al repetir la búsqueda quitando una o dos letras finales.
    # La fila solo se acepta después con validación exacta de nombre,
    # semana, unidad, programa y categoría.
    for recorte in (1, 2):
        if len(nombre_original) > recorte + 4:
            truncado = nombre_original[:-recorte].rstrip()
            if truncado and truncado not in terminos:
                terminos.append(truncado)

    nombre_n = normalizar_texto(
        nombre_original
    )

    # Algunos Retos Evaluativos no aparecen cuando el buscador de
    # PRIZMA recibe el título completo. Como fallback usamos frases
    # más cortas, pero la aceptación de la fila sigue siendo EXACTA
    # en analizar_resultados_pagina().
    palabras = [
        palabra
        for palabra in nombre_n.split()
        if len(palabra) >= 4
        and palabra not in {
            "analisis",
            "modelo",
            "modelos",
            "sobre",
            "para",
            "entre",
            "desde",
        }
    ]

    if palabras:
        # Preferimos las últimas palabras porque suelen ser las más
        # distintivas del reto (p. ej. "modelo de Piaget").
        frase_final = " ".join(
            palabras[-3:]
        )

        if (
            frase_final
            and frase_final not in terminos
        ):
            terminos.append(
                frase_final
            )

        ultima = palabras[-1]

        if (
            len(ultima) >= 5
            and ultima not in terminos
        ):
            terminos.append(
                ultima
            )

    return terminos


# ============================================================
# ESPERAS ADAPTATIVAS DE RENDIMIENTO
# ============================================================

def _esperar_resultados_busqueda(
    pagina,
    actividad,
    timeout_ms=2600,
    intervalo_ms=100,
    firma_previa=None,
):
    """Espera hasta que aparezcan resultados y la firma quede estable.

    firma_previa es la pantalla que habia ANTES de aplicar el filtro.
    Mientras la firma siga siendo esa, el filtro todavia no se aplico y
    seguir adelante significaria leer datos viejos. Ese era el origen de
    ERROR_RECUPERANDO_FILA cuando PRIZMA respondia lento.

    Si no aparece ningun resultado conserva el timeout completo, evitando
    falsos negativos por una respuesta lenta de PRIZMA.
    """
    transcurrido = 0
    firma_anterior = None
    estables = 0

    while True:
        firma = obtener_firma_pagina(pagina, actividad)

        if firma and firma != firma_previa:
            if firma == firma_anterior:
                estables += 1
            else:
                estables = 0

            if estables >= 2:
                return firma

        firma_anterior = firma

        if transcurrido >= timeout_ms:
            return firma or ""

        espera = min(intervalo_ms, timeout_ms - transcurrido)
        if espera <= 0:
            return firma or ""

        pagina.wait_for_timeout(espera)
        transcurrido += espera


def _esperar_cambio_firma(
    pagina,
    actividad,
    firma_anterior,
    timeout_ms=1200,
    intervalo_ms=100,
):
    """Espera un cambio real de pagina sin pagar 1.2 s fijos."""
    transcurrido = 0

    while True:
        firma = obtener_firma_pagina(pagina, actividad)
        if firma != firma_anterior:
            return True

        if transcurrido >= timeout_ms:
            return False

        espera = min(intervalo_ms, timeout_ms - transcurrido)
        if espera <= 0:
            return False

        pagina.wait_for_timeout(espera)
        transcurrido += espera


def _recuperar_fila_con_filtro(
    pagina,
    actividad,
    buscador,
    termino,
    timeout_ms,
):
    """Vuelve a filtrar y recorre las paginas hasta ver la fila exacta."""
    try:
        buscador.fill("")
        pagina.wait_for_timeout(400)
        asegurar_pagina_1(pagina)
        firma_sin_filtro = obtener_firma_pagina(pagina, actividad)
        buscador.fill(termino)
        _esperar_resultados_busqueda(
            pagina, actividad, timeout_ms=timeout_ms, firma_previa=firma_sin_filtro
        )
        asegurar_pagina_1(pagina)
    except Exception:
        return None

    firmas_vistas = set()
    numero_pagina = 1
    while numero_pagina <= 50:
        try:
            firma_actual = obtener_firma_pagina(pagina, actividad)
        except Exception:
            return None
        if firma_actual in firmas_vistas:
            break
        firmas_vistas.add(firma_actual)
        try:
            resultados = analizar_resultados_pagina(pagina, actividad)
        except Exception:
            return None
        finales = [resultado for resultado in resultados if resultado["coincide"]]
        if len(finales) > 1:
            return "DUPLICADA"
        if len(finales) == 1:
            print("✅ Fila recuperada por escaneo.")
            return finales[0]["fila"]
        siguiente_numero = numero_pagina + 1
        try:
            paginas = obtener_paginas_numericas(pagina)
        except Exception:
            break
        avanzo = False
        if siguiente_numero in paginas:
            try:
                avanzo = ir_a_pagina_numero(pagina, siguiente_numero)
            except Exception:
                avanzo = False
        if not avanzo:
            try:
                siguiente = encontrar_boton_siguiente(pagina)
            except Exception:
                siguiente = None
            if siguiente is None:
                break
            try:
                siguiente.click(timeout=5000)
                _esperar_cambio_firma(
                    pagina, actividad, firma_actual, timeout_ms=1500
                )
                avanzo = True
            except Exception:
                break
        if not avanzo:
            break
        try:
            firma_nueva = obtener_firma_pagina(pagina, actividad)
        except Exception:
            break
        if firma_nueva == firma_actual:
            break
        numero_pagina += 1
    return None


def _escanear_coincidencias_actuales(
    pagina,
    actividad,
    termino,
):
    """Recorre la paginación actual y devuelve solo coincidencias estrictas."""
    coincidencias = []
    claves = set()
    firmas_visitadas = set()
    numero_logico = 1

    while numero_logico <= 50:
        firma_actual = obtener_firma_pagina(
            pagina,
            actividad,
        )

        if firma_actual in firmas_visitadas:
            break
        firmas_visitadas.add(firma_actual)

        resultados = analizar_resultados_pagina(
            pagina,
            actividad,
        )

        print(
            "Página",
            numero_logico,
            "- candidatos:",
            len(resultados),
        )

        for resultado in resultados:
            if not resultado["coincide"]:
                continue

            clave = (
                numero_logico,
                normalizar_texto(resultado["texto"]),
            )
            if clave in claves:
                continue
            claves.add(clave)

            coincidencias.append(
                {
                    "pagina": numero_logico,
                    "texto": resultado["texto"],
                    "termino": termino,
                    "fila": resultado.get("fila"),
                }
            )

        siguiente_numero = numero_logico + 1
        paginas = obtener_paginas_numericas(pagina)

        if siguiente_numero in paginas:
            if ir_a_pagina_numero(pagina, siguiente_numero):
                firma_nueva = obtener_firma_pagina(
                    pagina,
                    actividad,
                )
                if firma_nueva != firma_actual:
                    numero_logico += 1
                    continue

        siguiente = encontrar_boton_siguiente(pagina)
        if siguiente is None:
            break

        try:
            siguiente.click(timeout=5000)
            _esperar_cambio_firma(
                pagina,
                actividad,
                firma_actual,
                timeout_ms=1500,
            )
        except Exception:
            break

        firma_nueva = obtener_firma_pagina(
            pagina,
            actividad,
        )
        if firma_nueva == firma_actual:
            break

        numero_logico += 1

    return coincidencias


def _buscar_termino_estable(
    pagina,
    buscador,
    actividad,
    termino,
):
    """Aplica un término y espera una vista estable antes de leerla."""
    print("Término de búsqueda:", termino)

    buscador.fill("")
    pagina.wait_for_timeout(400)
    asegurar_pagina_1(pagina)
    firma_sin_filtro = obtener_firma_pagina(
        pagina,
        actividad,
    )

    buscador.fill(termino)
    espera_busqueda = (
        4200
        if actividad["categoria_prizma"] == "CHALLENGE"
        else 2600
    )
    _esperar_resultados_busqueda(
        pagina,
        actividad,
        timeout_ms=espera_busqueda,
        firma_previa=firma_sin_filtro,
    )

    asegurar_pagina_1(pagina)
    pagina.wait_for_timeout(
        1000
        if actividad["categoria_prizma"] == "CHALLENGE"
        else 700
    )

    return _escanear_coincidencias_actuales(
        pagina,
        actividad,
        termino,
    )


def _buscar_exacta_dos_intentos(
    pagina,
    buscador,
    actividad,
):
    """Da dos oportunidades al título exacto antes de ampliar la búsqueda."""
    termino = str(actividad.get("nombre") or "").strip()
    primero = _buscar_termino_estable(
        pagina,
        buscador,
        actividad,
        termino,
    )
    if primero:
        return primero

    try:
        pagina.wait_for_timeout(900)
    except Exception:
        pass

    return _buscar_termino_estable(
        pagina,
        buscador,
        actividad,
        termino,
    )


def _error_duplicada(coincidencias):
    """ERROR_ACTIVIDAD_DUPLICADA con el texto de las filas repetidas.

    Sin el detalle el usuario solo ve "esta duplicada" y no sabe ni cuantas
    veces ni donde: con el texto que PRIZMA muestra (nombre, semana,
    unidad) puede ubicarlas y borrar la que sobra.
    """
    textos = []
    for coincidencia in coincidencias[:3]:
        texto = re.sub(
            r"\s+",
            " ",
            str(coincidencia.get("texto") or ""),
        ).strip()
        if texto and texto not in textos:
            textos.append(texto[:200])
    if not textos:
        return "ERROR_ACTIVIDAD_DUPLICADA"
    return (
        "ERROR_ACTIVIDAD_DUPLICADA: aparece "
        f"{len(coincidencias)} veces en PRIZMA — "
        + " | ".join(f"«{texto}»" for texto in textos)
    )


# ============================================================
# BUSCAR ACTIVIDAD
# ============================================================

def _tabla_actividades_vacia(pagina):
    """Detecta si el listado de actividades está completamente vacío.

    Sirve para distinguir "el curso no tiene la plantilla precargada" de
    "el nombre de la actividad no coincide". Cuenta los contenedores con
    exactamente 3 botones (misma heurística que encontrar_fila_desde_elemento)
    sin filtrar por nombre: si no hay ninguno, no hay actividades en la vista.

    Devuelve None si no se pudo medir (para no dar falsos positivos).
    """
    try:
        return pagina.evaluate(
            """() => {
                const candidatos = document.querySelectorAll('*');
                let filas = 0;
                for (const el of candidatos) {
                    if (el.querySelectorAll('button').length !== 3) continue;
                    // Nos quedamos con el más interno: si un descendiente
                    // también tiene 3 botones, este es un ancestro del mismo.
                    let interno = true;
                    for (const hijo of el.querySelectorAll('*')) {
                        if (hijo.querySelectorAll('button').length === 3) {
                            interno = false;
                            break;
                        }
                    }
                    if (interno) filas += 1;
                }
                return filas;
            }"""
        )
    except Exception:
        return None


def buscar_actividad_correcta(
    pagina,
    actividad,
):

    print()
    print(
        "Buscando actividad:",
        actividad["nombre"],
    )
    print(
        "Programa:",
        actividad["programa"],
    )
    print(
        "Semana:",
        actividad["semana"],
    )
    print(
        "Unidad:",
        actividad["unidad"],
    )
    print(
        "Categoría:",
        actividad["categoria_prizma"],
    )

    buscador = pagina.locator(
        'input[placeholder="Buscar..."]'
    )
    buscador.wait_for(
        state="visible",
        timeout=30000,
    )

    nombre_exacto = str(
        actividad.get("nombre") or ""
    ).strip()

    if actividad["categoria_prizma"] == "CHALLENGE":
        terminos_busqueda = obtener_terminos_busqueda_reto(
            nombre_exacto
        )
    else:
        terminos_busqueda = obtener_terminos_busqueda_general(
            nombre_exacto
        )

    # El título exacto recibe dos oportunidades estables antes de ampliar
    # la búsqueda. Esto evita falsos negativos cuando PRIZMA todavía muestra
    # la tabla anterior al momento de leerla.
    coincidencias = _buscar_exacta_dos_intentos(
        pagina,
        buscador,
        actividad,
    )
    termino_encontrado = nombre_exacto if coincidencias else None

    if len(coincidencias) > 1:
        print("Coincidencias exactas:", len(coincidencias))
        # La tabla puede seguir re-renderizandose (carga asincrona) y la
        # misma fila salir escaneada en dos paginas: un re-escaneo decide
        # si la duplicidad es real antes de fallar.
        reintento = _buscar_termino_estable(
            pagina,
            buscador,
            actividad,
            nombre_exacto,
        )
        print("Coincidencias tras re-escaneo:", len(reintento))
        if len(reintento) > 0:
            coincidencias = reintento
            termino_encontrado = nombre_exacto

    # Solo si los dos intentos exactos quedaron en cero usamos los fallbacks
    # existentes. La validación de la fila continúa siendo estricta.
    if not coincidencias:
        exacto_normalizado = normalizar_texto(nombre_exacto)
        for termino_busqueda in terminos_busqueda:
            if normalizar_texto(termino_busqueda) == exacto_normalizado:
                continue

            coincidencias_intento = _buscar_termino_estable(
                pagina,
                buscador,
                actividad,
                termino_busqueda,
            )

            if len(coincidencias_intento) > 1:
                print("Coincidencias exactas:", len(coincidencias_intento))
                return (
                    None,
                    _error_duplicada(coincidencias_intento),
                )

            if len(coincidencias_intento) == 1:
                coincidencias = coincidencias_intento
                termino_encontrado = termino_busqueda
                break

    print(
        "Coincidencias exactas:",
        len(coincidencias),
    )

    if len(coincidencias) == 0:

        # Distingue dos causas muy distintas con el mismo síntoma:
        #   - El listado está vacío: el curso no tiene la plantilla precargada.
        #   - Hay filas pero ninguna coincide: el nombre no coincide con PRIZMA.
        filas_visibles = _tabla_actividades_vacia(pagina)

        if filas_visibles == 0:

            print(
                "❌ El listado de actividades está vacío: "
                "el curso no tiene la plantilla precargada."
            )

            return (
                None,
                "ERROR_CURSO_SIN_PLANTILLA",
            )

        print(
            "El listado tiene filas; el nombre de la actividad "
            "no coincide con ninguno."
        )

        return (
            None,
            "ERROR_ACTIVIDAD_NO_ENCONTRADA",
        )

    if len(coincidencias) > 1:
        return (
            None,
            _error_duplicada(coincidencias),
        )

    coincidencia = coincidencias[0]
    if termino_encontrado is None:
        termino_encontrado = coincidencia.get(
            "termino",
            nombre_exacto,
        )

    # Camino rápido y más estable: la fila que acabamos de validar sigue
    # visible en la misma vista, así que no repetimos la búsqueda.
    fila_directa = coincidencia.get("fila")
    if fila_directa is not None:
        try:
            if fila_directa.is_visible():
                print("✅ Fila exacta reutilizada sin nueva búsqueda.")
                return (
                    fila_directa,
                    None,
                )
        except Exception:
            pass

    # Si PRIZMA refrescó el DOM entre la lectura y el clic, recuperamos la
    # misma fila con el mecanismo robusto existente.
    terminos_recuperacion = [termino_encontrado]
    for termino in terminos_busqueda:
        if termino not in terminos_recuperacion:
            terminos_recuperacion.append(termino)

    espera_base = (
        4200
        if actividad["categoria_prizma"] == "CHALLENGE"
        else 2600
    )

    for intento in range(1, 4):
        espera_intento = espera_base + (intento - 1) * 2500
        for termino in terminos_recuperacion:
            resultado_recuperacion = _recuperar_fila_con_filtro(
                pagina,
                actividad,
                buscador,
                termino,
                espera_intento,
            )
            if resultado_recuperacion == "DUPLICADA":
                return (
                    None,
                    "ERROR_ACTIVIDAD_DUPLICADA",
                )
            if resultado_recuperacion is not None:
                return (
                    resultado_recuperacion,
                    None,
                )

        print(
            "Reintentando recuperacion de fila (intento",
            intento,
            "de 3)...",
        )
        try:
            pagina.wait_for_timeout(700 * intento)
        except Exception:
            pass

    return (
        None,
        "ERROR_RECUPERANDO_FILA",
    )

# ============================================================
# ABRIR EDICIÓN
# ============================================================

def abrir_edicion(
    pagina,
    fila,
):

    botones = fila.locator(
        "button"
    )

    if botones.count() != 3:

        print(
            "❌ Cantidad inesperada de botones:",
            botones.count(),
        )

        return False

    try:

        botones.nth(1).click(
            timeout=10000
        )

    except PlaywrightTimeoutError:

        desbloquear_interfaz_post_guardado(
            pagina
        )

        try:

            botones.nth(1).click(
                timeout=10000
            )

        except Exception:
            return False

    except Exception:
        return False

    pagina.wait_for_timeout(
        1800
    )

    return True


# ============================================================
# CAMPO PRINCIPAL "RECURSO"
# ============================================================
#
# TODOS LOS ARCHIVOS SE CARGAN EN RECURSO.
#
# H5P -> RECURSO
# PDF -> RECURSO
#
# MATERIAL DESCARGABLE NO SE TOCA.
# ============================================================

def detectar_campo_recurso(
    pagina,
):

    archivos = pagina.locator(
        'input[type="file"]'
    )

    candidatos_recurso = []

    print()
    print(
        "Campos file detectados:",
        archivos.count(),
    )

    for i in range(
        archivos.count()
    ):

        archivo = archivos.nth(i)

        accept = (
            archivo.get_attribute(
                "accept"
            )
            or ""
        )

        accept_n = accept.lower()

        print()
        print(
            f"Campo #{i + 1}:"
        )

        print(
            "accept =",
            accept,
        )

        # El campo Recurso admite .h5p.
        if ".h5p" in accept_n:

            candidatos_recurso.append(
                archivo
            )

            print(
                "→ RECURSO PRINCIPAL"
            )

            continue

        # Material descargable.
        if accept_n.strip() == ".pdf":

            print(
                "→ MATERIAL DESCARGABLE - IGNORADO"
            )

            continue

        print(
            "→ CAMPO NO UTILIZADO"
        )

    if len(
        candidatos_recurso
    ) == 0:

        print(
            "❌ No se encontró RECURSO."
        )

        return None

    if len(
        candidatos_recurso
    ) > 1:

        print(
            "❌ Varios candidatos para RECURSO."
        )

        return None

    print()
    print(
        "✅ Campo RECURSO identificado."
    )

    return candidatos_recurso[0]


# ============================================================
# DESCRIPCIÓN
# ============================================================

def limpiar_descripcion(
    pagina,
):
    """Borra la descripcion placeholder (NO_DISPONIBLE) del editor rico.

    El editor de PRIZMA es un contenteditable que puede tardar en estar
    listo o re-renderizarse tras el primer cargue, asi que un simple
    Ctrl+A + Backspace a veces falla o lanza excepcion (el elemento se
    desengancha del DOM). Por eso se guarda el INDICE del editor (no el
    handle), se reintenta y, como respaldo, se vacia via JavaScript
    disparando los eventos de input que escucha el framework.
    """

    editores = pagina.locator('[contenteditable="true"]')

    objetivos = [
        normalizar_texto(texto)
        for texto in TEXTOS_DESCRIPCION_A_BORRAR
    ]

    exactos = []  # indices, no handles: evita elementos colgados
    peligrosos = []

    for i in range(editores.count()):

        try:

            editor = editores.nth(i)

            if not editor.is_visible():
                continue

            texto_original = editor.inner_text().strip()
            texto_n = normalizar_texto(texto_original)

            if texto_n in objetivos:
                exactos.append(i)
                continue

            for objetivo in objetivos:
                if objetivo and objetivo in texto_n:
                    peligrosos.append(texto_original)

        except Exception:
            pass

    if peligrosos:
        return (False, "ERROR_DESCRIPCION_CONTENIDO_ADICIONAL")

    if len(exactos) == 0:
        return (True, None)

    if len(exactos) > 1:
        return (False, "ERROR_DESCRIPCION_MULTIPLE")

    indice = exactos[0]
    ultimo_problema = "no intentado"

    for _intento in range(3):

        try:
            # Re-resuelve el editor en cada intento: el DOM puede cambiar.
            editor = editores.nth(indice)
            editor.wait_for(state="visible", timeout=10000)
            editor.click()

            editor.press("Control+A")
            editor.press("Backspace")
            pagina.wait_for_timeout(500)

            if not normalizar_texto(editor.inner_text()):
                print("✅ NO_DISPONIBLE eliminado.")
                return (True, None)

            ultimo_problema = "quedó texto tras Backspace"

            # Respaldo: vaciar por JS y disparar eventos para que el
            # framework (Quill/Vue) se entere del cambio y lo persista.
            pagina.evaluate(
                """(idx) => {
                    const eds = document.querySelectorAll('[contenteditable="true"]');
                    const ed = eds[idx];
                    if (!ed) return;
                    ed.focus();
                    ed.innerHTML = '';
                    ed.dispatchEvent(new Event('input', {bubbles: true}));
                    ed.dispatchEvent(new Event('change', {bubbles: true}));
                }""",
                indice,
            )
            pagina.wait_for_timeout(500)

            editor = editores.nth(indice)
            if not normalizar_texto(editor.inner_text()):
                print("✅ NO_DISPONIBLE eliminado (respaldo JS).")
                return (True, None)

            ultimo_problema = "el editor se volvió a llenar"

        except Exception as e:
            ultimo_problema = f"excepción: {e}"
            pagina.wait_for_timeout(1000)

    print(f"⚠️ No se pudo limpiar la descripción: {ultimo_problema}")
    return (False, "ERROR_LIMPIANDO_DESCRIPCION")


# ============================================================
# RETOS EVALUATIVOS
# ============================================================

def es_programa_posgrado(
    programa,
):

    programa_n = normalizar_texto(
        programa
    )

    return (
        programa_n.startswith(
            "especializacion "
        )
        or programa_n == "especializacion"
        or programa_n.startswith(
            "esp "
        )
        or programa_n == "esp"
        or programa_n.startswith(
            "maestria "
        )
        or programa_n == "maestria"
    )


def obtener_numero_unidad(
    unidad,
):

    unidad_n = normalizar_texto(
        unidad
    )

    coincidencia = re.search(
        r"\b(\d+)\b",
        unidad_n,
    )

    if not coincidencia:
        return None

    return int(
        coincidencia.group(1)
    )


def obtener_configuracion_reto(
    actividad,
):

    numero_unidad = obtener_numero_unidad(
        actividad.get(
            "unidad",
            "",
        )
    )

    if numero_unidad is None:
        return (
            None,
            "ERROR_UNIDAD_RETO_NO_RECONOCIDA",
        )

    posgrado = es_programa_posgrado(
        actividad.get(
            "programa",
            "",
        )
    )

    if posgrado:

        reglas = {
            1: (
                "Corte 1",
                "Nivel intermedio",
            ),
            2: (
                "Corte 1",
                "Nivel intermedio",
            ),
            3: (
                "Corte 1",
                "Nivel avanzado",
            ),
            4: (
                "Corte 1",
                "Nivel avanzado",
            ),
        }

    else:

        reglas = {
            1: (
                "Corte 1",
                "Nivel básico",
            ),
            2: (
                "Corte 2",
                "Nivel intermedio",
            ),
            3: (
                "Corte 3",
                "Nivel avanzado",
            ),
        }

    configuracion = reglas.get(
        numero_unidad
    )

    if configuracion is None:

        if posgrado:
            return (
                None,
                "ERROR_UNIDAD_POSGRADO_NO_SOPORTADA",
            )

        return (
            None,
            "ERROR_UNIDAD_PREGRADO_NO_SOPORTADA",
        )

    return (
        {
            "corte": configuracion[0],
            "nivel": configuracion[1],
            "posgrado": posgrado,
            "numero_unidad": numero_unidad,
        },
        None,
    )


def seleccionar_autocomplete_reto(
    pagina,
    placeholder,
    valor,
):

    campo = pagina.locator(
        f'input[placeholder="{placeholder}"]'
    )

    if campo.count() != 1:
        return False

    campo = campo.first
    valor_n = normalizar_texto(
        valor
    )

    # Los autocompletes MUI de PRIZMA a veces tardan en montar sus
    # opciones. Hacemos hasta 3 intentos, pero solo aceptamos una
    # opción cuyo texto normalizado coincida EXACTAMENTE con el valor
    # esperado. Nunca elegimos la primera opción a ciegas.
    for intento in range(1, 4):

        try:
            campo.wait_for(
                state="visible",
                timeout=10000,
            )

            # Si el valor ya quedó seleccionado por un intento previo,
            # no lo tocamos nuevamente.
            try:
                valor_actual = normalizar_texto(
                    campo.input_value()
                )
                if valor_actual == valor_n:
                    return True
            except Exception:
                pass

            campo.click(
                timeout=10000
            )

            pagina.wait_for_timeout(
                500 + (intento * 300)
            )

            opciones = pagina.locator(
                '[role="option"]'
            )

            candidatas = []

            for indice in range(
                opciones.count()
            ):
                try:
                    opcion = opciones.nth(
                        indice
                    )

                    if not opcion.is_visible():
                        continue

                    texto_opcion = normalizar_texto(
                        opcion.inner_text()
                    )

                    if texto_opcion == valor_n:
                        candidatas.append(
                            opcion
                        )
                except Exception:
                    pass

            if len(candidatas) == 1:
                candidatas[0].click(
                    timeout=10000
                )

                pagina.wait_for_timeout(
                    700
                )

                valor_final = normalizar_texto(
                    campo.input_value()
                )

                if valor_final == valor_n:
                    return True

            # Algunos MUI filtran las opciones solamente cuando el
            # input recibe texto. Lo intentamos sin relajar la igualdad.
            try:
                campo.fill(
                    valor
                )

                pagina.wait_for_timeout(
                    800
                )

                opciones = pagina.locator(
                    '[role="option"]'
                )

                candidatas = []

                for indice in range(
                    opciones.count()
                ):
                    try:
                        opcion = opciones.nth(
                            indice
                        )

                        if not opcion.is_visible():
                            continue

                        if normalizar_texto(
                            opcion.inner_text()
                        ) == valor_n:
                            candidatas.append(
                                opcion
                            )
                    except Exception:
                        pass

                if len(candidatas) == 1:
                    candidatas[0].click(
                        timeout=10000
                    )

                    pagina.wait_for_timeout(
                        700
                    )

                    valor_final = normalizar_texto(
                        campo.input_value()
                    )

                    if valor_final == valor_n:
                        return True

            except Exception:
                pass

        except Exception:
            pass

        try:
            pagina.keyboard.press(
                "Escape"
            )
        except Exception:
            pass

        pagina.wait_for_timeout(
            350
        )

    return False


def detectar_campo_contenido_reto(
    pagina,
):

    archivos = pagina.locator(
        'input[type="file"]'
    )

    candidatos = []

    for i in range(
        archivos.count()
    ):

        try:

            archivo = archivos.nth(i)

            accept = (
                archivo.get_attribute(
                    "accept"
                )
                or ""
            ).lower()

            if ".pdf" not in accept:
                continue

            candidatos.append(
                archivo
            )

        except Exception:
            pass

    if len(candidatos) != 1:
        return None

    return candidatos[0]


# ============================================================
# BOTÓN GUARDAR
# ============================================================

def encontrar_boton_guardar(
    pagina,
):

    botones = pagina.get_by_role(
        "button",
        name="Editar",
        exact=True,
    )

    for i in range(
        botones.count()
    ):

        try:

            boton = botones.nth(i)

            if not boton.is_visible():
                continue

            if not boton.is_enabled():
                continue

            tipo = (
                boton.get_attribute(
                    "type"
                )
                or ""
            ).lower()

            if tipo == "submit":

                return boton

        except Exception:
            pass

    return None


# ============================================================
# ARCHIVO VISIBLE
# ============================================================

def archivo_visible(
    pagina,
    nombre_archivo,
):

    referencia = pagina.get_by_text(
        nombre_archivo,
        exact=False,
    )

    for i in range(
        referencia.count()
    ):

        try:

            if referencia.nth(i).is_visible():

                return True

        except Exception:
            pass

    return False


def _esperar_archivo_visible(
    pagina,
    nombre_archivo,
    timeout_ms=2200,
    intervalo_ms=100,
):
    """Termina en cuanto PRIZMA muestre el nombre del archivo cargado."""
    transcurrido = 0

    while True:
        if archivo_visible(pagina, nombre_archivo):
            return True

        if transcurrido >= timeout_ms:
            return False

        espera = min(intervalo_ms, timeout_ms - transcurrido)
        if espera <= 0:
            return False

        pagina.wait_for_timeout(espera)
        transcurrido += espera


# ============================================================
# PATCH
# ============================================================

# PRIZMA no permite guardar un H5P si el campo "Html" de la actividad tiene
# contenido y lo indica con este texto bajo el campo (no es un alert MUI).
# Politica TEMPORAL (config.PRIZMA_VACIAR_HTML_CON_H5P): mientras se use H5P
# el motor vacia el Html antes de guardar; con la politica desactivada cancela
# la edicion sin guardar y reporta ERROR_HTML_Y_H5P.
TEXTO_CONFLICTO_HTML_H5P = "Solo puedes subir un H5P o completar el campo Html"

ERROR_CONFLICTO_HTML_H5P = (
    "ERROR_HTML_Y_H5P: la actividad tiene contenido en el campo Html y PRIZMA "
    "no permite Html y H5P a la vez. No se guardo nada; vacia el campo Html "
    "en PRIZMA y vuelve a cargar esta actividad."
)

ERROR_LIMPIANDO_HTML = (
    "ERROR_LIMPIANDO_HTML: el campo Html tiene contenido y no se pudo vaciar "
    "automaticamente. No se guardo nada; vacia el campo Html en PRIZMA y "
    "vuelve a cargar esta actividad."
)

# Editor del campo Html: el primer textarea / contenteditable / editor de
# codigo (CodeMirror, Monaco) que aparece DESPUES del titulo "Html" y ANTES
# del titulo "Recurso ...". Asi nunca se toca la Descripcion ni otros campos.
_CLASE_XPATH = "contains(concat(' ', normalize-space(@class), ' '), ' {} ')"
_XPATH_EDITOR_HTML = (
    "xpath=//*[normalize-space(text())='Html']/following::*["
    "self::textarea or @contenteditable='true' or "
    + _CLASE_XPATH.format("cm-content") + " or "
    + _CLASE_XPATH.format("monaco-editor") + " or "
    + _CLASE_XPATH.format("CodeMirror")
    + "][following::*[starts-with(normalize-space(text()), 'Recurso')]][1]"
)
_XPATH_SECCION_HTML = "xpath=//*[normalize-space(text())='Html']/ancestor::*[3]"


def _conflicto_html_h5p_visible(pagina):
    """True si PRIZMA muestra el aviso de que no se pueden tener Html y H5P a la vez."""
    try:
        aviso = pagina.get_by_text(TEXTO_CONFLICTO_HTML_H5P, exact=False)
        for indice in range(min(aviso.count(), 5)):
            if aviso.nth(indice).is_visible():
                return True
    except Exception:
        pass
    return False


def _localizar_editor_html(pagina):
    """Editor del campo Html de la actividad abierta, o None si no se encuentra."""
    try:
        editor = pagina.locator(_XPATH_EDITOR_HTML)
        if editor.count() == 0:
            return None
        return editor.first
    except Exception:
        return None


def _texto_editor_html(editor):
    """Contenido actual del editor Html ("" si no se puede leer)."""
    if editor is None:
        return ""
    try:
        if _es_campo_de_texto(editor):
            return editor.input_value(timeout=1000) or ""
        return editor.inner_text(timeout=1000) or ""
    except Exception:
        return ""


def _vaciar_editor_html(pagina, editor, aviso_inicial, timeout_ms=3000, intervalo_ms=200):
    """Vacia el editor Html y verifica el resultado.

    Exito: el aviso de conflicto Html/H5P no esta visible y ademas el editor
    quedo vacio o el aviso estaba visible antes (y desaparecio). La segunda
    condicion cubre editores de codigo cuyo texto incluye numeros de linea.
    """
    try:
        if _es_campo_de_texto(editor):
            editor.fill("")
        else:
            editor.click(timeout=5000)
            pagina.keyboard.press("Control+A")
            pagina.keyboard.press("Backspace")
    except Exception as error:
        print("No se pudo vaciar el campo Html:", type(error).__name__, str(error)[:200])
        return False

    transcurrido = 0
    while True:
        vacio = not normalizar_texto(_texto_editor_html(editor))
        sin_aviso = not _conflicto_html_h5p_visible(pagina)
        if sin_aviso and (vacio or aviso_inicial):
            return True
        if transcurrido >= timeout_ms:
            return False
        pagina.wait_for_timeout(intervalo_ms)
        transcurrido += intervalo_ms


def _guardar_estructura_campo_html(pagina, carpeta_temp, actividad):
    """Guarda el HTML de la seccion "Html" del formulario para ajustar el motor.

    Solo se usa cuando algo falla. Queda en temp/ (no se versiona).
    """
    fila = str((actividad or {}).get("fila_excel") or "x")
    return guardar_estructura(
        pagina,
        _XPATH_SECCION_HTML,
        carpeta_temp,
        "campo_html_fila_" + fila + ".html",
    )


def _resolver_campo_html(pagina, carpeta_temp, actividad, nombre_recurso):
    """Deja el campo Html listo para guardar un H5P, segun la politica vigente.

    Devuelve (resultado_error, texto_vaciado):
      - (None, None): no habia contenido ni aviso; no se toco nada.
      - (None, "texto previo"): se vacio el Html (el texto previo va al reporte).
      - (dict, None): no se puede guardar; la edicion ya quedo cancelada.
    """
    aviso = _conflicto_html_h5p_visible(pagina)
    editor = _localizar_editor_html(pagina)
    texto_previo = _texto_editor_html(editor).strip()

    if not aviso and not normalizar_texto(texto_previo):
        return None, None

    if not PRIZMA_VACIAR_HTML_CON_H5P:
        if not aviso:
            # Sin aviso de PRIZMA y sin permiso para borrar: se intenta guardar
            # igual; si PRIZMA lo bloquea, se reportara el conflicto.
            return None, None
        print("Conflicto Html/H5P detectado: no se guarda esta actividad.")
        cancelar_edicion_segura(pagina)
        return {
            "ok": False,
            "error": ERROR_CONFLICTO_HTML_H5P,
            "recurso": nombre_recurso,
        }, None

    print("Campo Html con contenido; se vacia para poder guardar el H5P:", texto_previo[:120])

    if editor is not None and _vaciar_editor_html(pagina, editor, aviso):
        print("Campo Html vaciado.")
        return None, texto_previo

    error = ERROR_LIMPIANDO_HTML
    ruta_estructura = _guardar_estructura_campo_html(pagina, carpeta_temp, actividad)
    if ruta_estructura:
        error += "; estructura: " + ruta_estructura
    ruta_captura = _guardar_captura_diagnostico(pagina, carpeta_temp, actividad)
    if ruta_captura:
        error += "; captura: " + ruta_captura
    print(error)

    cancelar_edicion_segura(pagina)
    return {
        "ok": False,
        "error": error,
        "recurso": nombre_recurso,
    }, None


ERROR_RECURSO_Y_HTML = (
    "ERROR_RECURSO_Y_HTML: la actividad tiene un archivo en Recurso y PRIZMA no "
    "permite archivo y Html a la vez. No se guardo nada; quita el archivo en "
    "PRIZMA y vuelve a cargar esta actividad."
)

ERROR_QUITANDO_RECURSO = (
    "ERROR_QUITANDO_RECURSO: la actividad tiene un archivo en Recurso y no se "
    "pudo quitar automaticamente. No se guardo nada; quitalo en PRIZMA y vuelve "
    "a cargar esta actividad."
)


def _esperar_sin_conflicto_html_h5p(pagina, timeout_ms=3000, intervalo_ms=200):
    """Espera a que PRIZMA deje de mostrar el aviso de conflicto Html/H5P."""
    transcurrido = 0
    while True:
        if not _conflicto_html_h5p_visible(pagina):
            return True
        if transcurrido >= timeout_ms:
            return False
        pagina.wait_for_timeout(intervalo_ms)
        transcurrido += intervalo_ms


def _resolver_recurso_para_html(pagina, carpeta_temp, actividad, nombre_recurso):
    """Modo HTML: deja el campo "Recurso" sin archivo para poder guardar el Html.

    PRIZMA no permite un archivo en "Recurso" y Html a la vez. Se actua si hay
    un archivo en "Recurso" o si PRIZMA muestra el aviso de conflicto. Se
    quita cualquier archivo, sea H5P o no (decision del usuario, 2026-10-07),
    si PRIZMA_VACIAR_RECURSO_CON_HTML lo permite. Nunca toca el campo Html.

    Returns:
        (resultado_error, archivo_quitado):
          - (None, ""): no habia archivo ni aviso; no se toco nada.
          - (None, nombre): se quito ese archivo (queda en el reporte).
          - (dict, ""): no se puede guardar; la edicion ya quedo cancelada.
    """
    aviso = _conflicto_html_h5p_visible(pagina)
    archivo = archivo_en_recurso(pagina)

    if not aviso and not archivo:
        return None, ""

    if not PRIZMA_VACIAR_RECURSO_CON_HTML:
        print("La actividad tiene un archivo en Recurso y la politica no permite quitarlo: no se guarda.")
        cancelar_edicion_segura(pagina)
        return {
            "ok": False,
            "error": ERROR_RECURSO_Y_HTML,
            "recurso": nombre_recurso,
        }, ""

    if archivo:
        print("La actividad tiene un archivo en Recurso (" + archivo + "); se quita para poder guardar el HTML.")
        quitado, motivo = quitar_archivo_de_recurso(pagina)
    else:
        quitado, motivo = False, "PRIZMA avisa del conflicto Html/H5P pero no se ve ningun archivo en Recurso"

    if quitado and (not aviso or _esperar_sin_conflicto_html_h5p(pagina)):
        print("Archivo quitado de Recurso:", archivo)
        return None, archivo

    error = ERROR_QUITANDO_RECURSO + " (" + (motivo or "PRIZMA sigue mostrando el aviso Html/H5P") + ")"
    fila = str((actividad or {}).get("fila_excel") or "x")
    ruta_estructura = guardar_estructura(
        pagina,
        XPATH_SECCION_RECURSO,
        carpeta_temp,
        "campo_recurso_fila_" + fila + ".html",
    )
    if ruta_estructura:
        error += "; estructura: " + ruta_estructura
    ruta_captura = _guardar_captura_diagnostico(pagina, carpeta_temp, actividad)
    if ruta_captura:
        error += "; captura: " + ruta_captura
    print(error)

    cancelar_edicion_segura(pagina)
    return {
        "ok": False,
        "error": error,
        "recurso": nombre_recurso,
    }, ""


def _cargar_html_en_formulario(pagina, contenido, carpeta_temp, actividad, nombre_recurso):
    """Modo HTML: pega el codigo completo del archivo en el campo "Html".

    contenido es el codigo ya leido con leer_archivo_html. Verifica que el
    campo quedo identico al archivo antes de guardar. Ante cualquier
    problema cancela la edicion sin guardar.

    Returns:
        (resultado_error, html_pegado): resultado_error es None si se puede
        guardar; html_pegado = {"caracteres": n, "reemplazo_previo": bool}.
    """
    editor = _localizar_editor_html(pagina)
    texto_previo = ""
    if editor is None:
        escrito, motivo = False, "no se encontro el campo Html del formulario"
    else:
        texto_previo = _texto_editor_html(editor)
        escrito, motivo = escribir_html_en_editor(editor, contenido)

    if not escrito:
        error = "ERROR_ESCRIBIENDO_HTML: " + motivo + ". No se guardo nada"
        ruta_estructura = _guardar_estructura_campo_html(pagina, carpeta_temp, actividad)
        if ruta_estructura:
            error += "; estructura: " + ruta_estructura
        ruta_captura = _guardar_captura_diagnostico(pagina, carpeta_temp, actividad)
        if ruta_captura:
            error += "; captura: " + ruta_captura
        print(error)
        cancelar_edicion_segura(pagina)
        return {
            "ok": False,
            "error": error,
            "recurso": nombre_recurso,
        }, None

    caracteres = len(normalizar_saltos(contenido))
    print("HTML pegado y verificado:", caracteres, "caracteres.")
    return None, {
        "caracteres": caracteres,
        "reemplazo_previo": bool(texto_previo.strip()),
    }


def _intentar_guardado(pagina, boton_guardar, respuestas, captura):
    """Pulsa Guardar y espera el PATCH de la actividad.

    Devuelve (respuesta_patch, detalle_clic, error_clic):
      - respuesta_patch: dict metodo/status/url, o None si no llego.
      - detalle_clic: motivo si el clic no se pudo completar ("" si se hizo).
      - error_clic: excepcion inesperada al pulsar, o None.
    """
    respuestas.clear()
    captura["fallidas"] = []

    captura[
        "activa"
    ] = True

    print(
        "💾 Guardando..."
    )

    # La confirmacion principal queda enlazada al mismo clic que guarda.
    # Esto evita depender exclusivamente del listener global y elimina una
    # posible condicion de carrera entre el clic y el registro de respuestas.
    respuesta_patch = None
    error_clic_guardar = None
    # Si el TimeoutError viene del clic (boton tapado/deshabilitado) y no de
    # la espera del PATCH, aqui queda el motivo para el diagnostico.
    detalle_clic = ""

    try:
        with pagina.expect_response(
            _es_patch_actividad_response,
            timeout=45000,
        ) as respuesta_info:
            boton_guardar.click(
                timeout=10000
            )

        respuesta_obj = respuesta_info.value
        respuesta_patch = {
            "metodo": respuesta_obj.request.method.upper(),
            "status": respuesta_obj.status,
            "url": respuesta_obj.url,
        }
        print(
            "✅ PATCH capturado directamente:",
            respuesta_patch["status"],
        )

    except PlaywrightTimeoutError as error_espera:
        detalle_clic = _resumir_error_clic(error_espera)
        print(
            "⚠️ No se capturó el PATCH con expect_response; "
            "revisando el listener global..."
        )
        if detalle_clic:
            print("Clic en Guardar no completado:", detalle_clic)

        # El listener global se conserva como respaldo y diagnostico.
        respuesta_patch = esperar_patch(
            pagina,
            respuestas,
            timeout_ms=2000,
        )

    except Exception as error_guardado:
        error_clic_guardar = error_guardado

    captura[
        "activa"
    ] = False

    return respuesta_patch, detalle_clic, error_clic_guardar


def _es_patch_actividad_response(response):
    """Identifica exclusivamente la respuesta PATCH que actualiza una actividad."""
    try:
        return (
            response.request.method.upper() == "PATCH"
            and "/academic/activity/" in response.url
        )
    except Exception:
        return False


def esperar_patch(
    pagina,
    respuestas,
    timeout_ms=15000,
):
    """Fallback historico sobre el listener global de respuestas."""

    esperado = 0

    while esperado < timeout_ms:

        for respuesta in respuestas:

            if respuesta[
                "metodo"
            ] != "PATCH":
                continue

            if "/academic/activity/" not in respuesta[
                "url"
            ]:
                continue

            return respuesta

        pagina.wait_for_timeout(
            250
        )

        esperado += 250

    return None


def verificar_guardado_por_reapertura(
    pagina,
    actividad,
    nombre_recurso,
    estado_navegacion=None,
):
    """Confirma el guardado reabriendo la misma actividad una sola vez.

    Se usa solo como respaldo cuando no se capturo el PATCH de PRIZMA.
    Nunca sustituye un PATCH HTTP con error: solo cubre el caso donde la
    respuesta no fue observada pero el archivo si quedo guardado.
    """
    print("⚠️ PATCH no observado. Verificando el guardado al reabrir la actividad...")

    if not asegurar_listado(pagina):
        return False

    if not entrar_categoria(
        pagina,
        actividad["categoria_prizma"],
        estado_navegacion,
    ):
        return False

    fila, error = buscar_actividad_correcta(pagina, actividad)
    if error or fila is None:
        return False

    if not abrir_edicion(pagina, fila):
        return False

    # PRIZMA puede tardar un instante en reconstruir la tarjeta del recurso
    # despues de guardar. Esperamos de forma adaptativa en vez de consultar
    # una sola vez inmediatamente despues de abrir la edicion.
    confirmado = _esperar_archivo_visible(
        pagina,
        nombre_recurso,
        timeout_ms=6000,
        intervalo_ms=200,
    )

    # Esta reapertura es solo de verificacion: no guardamos otra vez.
    cancelar_edicion_segura(pagina)
    asegurar_listado(pagina)

    if confirmado:
        print("✅ Guardado confirmado al reabrir: el recurso ya aparece en PRIZMA.")

    return confirmado


# ============================================================
# LOGIN AUTOMÁTICO
# ============================================================


def validar_credenciales_prizma(
    usuario_prizma,
    contrasena_prizma,
    timeout_ms=30000,
):
    """Valida credenciales sin iniciar un cargue.

    Retorna un dict con ok/tipo/mensaje para distinguir credenciales
    rechazadas de una indisponibilidad técnica de PRIZMA.
    """
    usuario_prizma = str(usuario_prizma or "").strip()
    contrasena_prizma = str(contrasena_prizma or "")

    if not usuario_prizma or not contrasena_prizma:
        return {
            "ok": False,
            "tipo": "credenciales",
            "mensaje": "Debes ingresar usuario y contraseña de PRIZMA.",
        }

    try:
        with sync_playwright() as p:
            navegador = p.chromium.launch(
                headless=True,
                args=list(CHROMIUM_ARGS),
            )
            try:
                pagina = navegador.new_page(
                    viewport=dict(VIEWPORT_NAVEGADOR)
                )
                pagina.goto(
                    URL_PRIZMA,
                    wait_until="domcontentloaded",
                    timeout=timeout_ms,
                )

                usuario = pagina.locator('input[name="identification_number"]')
                clave = pagina.locator('input[name="password"]')
                boton = pagina.get_by_role(
                    "button", name="Iniciar sesión", exact=True
                )

                usuario.wait_for(state="visible", timeout=timeout_ms)
                clave.wait_for(state="visible", timeout=timeout_ms)
                boton.wait_for(state="visible", timeout=timeout_ms)

                usuario.fill(usuario_prizma)
                clave.fill(contrasena_prizma)
                boton.click(timeout=10000)

                bienvenida = pagina.get_by_text(
                    "Bienvenido a Prizma admin", exact=False
                )

                # PRIZMA muestra el rechazo de credenciales como un toast breve.
                # Antes esperábamos hasta 15 s solo por la bienvenida; para cuando
                # revisábamos el body, el toast ya podía haber desaparecido.
                # Ahora observamos en paralelo éxito y rechazo desde el primer
                # instante posterior al clic.
                error_credenciales = pagina.get_by_text(
                    "Credenciales incorrectas", exact=False
                )

                limite = 15000
                transcurrido = 0
                intervalo = 250

                while transcurrido < limite:
                    try:
                        if bienvenida.count() > 0 and bienvenida.first.is_visible():
                            return {
                                "ok": True,
                                "tipo": "ok",
                                "mensaje": "Credenciales PRIZMA verificadas correctamente.",
                            }
                    except Exception:
                        pass

                    try:
                        if (
                            error_credenciales.count() > 0
                            and error_credenciales.first.is_visible()
                        ):
                            return {
                                "ok": False,
                                "tipo": "credenciales",
                                "mensaje": (
                                    "No se pudo iniciar sesión en PRIZMA. "
                                    "Verifica tu usuario y contraseña."
                                ),
                            }
                    except Exception:
                        pass

                    # Fallback por si PRIZMA cambia el contenedor del toast pero
                    # conserva el texto o una variante equivalente.
                    try:
                        texto = normalizar_texto(
                            pagina.locator("body").inner_text(timeout=1500)
                        )
                        senales_credenciales = [
                            "credenciales incorrectas",
                            "credenciales invalidas",
                            "usuario o contrasena",
                            "usuario y contrasena",
                            "contrasena incorrecta",
                            "usuario incorrecto",
                            "datos incorrectos",
                            "no autorizado",
                            "unauthorized",
                            "invalid credentials",
                        ]
                        if any(s in texto for s in senales_credenciales):
                            return {
                                "ok": False,
                                "tipo": "credenciales",
                                "mensaje": (
                                    "No se pudo iniciar sesión en PRIZMA. "
                                    "Verifica tu usuario y contraseña."
                                ),
                            }
                    except Exception:
                        pass

                    pagina.wait_for_timeout(intervalo)
                    transcurrido += intervalo

                # Si no vimos ni bienvenida ni el toast de rechazo, no culpamos
                # al usuario: puede ser lentitud, mantenimiento o fallo del backend.
                return {
                    "ok": False,
                    "tipo": "servicio",
                    "mensaje": (
                        "PRIZMA no respondió correctamente durante la validación "
                        "del acceso. Intenta nuevamente en unos minutos."
                    ),
                }
            finally:
                try:
                    navegador.close()
                except Exception:
                    pass

    except Exception as e:
        return {
            "ok": False,
            "tipo": "servicio",
            "mensaje": (
                "No fue posible conectar con PRIZMA para validar el acceso. "
                "Intenta nuevamente en unos minutos."
            ),
            "detalle": str(e),
        }


def iniciar_sesion_prizma(
    pagina,
    usuario_prizma,
    contrasena_prizma,
):

    print()
    print("--------------------------------------")
    print("LOGIN AUTOMÁTICO PRIZMA")
    print("--------------------------------------")

    pagina.goto(
        URL_PRIZMA,
        wait_until="domcontentloaded",
        timeout=60000,
    )

    campo_usuario = pagina.locator(
        'input[name="identification_number"]'
    )

    campo_contrasena = pagina.locator(
        'input[name="password"]'
    )

    boton_login = pagina.get_by_role(
        "button",
        name="Iniciar sesión",
        exact=True,
    )

    campo_usuario.wait_for(
        state="visible",
        timeout=30000,
    )

    campo_contrasena.wait_for(
        state="visible",
        timeout=30000,
    )

    boton_login.wait_for(
        state="visible",
        timeout=30000,
    )

    print(
        "✅ Formulario de login detectado."
    )

    campo_usuario.fill(
        usuario_prizma
    )

    campo_contrasena.fill(
        contrasena_prizma
    )

    print(
        "✅ Credenciales introducidas."
    )

    print(
        "Iniciando sesión..."
    )

    boton_login.click(
        timeout=10000
    )

    # --------------------------------------------------------
    # CONFIRMAR PANEL
    # --------------------------------------------------------

    bienvenida = pagina.get_by_text(
        "Bienvenido a Prizma admin",
        exact=False,
    )

    try:

        bienvenida.wait_for(
            state="visible",
            timeout=60000,
        )

    except Exception:

        try:

            texto_visible = pagina.locator(
                "body"
            ).inner_text()

            print()
            print(
                "No se pudo confirmar el panel."
            )

            print(
                texto_visible[:1200]
            )

        except Exception:
            pass

        raise RuntimeError(
            "ERROR_LOGIN_PRIZMA"
        )

    print(
        "✅ Login automático confirmado."
    )

    print(
        "URL después del login:",
        pagina.url,
    )

    # --------------------------------------------------------
    # ENTRAR A ACTIVIDADES
    # --------------------------------------------------------

    tarjeta_actividades = pagina.get_by_role(
        "button"
    ).filter(
        has_text="Actividades"
    ).first

    tarjeta_actividades.wait_for(
        state="visible",
        timeout=30000,
    )

    print(
        "✅ Tarjeta Actividades detectada."
    )

    tarjeta_actividades.click(
        timeout=10000
    )

    # --------------------------------------------------------
    # CONFIRMAR LISTADO
    # --------------------------------------------------------

    buscador = pagina.locator(
        'input[placeholder="Buscar..."]'
    )

    buscador.wait_for(
        state="visible",
        timeout=60000,
    )

    print(
        "✅ Módulo Actividades abierto."
    )

    return True


# ============================================================
# PROCESAR ACTIVIDAD
# ============================================================

def _cargar_archivo_en_recurso(
    pagina,
    campo_recurso,
    ruta_recurso,
    nombre_recurso,
    carpeta_temp,
    actividad,
):
    """Pasos 8-9b del modo archivo (H5P/PDF) de procesar_actividad.

    Sube el archivo al campo "Recurso", espera a que PRIZMA lo muestre y
    aplica la politica del campo Html (PRIZMA_VACIAR_HTML_CON_H5P).

    Returns:
        (resultado_error, html_vaciado): resultado_error es None si se puede
        guardar; html_vaciado es el texto previo del Html si se vacio.
    """

    # 8. CARGAR

    print()
    print(
        "Cargando:",
        nombre_recurso,
    )

    try:

        campo_recurso.set_input_files(
            ruta_recurso
        )

    except Exception:

        cancelar_edicion_segura(
            pagina
        )

        return {
            "ok": False,
            "error":
                "ERROR_SET_INPUT_FILES",
            "recurso":
                nombre_recurso,
        }, None

    # 9. VISIBLE - espera adaptativa; si PRIZMA responde rapido no pagamos
    # los 2.2 segundos fijos de la version anterior.
    if not _esperar_archivo_visible(
        pagina,
        nombre_recurso,
        timeout_ms=2200,
    ):

        cancelar_edicion_segura(
            pagina
        )

        return {
            "ok": False,
            "error":
                "ERROR_RECURSO_NO_VISIBLE",
            "recurso":
                nombre_recurso,
        }, None

    print(
        "✅ Recurso visible."
    )

    # 9b. CAMPO HTML / H5P: PRIZMA no permite guardar un H5P si el campo
    # Html tiene contenido. Politica temporal (PRIZMA_VACIAR_HTML_CON_H5P):
    # se vacia el Html antes de guardar y el contenido previo va al reporte.
    html_vaciado = None

    if (
        str(actividad.get("tipo_archivo") or "").upper() == "H5P"
        or _conflicto_html_h5p_visible(pagina)
    ):
        error_html, html_vaciado = _resolver_campo_html(
            pagina,
            carpeta_temp,
            actividad,
            nombre_recurso,
        )
        if error_html:
            return error_html, None

    return None, html_vaciado


def procesar_actividad(
    pagina,
    actividad,
    indice_recursos,
    carpeta_temp,
    respuestas,
    captura,
    actividades_curso,
    estado_navegacion=None,
    recurso_manual=None,
):

    print()
    print()
    print("======================================")

    print(
        "ACTIVIDAD:",
        actividad["nombre"],
    )

    print(
        "FILA:",
        actividad["fila_excel"],
    )

    print(
        "CATEGORÍA:",
        actividad["categoria_prizma"],
    )

    print(
        "TIPO:",
        actividad["tipo_archivo"],
    )

    print("======================================")

    # 1. VALIDAR CARPETA DE RECURSO EN GOOGLE DRIVE (SI EXISTE)
    # 2. RECURSO DEL ZIP
    #
    # En el reintento individual el archivo lo entrega el usuario
    # (recurso_manual): se usa tal cual, sin buscar en Drive ni en el ZIP.

    if recurso_manual:

        ruta_recurso = str(recurso_manual)
        nombre_recurso = os.path.basename(ruta_recurso)

        if not os.path.isfile(ruta_recurso):

            return {
                "ok": False,
                "error":
                    "ERROR_ARCHIVO_RECURSO_NO_VALIDO",
                "recurso":
                    nombre_recurso,
            }

    else:

        error_carpeta = validar_carpeta_drive_para_actividad(
            actividad
        )

        if error_carpeta:
            return {
                "ok": False,
                "error": error_carpeta,
                "recurso": "",
            }

        recurso, error = resolver_recurso(
            actividad,
            indice_recursos,
            carpeta_temp,
            actividades_curso,
        )

        if error:

            return {
                "ok": False,
                "error": error,
                "recurso": "",
            }

        ruta_recurso = recurso[
            "ruta"
        ]

        nombre_recurso = recurso[
            "nombre_original"
        ]

    # Modo HTML: el codigo del archivo va al campo "Html" (no al campo "Recurso").
    es_html = str(actividad.get("tipo_archivo") or "").upper() == "HTML"

    # 2. LISTADO

    if not asegurar_listado(
        pagina
    ):

        return {
            "ok": False,
            "error":
                "ERROR_NO_SE_PUDO_VOLVER_AL_LISTADO",
            "recurso":
                nombre_recurso,
        }

    # 3. CATEGORÍA

    if not entrar_categoria(
        pagina,
        actividad["categoria_prizma"],
        estado_navegacion,
    ):

        return {
            "ok": False,
            "error":
                "ERROR_ENTRANDO_CATEGORIA",
            "recurso":
                nombre_recurso,
        }

    # 4. BUSCAR

    fila, error = buscar_actividad_correcta(
        pagina,
        actividad,
    )

    if error:

        return {
            "ok": False,
            "error": error,
            "recurso":
                nombre_recurso,
        }

    # 5. EDITAR

    if not abrir_edicion(
        pagina,
        fila,
    ):

        return {
            "ok": False,
            "error":
                "ERROR_ABRIENDO_EDICION",
            "recurso":
                nombre_recurso,
        }

    # En un cargue automático repetido, no volver a subir el mismo recurso
    # cuando PRIZMA ya muestra ese archivo en la actividad. Un reintento manual
    # con archivo explícito sí debe respetar la intención de reemplazarlo.
    recurso_ya_montado = (
        not recurso_manual and archivo_visible(pagina, nombre_recurso)
    )
    if not recurso_manual and es_html:
        contenido_esperado, error_lectura = leer_archivo_html(ruta_recurso)
        editor_html = _localizar_editor_html(pagina)
        if not error_lectura and editor_html is not None:
            recurso_ya_montado = (
                normalizar_saltos(_texto_editor_html(editor_html))
                == normalizar_saltos(contenido_esperado)
            )

    if recurso_ya_montado:
        print("Recurso ya montado; se omite:", nombre_recurso)
        cancelar_edicion_segura(pagina)
        return {
            "ok": True,
            "omitido": True,
            "recurso": nombre_recurso,
        }

    # 6. PREPARAR CAMPOS SEGÚN CATEGORÍA

    if actividad[
        "categoria_prizma"
    ] == "CHALLENGE":

        configuracion_reto, error_reto = (
            obtener_configuracion_reto(
                actividad
            )
        )

        if error_reto:

            cancelar_edicion_segura(
                pagina
            )

            return {
                "ok": False,
                "error": error_reto,
                "recurso": nombre_recurso,
            }

        if not seleccionar_autocomplete_reto(
            pagina,
            "Corte",
            configuracion_reto["corte"],
        ):

            cancelar_edicion_segura(
                pagina
            )

            return {
                "ok": False,
                "error":
                    "ERROR_SELECCIONANDO_CORTE",
                "recurso": nombre_recurso,
            }

        if not seleccionar_autocomplete_reto(
            pagina,
            "Nivel de dificultad",
            configuracion_reto["nivel"],
        ):

            cancelar_edicion_segura(
                pagina
            )

            return {
                "ok": False,
                "error":
                    "ERROR_SELECCIONANDO_NIVEL_DIFICULTAD",
                "recurso": nombre_recurso,
            }

        campo_recurso = detectar_campo_contenido_reto(
            pagina
        )

        if campo_recurso is None:

            cancelar_edicion_segura(
                pagina
            )

            return {
                "ok": False,
                "error":
                    "ERROR_CAMPO_CONTENIDO_RETO_NO_ENCONTRADO",
                "recurso": nombre_recurso,
            }

        print()
        print(
            "✅ Corte:",
            configuracion_reto["corte"],
        )
        print(
            "✅ Nivel:",
            configuracion_reto["nivel"],
        )
        print(
            "✅ El PDF será cargado en CONTENIDO."
        )

    elif es_html:

        # El HTML se pega en el campo "Html"; no se usa el campo de archivo.
        campo_recurso = None

        print()
        print("El codigo HTML se pegara en el campo Html.")

    else:

        campo_recurso = detectar_campo_recurso(
            pagina
        )

        if campo_recurso is None:

            cancelar_edicion_segura(
                pagina
            )

            return {
                "ok": False,
                "error":
                    "ERROR_CAMPO_RECURSO_NO_ENCONTRADO",
                "recurso":
                    nombre_recurso,
            }

        print()
        print(
            "✅ El archivo será cargado en RECURSO."
        )

        print(
            "Material descargable: IGNORADO"
        )

    # 7. DESCRIPCIÓN

    descripcion_ok, error_descripcion = limpiar_descripcion(
        pagina
    )

    if not descripcion_ok:

        cancelar_edicion_segura(
            pagina
        )

        return {
            "ok": False,
            "error":
                error_descripcion,
            "recurso":
                nombre_recurso,
        }

    # 8-9. CARGAR EL RECURSO
    html_vaciado = None
    html_pegado = None
    recurso_quitado = ""

    if es_html:
        # Modo HTML: leer el .html ANTES de tocar el formulario (si no sirve,
        # no se quita nada de "Recurso"), quitar el archivo previo de
        # "Recurso" (si lo hay) y pegar el codigo en el campo Html,
        # verificando que quedo identico.
        contenido_html, error_lectura = leer_archivo_html(ruta_recurso)
        if error_lectura:
            cancelar_edicion_segura(pagina)
            return {
                "ok": False,
                "error": error_lectura,
                "recurso": nombre_recurso,
            }

        error_recurso, recurso_quitado = _resolver_recurso_para_html(
            pagina,
            carpeta_temp,
            actividad,
            nombre_recurso,
        )
        if error_recurso:
            return error_recurso

        error_carga, html_pegado = _cargar_html_en_formulario(
            pagina,
            contenido_html,
            carpeta_temp,
            actividad,
            nombre_recurso,
        )
    else:
        error_carga, html_vaciado = _cargar_archivo_en_recurso(
            pagina,
            campo_recurso,
            ruta_recurso,
            nombre_recurso,
            carpeta_temp,
            actividad,
        )

    if error_carga:
        return error_carga

    # 10. GUARDAR

    boton_guardar = encontrar_boton_guardar(
        pagina
    )

    if boton_guardar is None:

        cancelar_edicion_segura(
            pagina
        )

        return {
            "ok": False,
            "error":
                "ERROR_BOTON_GUARDAR",
            "recurso":
                nombre_recurso,
        }

    respuesta_patch, detalle_clic, error_clic_guardar = _intentar_guardado(
        pagina,
        boton_guardar,
        respuestas,
        captura,
    )

    # 10b. PRIZMA puede mostrar el aviso Html/H5P solo al intentar guardar.
    # Se resuelve el campo Html (o se cancela segun la politica) y se
    # reintenta el guardado UNA vez.
    if (
        error_clic_guardar is None
        and respuesta_patch is None
        and _conflicto_html_h5p_visible(pagina)
    ):
        if es_html:
            # Modo HTML: el aviso indica que sigue habiendo un archivo en
            # "Recurso". Nunca se borra el Html que se acaba de pegar.
            error_conflicto, quitado = _resolver_recurso_para_html(
                pagina,
                carpeta_temp,
                actividad,
                nombre_recurso,
            )
            if error_conflicto:
                return error_conflicto

            recurso_quitado = recurso_quitado or quitado
        else:
            error_html, texto_vaciado = _resolver_campo_html(
                pagina,
                carpeta_temp,
                actividad,
                nombre_recurso,
            )
            if error_html:
                return error_html

            if texto_vaciado is not None:
                html_vaciado = texto_vaciado

        boton_guardar = encontrar_boton_guardar(
            pagina
        )

        if boton_guardar is None:
            cancelar_edicion_segura(
                pagina
            )
            return {
                "ok": False,
                "error": "ERROR_BOTON_GUARDAR",
                "recurso": nombre_recurso,
            }

        print("Reintentando el guardado tras resolver el conflicto Html/H5P...")

        respuesta_patch, detalle_clic, error_clic_guardar = _intentar_guardado(
            pagina,
            boton_guardar,
            respuestas,
            captura,
        )

        if (
            error_clic_guardar is None
            and respuesta_patch is None
            and _conflicto_html_h5p_visible(pagina)
        ):
            cancelar_edicion_segura(
                pagina
            )
            return {
                "ok": False,
                "error": ERROR_QUITANDO_RECURSO if es_html else ERROR_LIMPIANDO_HTML,
                "recurso": nombre_recurso,
            }

    if error_clic_guardar is not None:
        cancelar_edicion_segura(
            pagina
        )

        print(
            "❌ Error al ejecutar el guardado:",
            type(error_clic_guardar).__name__,
            str(error_clic_guardar),
        )

        return {
            "ok": False,
            "error":
                "ERROR_CLIC_GUARDAR: "
                + type(error_clic_guardar).__name__
                + ": "
                + (str(error_clic_guardar).strip().splitlines() or [""])[0][:200],
            "recurso":
                nombre_recurso,
        }

    # 11. PATCH

    if respuesta_patch is None:

        # Se toma ANTES de reabrir la actividad: la reapertura navega y
        # haria desaparecer las alertas y errores del formulario.
        diagnostico = _diagnostico_guardado_sin_patch(
            pagina,
            respuestas,
            captura,
            detalle_clic=detalle_clic,
        )
        ruta_captura = _guardar_captura_diagnostico(
            pagina,
            carpeta_temp,
            actividad,
        )
        if ruta_captura:
            diagnostico += "; captura: " + ruta_captura
        print("Diagnostico de guardado sin PATCH:", diagnostico)

        if verificar_guardado_por_reapertura(
            pagina,
            actividad,
            nombre_recurso,
            estado_navegacion,
        ):
            return {
                "ok": True,
                "error": "",
                "recurso": nombre_recurso,
                "html_vaciado": html_vaciado,
                "html_pegado": html_pegado,
                "recurso_quitado": recurso_quitado,
                "confirmacion": "REAPERTURA",
            }

        return {
            "ok": False,
            "error":
                "ERROR_PATCH_NO_CONFIRMADO: " + diagnostico,
            "recurso":
                nombre_recurso,
        }

    if not (
        200
        <= respuesta_patch[
            "status"
        ]
        < 300
    ):

        asegurar_listado(
            pagina
        )

        return {
            "ok": False,
            "error":
                "ERROR_PATCH_HTTP_"
                + str(
                    respuesta_patch["status"]
                ),
            "recurso":
                nombre_recurso,
        }

    print(
        "✅ PATCH confirmado:",
        respuesta_patch["status"],
    )

    # ========================================================
    # NO REABRIR.
    # DIRECTO A SIGUIENTE ACTIVIDAD.
    # ========================================================

    if not asegurar_listado(
        pagina
    ):

        return {
            "ok": False,
            "error":
                "ERROR_POST_GUARDADO_LISTADO",
            "recurso":
                nombre_recurso,
        }

    return {
        "ok": True,
        "error": "",
        "recurso":
            nombre_recurso,
        "html_vaciado": html_vaciado,
        "html_pegado": html_pegado,
        "recurso_quitado": recurso_quitado,
        "confirmacion": "PATCH",
    }


def _observacion_guardado(resultado):
    if str((resultado or {}).get("confirmacion") or "").upper() == "REAPERTURA":
        observacion = "Carga guardada - verificada al reabrir la actividad"
    else:
        observacion = "Carga guardada - PATCH confirmado"

    # Politica temporal Html/H5P: deja constancia del contenido eliminado.
    html_vaciado = (resultado or {}).get("html_vaciado")
    if html_vaciado is not None:
        previo = re.sub(r"\s+", " ", str(html_vaciado)).strip()
        observacion += (
            " - se vacio el campo Html (contenido previo: \""
            + (previo[:300] if previo else "no legible")
            + "\")"
        )

    # Modo HTML: tamano del codigo pegado y archivo previo quitado de "Recurso".
    html_pegado = (resultado or {}).get("html_pegado")
    if html_pegado:
        observacion += (
            " - HTML pegado (" + str(html_pegado.get("caracteres", 0)) + " caracteres"
            + (", reemplaza el Html previo" if html_pegado.get("reemplazo_previo") else "")
            + ")"
        )

    recurso_quitado = (resultado or {}).get("recurso_quitado")
    if recurso_quitado:
        observacion += " - se quito el archivo previo de Recurso"
        if recurso_quitado != NOMBRE_ARCHIVO_DESCONOCIDO:
            observacion += " (\"" + str(recurso_quitado) + "\")"

    return observacion


def _registrar_respuesta_prizma(response, respuestas, captura):
    """Guarda y muestra las respuestas no-GET mientras se confirma un guardado."""
    if not captura.get("activa"):
        return

    try:
        metodo = response.request.method.upper()
        if metodo == "GET":
            return

        item = {
            "metodo": metodo,
            "status": response.status,
            "url": response.url,
        }
        respuestas.append(item)
        print("PRIZMA HTTP ->", metodo, response.status, str(response.url).split("?", 1)[0])
    except Exception:
        pass


def _registrar_fallo_prizma(request, captura):
    """Guarda las peticiones no-GET que fallan sin respuesta mientras se confirma un guardado.

    Se guardan aparte (captura["fallidas"]) para no mezclarlas con las
    respuestas que revisa esperar_patch.
    """
    if not captura.get("activa"):
        return

    try:
        metodo = request.method.upper()
        if metodo == "GET":
            return

        fallo = request.failure
        motivo = fallo if isinstance(fallo, str) else str(
            getattr(fallo, "error_text", "") or fallo or ""
        )
        captura.setdefault("fallidas", []).append({
            "metodo": metodo,
            "url": request.url,
            "motivo": motivo,
        })
        print("PRIZMA HTTP FALLIDA ->", metodo, str(request.url).split("?", 1)[0], motivo)
    except Exception:
        pass


# Alertas/toasts de PRIZMA (MUI) y errores de validacion de los campos.
_SELECTORES_ALERTAS_PRIZMA = '[role="alert"], .MuiFormHelperText-root.Mui-error, p.form-error'


def _ruta_sin_dominio(url):
    """Ruta de una URL sin dominio ni query string (no expone tokens en reportes)."""
    texto = str(url or "").split("?", 1)[0]
    coincidencia = re.match(r"^[a-zA-Z][a-zA-Z0-9+.-]*://[^/]+(/.*)?$", texto)
    if coincidencia:
        return coincidencia.group(1) or "/"
    return texto


def _textos_alerta_visibles(pagina, limite=3):
    """Textos visibles de alertas y errores de formulario en la pagina actual."""
    textos = []

    try:
        elementos = pagina.locator(_SELECTORES_ALERTAS_PRIZMA)
        cantidad = min(elementos.count(), 10)
    except Exception:
        return textos

    for indice in range(cantidad):
        try:
            elemento = elementos.nth(indice)
            if not elemento.is_visible():
                continue
            texto = re.sub(r"\s+", " ", elemento.inner_text(timeout=500) or "").strip()
        except Exception:
            continue

        texto = texto[:160]
        if texto and not any(texto in previo or previo in texto for previo in textos):
            textos.append(texto)

        if len(textos) >= limite:
            break

    return textos


_SENALES_CLIC_BLOQUEADO = (
    "intercepts pointer events",
    "not enabled",
    "not visible",
    "not stable",
    "detached",
    "outside of the viewport",
)


def _resumir_error_clic(error):
    """Motivo legible si el TimeoutError vino del clic en Guardar.

    Devuelve "" cuando el timeout fue la espera del PATCH (el clic si se hizo).
    Playwright incluye en el mensaje un "call log" con la causa, por ejemplo
    '<div class="MuiBox-root ..."> intercepts pointer events'.
    """
    texto = str(error or "")
    if "click" not in texto.lower():
        return ""

    for linea in texto.splitlines():
        limpia = linea.strip(" -\t")
        if any(senal in limpia for senal in _SENALES_CLIC_BLOQUEADO):
            return limpia[:200]

    primera = texto.strip().splitlines()[0] if texto.strip() else ""
    return primera.strip()[:200]


def _guardar_captura_diagnostico(pagina, carpeta_temp, actividad):
    """Guarda una captura de pantalla del fallo en la carpeta temporal del trabajo.

    Devuelve la ruta relativa al proyecto (o "" si no se pudo). La captura
    queda solo en disco local (temp/ no se versiona).
    """
    try:
        if not carpeta_temp:
            return ""
        os.makedirs(carpeta_temp, exist_ok=True)
        fila = str((actividad or {}).get("fila_excel") or "x")
        ruta = os.path.join(carpeta_temp, "sin_patch_fila_" + fila + ".png")
        pagina.screenshot(path=ruta, full_page=True, timeout=5000)
        try:
            return os.path.relpath(ruta)
        except ValueError:
            return ruta
    except Exception:
        return ""


def _diagnostico_guardado_sin_patch(pagina, respuestas, captura, detalle_clic=""):
    """Explica por que no se confirmo el guardado cuando PRIZMA no envio el PATCH.

    Reune lo observable sin modificar nada: si el clic en Guardar no se pudo
    completar, alertas visibles, respuestas HTTP no-GET y peticiones fallidas
    registradas durante el guardado.
    """
    partes = []

    if detalle_clic:
        partes.append("el clic en Guardar no se completo (" + detalle_clic + ")")

    alertas = _textos_alerta_visibles(pagina)
    if alertas:
        partes.append("PRIZMA mostro: " + " | ".join(alertas))

    http = [
        str(r.get("metodo")) + " " + str(r.get("status")) + " " + _ruta_sin_dominio(r.get("url"))
        for r in list(respuestas or [])[:5]
    ]
    if http:
        partes.append("respuestas: " + ", ".join(http))

    fallidas = [
        str(f.get("metodo")) + " " + _ruta_sin_dominio(f.get("url"))
        + " (" + str(f.get("motivo") or "sin detalle") + ")"
        for f in list((captura or {}).get("fallidas") or [])[:3]
    ]
    if fallidas:
        partes.append("peticiones fallidas: " + ", ".join(fallidas))

    if not partes:
        partes.append(
            "PRIZMA no envio ninguna peticion de guardado ni mostro alertas "
            "(posible validacion del formulario o boton Guardar sin efecto)"
        )

    return "; ".join(partes)


# ============================================================
# SELECCION DE ACTIVIDADES POST-ANALISIS
# ============================================================


def filtrar_actividades_seleccionadas(actividades, seleccion_actividades=None):
    """Devuelve solo las actividades elegidas por hoja + fila.

    ``None`` conserva el comportamiento historico y procesa todas. Una lista
    vacia significa que el usuario no eligio ninguna actividad.
    """
    if seleccion_actividades is None:
        return list(actividades or [])

    claves = set()
    for item in seleccion_actividades or []:
        if not isinstance(item, (list, tuple)) or len(item) != 2:
            continue
        hoja = str(item[0] or "")
        try:
            fila = int(item[1])
        except (TypeError, ValueError):
            continue
        claves.add((hoja, fila))

    return [
        actividad
        for actividad in (actividades or [])
        if (
            str(actividad.get("hoja") or ""),
            int(actividad.get("fila_excel") or 0),
        ) in claves
    ]


# ============================================================
# MOTOR COMPLETO
# ============================================================


def _es_error_estructural_prizma(error):
    texto = str(error or "")
    senales = [
        "Page crashed",
        "Target crashed",
        "Target page, context or browser has been closed",
        "Timeout 30000ms exceeded",
        "Timeout 60000ms exceeded",
        "ERROR_ENTRANDO_CATEGORIA",
        "ERROR_NO_SE_PUDO_VOLVER_AL_LISTADO",
        "ERROR_RECUPERANDO_TRAS_PAGE_CRASH",
        "ERROR_RECUPERANDO_PAGINA",
        "ERROR_RECUPERANDO_FILA",
        "ERROR_LOGIN_PRIZMA",
    ]
    return any(senal in texto for senal in senales)


def _cancelacion_solicitada(estado):
    return bool(estado.get("cancelar_solicitado"))


def _abrir_sesion_inicial(
    abrir_sesion,
    estado,
    esperas=PRIZMA_ESPERAS_REINTENTO_SESION,
    dormir=time.sleep,
):
    """Abre la sesion de PRIZMA al inicio del cargue, reintentando fallos momentaneos.

    Solo reintenta errores estructurales (p. ej. "Page.goto: Timeout 60000ms
    exceeded" cuando PRIZMA tarda en responder) y nunca ERROR_LOGIN_PRIZMA,
    para no repetir un inicio de sesion rechazado.

    Devuelve la pagina con sesion iniciada, o None si el usuario cancelo
    durante una espera. Si se agotan los intentos, relanza el ultimo error.
    """
    total = len(esperas) + 1

    for intento in range(1, total + 1):
        motivo = (
            "inicio del cargue"
            if intento == 1
            else "reintento " + str(intento) + " de " + str(total) + " al iniciar"
        )

        try:
            return abrir_sesion(motivo)
        except Exception as error:
            texto = str(error)
            reintentable = (
                _es_error_estructural_prizma(texto)
                and "ERROR_LOGIN_PRIZMA" not in texto
            )
            if not reintentable or intento == total:
                raise

            espera = int(esperas[intento - 1])
            print(
                "PRIZMA no respondio al iniciar sesion (intento",
                intento,
                "de",
                str(total) + "). Reintentando en",
                espera,
                "s:",
                (texto.strip().splitlines() or [""])[0][:200],
            )
            actualizar_estado(
                estado,
                mensaje=(
                    "PRIZMA no respondió al iniciar sesión "
                    f"(intento {intento} de {total}). "
                    f"Reintentando en {espera} s..."
                ),
            )

            # Espera en pasos de 1 s para atender una cancelacion a tiempo.
            for _ in range(espera):
                if _cancelacion_solicitada(estado):
                    return None
                dormir(1)

    return None


# ============================================================
# REINTENTO INDIVIDUAL: UNA ACTIVIDAD, UN RECURSO ELEGIDO A MANO
# ============================================================

def subir_actividad_manual(
    usuario_prizma,
    contrasena_prizma,
    actividad,
    ruta_archivo,
    carpeta_temp,
    estado,
    indice_recursos=None,
):
    """Sube UN recurso (archivo entregado por el usuario) a UNA actividad.

    Mismo camino que el cargue completo pero para una sola actividad: login,
    listado, categoria, edicion y guardado, con un reintento en sesion limpia
    si el fallo es estructural. El archivo se usa tal cual (recurso_manual),
    sin buscar en Drive ni en el ZIP. Si no se entrega ruta_archivo, resuelve
    el recurso desde indice_recursos (ZIP original).

    Returns:
        dict con {"ok", "error", "recurso"} como procesar_actividad.
    """
    nombre_actividad = str(actividad.get("nombre") or "")
    respuestas = []
    captura = {"activa": False}
    estado_navegacion = {"categoria_actual": None}
    resultado = None

    actualizar_estado(
        estado,
        etapa="login",
        mensaje="Iniciando sesión en PRIZMA...",
        total=1,
        procesadas=0,
        exitosas=0,
        errores=0,
        terminado=False,
    )

    with sync_playwright() as p:
        navegador = None

        def registrar_respuesta(response):
            _registrar_respuesta_prizma(
                response,
                respuestas,
                captura,
            )

        def registrar_fallo(request):
            _registrar_fallo_prizma(
                request,
                captura,
            )

        def abrir_sesion(motivo=""):
            nonlocal navegador
            estado_navegacion["categoria_actual"] = None
            if motivo:
                print("⚠️ Reiniciando sesión PRIZMA:", motivo)
            if navegador is not None:
                try:
                    navegador.close()
                except Exception:
                    pass
            navegador = p.chromium.launch(
                headless=True,
                args=list(CHROMIUM_ARGS),
            )
            pagina = navegador.new_page(
                viewport=dict(VIEWPORT_NAVEGADOR)
            )
            pagina.on("response", registrar_respuesta)
            pagina.on("requestfailed", registrar_fallo)
            iniciar_sesion_prizma(
                pagina,
                usuario_prizma,
                contrasena_prizma,
            )
            return pagina

        try:
            pagina = _abrir_sesion_inicial(abrir_sesion, estado)
        except Exception as error:
            texto = str(error).strip()
            if navegador is not None:
                try:
                    navegador.close()
                except Exception:
                    pass
            if "ERROR_LOGIN_PRIZMA" in texto:
                marca = "ERROR_LOGIN_PRIZMA"
            else:
                marca = "ERROR_CONECTANDO_PRIZMA"
            return {
                "ok": False,
                "error": marca + ": " + (texto.splitlines() or [""])[0][:200],
                "recurso": "",
            }

        if pagina is None:
            # Cancelado durante la espera de la sesion.
            return {
                "ok": False,
                "error": "ERROR_CANCELADO",
                "recurso": "",
            }

        actualizar_estado(
            estado,
            etapa="procesando",
            mensaje="Subiendo el recurso a: " + nombre_actividad,
            procesadas=0,
        )

        try:
            # Preflight: si no estamos en el listado, sesion limpia.
            try:
                listado_ok = asegurar_listado(pagina)
            except Exception:
                listado_ok = False
            if not listado_ok:
                pagina = abrir_sesion(
                    "no fue posible confirmar el listado antes de la actividad"
                )

            for intento in (1, 2):
                if _cancelacion_solicitada(estado):
                    return {
                        "ok": False,
                        "error": "ERROR_CANCELADO",
                        "recurso": "",
                    }

                try:
                    resultado = procesar_actividad(
                        pagina,
                        actividad,
                        indice_recursos or [],
                        carpeta_temp,
                        respuestas,
                        captura,
                        [actividad],
                        estado_navegacion,
                        recurso_manual=ruta_archivo or None,
                    )
                except Exception as error:
                    traceback.print_exc()
                    resultado = {
                        "ok": False,
                        "error": (
                            "ERROR_PROCESANDO_ACTIVIDAD: "
                            + str(error).strip()[:200]
                        ),
                        "recurso": "",
                    }

                if resultado.get("ok"):
                    break

                error_actual = str(resultado.get("error") or "")
                es_ultimo = intento == 2
                if es_ultimo or not _es_error_estructural_prizma(error_actual):
                    break

                # Mismo criterio del cargue: un fallo estructural se
                # reintenta una sola vez con navegador limpio.
                try:
                    pagina = abrir_sesion(
                        "reintento tras " + error_actual
                    )
                except Exception:
                    traceback.print_exc()
                    break
        finally:
            try:
                if navegador is not None:
                    navegador.close()
            except Exception:
                pass

    return resultado


# ============================================================
# VERIFICACION (SOLO LECTURA): ¿LA ACTIVIDAD APARECE EN PRIZMA?
# ============================================================

def verificar_actividades_prizma(
    usuario_prizma,
    contrasena_prizma,
    actividades,
    estado,
):
    """Comprueba, sin tocar formularios, que cada actividad aparezca en PRIZMA.

    ``actividades`` es [(clave, actividad), ...]. Por cada una entra a su
    categoria, busca la fila y anota:

    - "existe": aparece en el listado (si sale mas de una vez se anota),
    - "no_esta": el curso no tiene esa actividad,
    - "no_se_pudo": fallo la consulta (sesion, categoria, error de tabla).

    Returns:
        {"ok", "error", "resultados": [{"clave", "nombre", "verificacion",
        "detalle"}]}
    """
    total = len(actividades)
    resultados = []

    actualizar_estado(
        estado,
        etapa="login",
        mensaje="Iniciando sesión en PRIZMA para verificar el curso...",
        total=total,
        procesadas=0,
        exitosas=0,
        errores=0,
        terminado=False,
    )

    with sync_playwright() as p:
        navegador = None
        estado_navegacion = {"categoria_actual": None}

        def abrir_sesion(motivo=""):
            nonlocal navegador
            estado_navegacion["categoria_actual"] = None
            if motivo:
                print("⚠️ Reiniciando sesión PRIZMA:", motivo)
            if navegador is not None:
                try:
                    navegador.close()
                except Exception:
                    pass
            navegador = p.chromium.launch(
                headless=True,
                args=list(CHROMIUM_ARGS),
            )
            pagina = navegador.new_page(
                viewport=dict(VIEWPORT_NAVEGADOR)
            )
            iniciar_sesion_prizma(
                pagina,
                usuario_prizma,
                contrasena_prizma,
            )
            return pagina

        try:
            pagina = _abrir_sesion_inicial(abrir_sesion, estado)
        except Exception as error:
            texto = str(error).strip()
            if navegador is not None:
                try:
                    navegador.close()
                except Exception:
                    pass
            if "ERROR_LOGIN_PRIZMA" in texto:
                marca = "ERROR_LOGIN_PRIZMA"
            else:
                marca = "ERROR_CONECTANDO_PRIZMA"
            return {
                "ok": False,
                "error": marca + ": " + (texto.splitlines() or [""])[0][:200],
                "resultados": resultados,
            }

        if pagina is None:
            return {
                "ok": False,
                "error": "ERROR_CANCELADO",
                "resultados": resultados,
            }

        try:
            for numero, (clave, actividad) in enumerate(actividades, start=1):
                if _cancelacion_solicitada(estado):
                    return {
                        "ok": False,
                        "error": "ERROR_CANCELADO",
                        "resultados": resultados,
                    }

                nombre = str(actividad.get("nombre") or "")
                actualizar_estado(
                    estado,
                    etapa="verificando",
                    mensaje=(
                        "Verificando " + str(numero) + " de " + str(total)
                        + ": " + nombre
                    ),
                    procesadas=numero - 1,
                )

                entrada = {
                    "clave": clave,
                    "nombre": nombre,
                    "verificacion": "no_se_pudo",
                    "detalle": "",
                }

                try:
                    try:
                        listado_ok = asegurar_listado(pagina)
                    except Exception:
                        listado_ok = False
                    if not listado_ok:
                        pagina = abrir_sesion(
                            "listado no disponible durante la verificacion"
                        )

                    if not entrar_categoria(
                        pagina,
                        actividad["categoria_prizma"],
                        estado_navegacion,
                    ):
                        entrada["detalle"] = (
                            "No se pudo abrir la categoría "
                            + str(actividad["categoria_prizma"]) + "."
                        )
                    else:
                        fila, error = buscar_actividad_correcta(
                            pagina,
                            actividad,
                        )
                        if error:
                            codigo = str(error)
                            if (
                                "ERROR_ACTIVIDAD_NO_ENCONTRADA" in codigo
                                or "ERROR_CURSO_SIN_PLANTILLA" in codigo
                            ):
                                entrada["verificacion"] = "no_esta"
                            elif "ERROR_ACTIVIDAD_DUPLICADA" in codigo:
                                # Varias filas con el nombre: si esta, mas de
                                # una vez; se conserva el detalle del motivo.
                                entrada["verificacion"] = "existe"
                            entrada["detalle"] = mensaje_error(codigo)
                        else:
                            entrada["verificacion"] = "existe"
                            entrada["detalle"] = "Aparece en PRIZMA."
                            if str(actividad.get("tipo_archivo") or "").upper() == "H5P":
                                entrada["detalle"] += " " + _inspeccionar_h5p_montado(
                                    pagina, fila, actividad
                                )
                except Exception:
                    traceback.print_exc()
                    entrada["verificacion"] = "no_se_pudo"
                    entrada["detalle"] = (
                        "No se pudo consultar la actividad (error inesperado)."
                    )

                resultados.append(entrada)

                exitosas = sum(
                    1 for r in resultados
                    if r["verificacion"] == "existe"
                )
                errores = len(resultados) - exitosas
                actualizar_estado(
                    estado,
                    procesadas=numero,
                    exitosas=exitosas,
                    errores=errores,
                )
        finally:
            try:
                if navegador is not None:
                    navegador.close()
            except Exception:
                pass

    return {
        "ok": True,
        "error": "",
        "resultados": resultados,
    }


def _inspeccionar_h5p_montado(pagina, fila, actividad):
    """Lee en modo solo lectura el H5P adjunto y contrasta su título con la matriz."""
    ruta_temporal = ""
    if not abrir_edicion(pagina, fila):
        return "No se pudo abrir la actividad para revisar el H5P adjunto."
    try:
        nombre_archivo = archivo_en_recurso(pagina)
        if not nombre_archivo:
            return "No se detectó un archivo adjunto en el campo Recurso."

        tarjeta = pagina.locator(XPATH_TARJETA_ARCHIVO_RECURSO)
        botones = tarjeta.locator("button")
        if botones.count() < 1:
            return "Se encontró " + nombre_archivo + ", pero no su botón de vista previa/descarga."

        descriptor, ruta_temporal = tempfile.mkstemp(suffix=".h5p", prefix="prizma_h5p_")
        os.close(descriptor)
        try:
            with pagina.expect_download(timeout=12000) as descarga_info:
                botones.first.click(timeout=5000)
            descarga_info.value.save_as(ruta_temporal)
        except Exception:
            return "Se encontró " + nombre_archivo + ", pero PRIZMA no permitió leer su H5P para comparar el título."

        # Se llama a la misma lógica de títulos usada por la pestaña Auditor.
        from auto_prizma.auditor.procesar import clave as clave_auditor
        from auto_prizma.auditor.procesar import leer_titulo_h5p, titulos_visibles_h5p

        titulo, error = leer_titulo_h5p(ruta_temporal)
        if error:
            return "Archivo adjunto: " + nombre_archivo + "; H5P ilegible: " + str(error) + "."
        titulo = str(titulo or "").strip()
        esperado = str(actividad.get("nombre") or "").strip()
        if titulo and clave_auditor(titulo) == clave_auditor(esperado):
            return (
                "Archivo adjunto: " + nombre_archivo
                + '; título H5P coincide con la matriz ("' + titulo + '").'
            )

        visibles = titulos_visibles_h5p(ruta_temporal)
        coincide_contenido = next(
            (t for t in visibles if clave_auditor(t) == clave_auditor(esperado)),
            "",
        )
        if coincide_contenido:
            return (
                "Archivo adjunto: " + nombre_archivo
                + '; el título del paquete es "' + (titulo or "sin título")
                + '", pero el contenido sí muestra el título de la matriz: "'
                + coincide_contenido + '".'
            )
        visible = " · ".join(visibles[:3])
        return (
            "Archivo adjunto: " + nombre_archivo
            + '; REVISAR: título del paquete "' + (titulo or "sin título")
            + '" no coincide con matriz "' + esperado + '"'
            + ("; títulos visibles del contenido: " + visible + "." if visible else ".")
        )
    except Exception as error:
        print("No se pudo inspeccionar el H5P montado:", type(error).__name__)
        return "Aparece en PRIZMA, pero no se pudo previsualizar el título interno del H5P."
    finally:
        try:
            cancelar_edicion_segura(pagina)
        finally:
            if ruta_temporal:
                try:
                    os.remove(ruta_temporal)
                except OSError:
                    pass


def ejecutar_cargue(
    ruta_excel,
    ruta_zip,
    carpeta_temp,
    ruta_reporte,
    procesar_ovi,
    procesar_ova,
    procesar_retos,
    usuario_prizma,
    contrasena_prizma,
    estado,
    seleccion_actividades=None,
    formato_recurso=FORMATO_RECURSO_POR_DEFECTO,
):
    navegador = None

    try:
        print()
        print("======================================")
        print("AUTO PRIZMA PRO")
        print("======================================")
        print("VERSION_SCRIPT =", VERSION_SCRIPT)
        print("BUILD_INTERNO =", BUILD_INTERNO)
        print("URL =", URL_PRIZMA)
        print("MODO = CURSO COMPLETO")
        print("NAVEGADOR = HEADLESS")
        print("LOGIN = AUTOMÁTICO")
        print()
        print("H5P -> RECURSO")
        print("PDF -> RECURSO")
        print("MATERIAL DESCARGABLE -> NO TOCAR")

        actualizar_estado(
            estado,
            etapa="preparando",
            mensaje="Leyendo Excel y preparando recursos...",
            terminado=False,
        )

        actividades = leer_actividades_excel(
            ruta_excel,
            procesar_ovi,
            procesar_ova,
            procesar_retos,
        )

        # Modo HTML: las actividades H5P pasan a HTML (los PDF no cambian).
        actividades = aplicar_formato_recurso(actividades, formato_recurso)

        actividades = filtrar_actividades_seleccionadas(
            actividades,
            seleccion_actividades,
        )

        if not actividades:
            actualizar_estado(
                estado,
                etapa="error",
                mensaje=(
                    "No se encontraron actividades seleccionadas para cargar "
                    "en PRIZMA."
                ),
                terminado=True,
            )
            return

        indice_recursos = crear_indice_recursos(ruta_zip)
        if not indice_recursos:
            actualizar_estado(
                estado,
                etapa="error",
                mensaje=(
                    "No se encontraron archivos "
                    + ("HTML" if str(formato_recurso or "").upper() == "HTML" else "H5P")
                    + "/PDF dentro del ZIP."
                ),
                terminado=True,
            )
            return

        actividades_a_procesar = actividades
        preparar_detalle_actividades(estado, actividades_a_procesar)

        actualizar_estado(
            estado,
            etapa="login",
            mensaje="Iniciando sesión automáticamente en PRIZMA...",
            total=len(actividades_a_procesar),
            procesadas=0,
            exitosas=0,
            errores=0,
        )

        with sync_playwright() as p:
            respuestas = []
            captura = {"activa": False}
            estado_navegacion = {"categoria_actual": None}

            def registrar_respuesta(response):
                _registrar_respuesta_prizma(
                    response,
                    respuestas,
                    captura,
                )

            def registrar_fallo(request):
                _registrar_fallo_prizma(
                    request,
                    captura,
                )

            def abrir_sesion_limpia(motivo=""):
                nonlocal navegador
                estado_navegacion["categoria_actual"] = None
                if motivo:
                    print("⚠️ Reiniciando sesión PRIZMA:", motivo)
                try:
                    if navegador is not None:
                        navegador.close()
                except Exception:
                    pass

                navegador = p.chromium.launch(
                    headless=True,
                    args=list(CHROMIUM_ARGS),
                )
                nueva_pagina = navegador.new_page(
                    viewport=dict(VIEWPORT_NAVEGADOR)
                )
                nueva_pagina.on("response", registrar_respuesta)
                nueva_pagina.on("requestfailed", registrar_fallo)
                iniciar_sesion_prizma(
                    nueva_pagina,
                    usuario_prizma,
                    contrasena_prizma,
                )
                return nueva_pagina

            pagina = _abrir_sesion_inicial(abrir_sesion_limpia, estado)

            if pagina is None:
                actualizar_estado(
                    estado,
                    etapa="cancelado",
                    mensaje="Proceso cancelado por el usuario antes de iniciar sesión en PRIZMA.",
                    terminado=True,
                )
                return

            actualizar_estado(
                estado,
                etapa="procesando",
                mensaje="Login correcto. Módulo Actividades abierto.",
            )

            exitosas = 0
            errores = 0
            procesadas = 0
            errores_estructurales_consecutivos = 0

            for numero, actividad in enumerate(actividades_a_procesar, start=1):
                if _cancelacion_solicitada(estado):
                    actualizar_estado(
                        estado,
                        etapa="cancelado",
                        mensaje=(
                            "Proceso cancelado por el usuario. "
                            f"Se alcanzaron a procesar {procesadas} de "
                            f"{len(actividades_a_procesar)} actividades."
                        ),
                        terminado=True,
                    )
                    return

                print()
                print("######################################")
                print(f"ACTIVIDAD {numero} DE {len(actividades_a_procesar)}")
                print("######################################")

                actualizar_estado(
                    estado,
                    etapa="procesando",
                    mensaje=(
                        f"Procesando {numero} de {len(actividades_a_procesar)}: "
                        + actividad["nombre"]
                    ),
                )
                actualizar_detalle_actividad(estado, numero, "procesando")

                # Si la sesión quedó dañada en la actividad anterior, no se reutiliza.
                if errores_estructurales_consecutivos > 0:
                    try:
                        pagina = abrir_sesion_limpia(
                            "recuperación preventiva tras error estructural"
                        )
                        errores_estructurales_consecutivos = 0
                    except Exception as e:
                        print(traceback.format_exc())
                        resultado = {
                            "ok": False,
                            "error": "ERROR_RECUPERANDO_SESION: " + str(e),
                            "recurso": "",
                        }
                    else:
                        resultado = None
                else:
                    resultado = None

                if resultado is None:
                    # Preflight: si no estamos realmente en el listado, reiniciar antes de tocar la actividad.
                    try:
                        if not asegurar_listado(pagina):
                            pagina = abrir_sesion_limpia(
                                "no fue posible confirmar el listado antes de la actividad"
                            )
                    except Exception:
                        pagina = abrir_sesion_limpia(
                            "fallo comprobando el listado antes de la actividad"
                        )

                    resultado = None
                    ultimo_error = None

                    # Una actividad puede reintentarse una sola vez, siempre con navegador limpio.
                    for intento in (1, 2):
                        if _cancelacion_solicitada(estado):
                            actualizar_estado(
                                estado,
                                etapa="cancelado",
                                mensaje="Proceso cancelado por el usuario.",
                                terminado=True,
                            )
                            return

                        try:
                            resultado = procesar_actividad(
                                pagina,
                                actividad,
                                indice_recursos,
                                carpeta_temp,
                                respuestas,
                                captura,
                                actividades_a_procesar,
                                estado_navegacion,
                            )
                        except Exception as e:
                            print(traceback.format_exc())
                            ultimo_error = "ERROR_NO_CONTROLADO: " + str(e)
                            resultado = {
                                "ok": False,
                                "error": ultimo_error,
                                "recurso": "",
                            }

                        if resultado.get("ok"):
                            break

                        error_actual = resultado.get("error", "")
                        if intento == 1 and _es_error_estructural_prizma(error_actual):
                            actualizar_estado(
                                estado,
                                mensaje=(
                                    "PRIZMA o el navegador perdieron estabilidad. "
                                    "Reiniciando la sesión y reintentando la actividad actual..."
                                ),
                            )
                            try:
                                pagina = abrir_sesion_limpia(error_actual)
                                continue
                            except Exception as e_reinicio:
                                resultado = {
                                    "ok": False,
                                    "error": (
                                        "ERROR_RECUPERANDO_SESION: "
                                        + str(e_reinicio)
                                    ),
                                    "recurso": resultado.get("recurso", ""),
                                }
                        break

                procesadas += 1

                if resultado["ok"]:
                    exitosas += 1
                    errores_estructurales_consecutivos = 0
                    actualizar_detalle_actividad(estado, numero, "ok")
                    guardar_resultado(
                        ruta_reporte,
                        actividad,
                        "OMITIDO" if resultado.get("omitido") else "OK",
                        (
                            "El recurso ya estaba montado en PRIZMA."
                            if resultado.get("omitido")
                            else _observacion_guardado(resultado)
                        ),
                        resultado["recurso"],
                    )
                else:
                    errores += 1
                    error_actual = resultado.get("error", "ERROR_DESCONOCIDO")
                    if _es_error_estructural_prizma(error_actual):
                        errores_estructurales_consecutivos += 1
                    else:
                        errores_estructurales_consecutivos = 0

                    actualizar_detalle_actividad(
                        estado, numero, "error", error_actual
                    )
                    guardar_resultado(
                        ruta_reporte,
                        actividad,
                        "ERROR",
                        error_actual,
                        resultado.get("recurso", ""),
                    )

                actualizar_estado(
                    estado,
                    procesadas=procesadas,
                    exitosas=exitosas,
                    errores=errores,
                    mensaje=(
                        f"Procesadas {procesadas} de {len(actividades_a_procesar)}"
                    ),
                )

                if _cancelacion_solicitada(estado):
                    actualizar_estado(
                        estado,
                        etapa="cancelado",
                        mensaje="Proceso cancelado por el usuario.",
                        terminado=True,
                    )
                    return

            try:
                if navegador is not None:
                    navegador.close()
            except Exception:
                pass

        actualizar_estado(
            estado,
            etapa="finalizado",
            mensaje=(
                "Cargue completo finalizado. "
                f"Exitosas: {exitosas}. Errores: {errores}."
            ),
            terminado=True,
        )

    except Exception as e:
        print()
        print(traceback.format_exc())
        mensaje_error = str(e)
        if "ERROR_LOGIN_PRIZMA" in mensaje_error:
            mensaje_error = (
                "No fue posible iniciar sesión en PRIZMA. "
                "Verifica el usuario y la contraseña."
            )
        elif "Page.goto" in mensaje_error and "Timeout" in mensaje_error:
            mensaje_error = (
                "PRIZMA no respondió al abrir la página de inicio de sesión "
                f"después de {len(PRIZMA_ESPERAS_REINTENTO_SESION) + 1} intentos. "
                "Puede estar lento o caído; intenta nuevamente en unos minutos. "
                "(Detalle: " + mensaje_error.strip().splitlines()[0][:200] + ")"
            )
        actualizar_estado(
            estado,
            etapa="error",
            mensaje=mensaje_error,
            terminado=True,
        )
    finally:
        try:
            if navegador is not None:
                navegador.close()
        except Exception:
            pass
