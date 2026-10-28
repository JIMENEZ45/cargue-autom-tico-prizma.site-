"""Analisis, almacenamiento temporal y estado de cargues Moodle."""

import csv
from difflib import SequenceMatcher
import hashlib
import io
import os
import queue
import re
import shutil
import threading
import time
import unicodedata
import uuid
import zipfile
from datetime import datetime, timezone

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from auto_prizma.core.config import DATA_DIR, UPLOADS_DIR
from auto_prizma.core.storage import read_json, write_json_atomic
from auto_prizma.auditor.procesar import clave as h5p_title_key
from auto_prizma.auditor.procesar import (
    leer_titulo_h5p,
    titulos_principales_h5p,
    titulos_visibles_h5p,
)
from auto_prizma.domain.moodle_resources import (
    BLOCK_LABELS,
    destination_for,
    index_resource_files,
    label_cleanup_plan,
    normalize_key,
    publish_order,
    resolve_resource,
    resource_kind,
)

MOODLE_JOBS = {}
MOODLE_JOBS_LOCK = threading.RLock()
MOODLE_WORKERS = {}
MOODLE_WORKERS_LOCK = threading.Lock()
MAX_ARCHIVE_UNPACKED_BYTES = 4 * 1024 * 1024 * 1024
MAX_MANUAL_RESOURCE_BYTES = 200 * 1024 * 1024
REQUIRED_HEADERS = (
    "Semana correspondiente",
    "Unidad",
    "Nombre del recurso",
    "Categoría",
    "Tipo de recurso",
    "Enlace",
)
REQUIRED_ROW_FIELDS = (
    ("title", "Nombre del recurso"),
    ("category", "Categoría"),
    ("resourceType", "Tipo de recurso"),
    ("linkLabel", "Enlace"),
)
RUNNING_STATUSES = {"queued", "logging_in", "publishing", "paused"}
PAUSE_POLL_SECONDS = 1.0


def _header_key(value):
    return normalize_key(value)


def _read_matrix(filename, content):
    extension = os.path.splitext(str(filename or ""))[1].lower()
    if extension == ".csv":
        texto = content.decode("utf-8-sig")
        try:
            dialect = csv.Sniffer().sniff(texto[:8192], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        filas = list(csv.reader(io.StringIO(texto), dialect))
    elif extension == ".xlsx":
        try:
            libro = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        except (InvalidFileException, zipfile.BadZipFile) as exc:
            raise ValueError("El archivo XLSX no es válido o está dañado.") from exc
        try:
            hoja = libro[libro.sheetnames[0]]
            filas = [list(fila) for fila in hoja.iter_rows(values_only=True)]
        finally:
            libro.close()
    else:
        raise ValueError("La matriz debe ser un archivo CSV o XLSX.")

    required = {_header_key(nombre): nombre for nombre in REQUIRED_HEADERS}
    header_index = None
    header_indexes = {}
    for index, fila in enumerate(filas):
        encontrados = {_header_key(valor): columna for columna, valor in enumerate(fila)}
        if all(key in encontrados for key in required):
            header_index = index
            header_indexes = {nombre: encontrados[key] for key, nombre in required.items()}
            break
    if header_index is None:
        raise ValueError(
            "No encontré los encabezados requeridos: "
            + ", ".join(REQUIRED_HEADERS)
        )

    def value(row, header):
        columna = header_indexes[header]
        return str(row[columna] if columna < len(row) and row[columna] is not None else "").strip()

    rows = []
    issues = []
    for index, record in enumerate(filas[header_index + 1:], start=header_index + 2):
        row = {
            "sourceRow": index,
            "week": value(record, "Semana correspondiente"),
            "unit": value(record, "Unidad"),
            "title": value(record, "Nombre del recurso"),
            "category": value(record, "Categoría"),
            "resourceType": value(record, "Tipo de recurso"),
            "linkLabel": value(record, "Enlace"),
        }
        if not any(row[key] for key in ("title", "category", "resourceType", "linkLabel")):
            continue
        if row["title"] and not any(row[key] for key in ("category", "resourceType", "linkLabel")):
            # Fila de tema de la matriz (p. ej. "1.1 Escritura académica"): no es un recurso ni un error.
            issues.append({"sourceRow": index, "reason": "section-heading", "title": row["title"], "week": row["week"]})
            continue
        missing = [label for key, label in REQUIRED_ROW_FIELDS if not row[key]]
        if missing:
            issues.append({
                "sourceRow": index,
                "reason": "missing-required-field",
                "missing": missing,
                "title": row["title"],
                "week": row["week"],
            })
            continue
        rows.append(row)
    return rows, issues


def _safe_extract_resources(archive_bytes, destination):
    files_root = os.path.join(destination, "resources")
    os.makedirs(files_root, exist_ok=True)
    unpacked_bytes = 0
    extracted = 0
    try:
        with zipfile.ZipFile(io.BytesIO(archive_bytes), "r") as archive:
            for info in archive.infolist():
                if info.is_dir():
                    continue
                name = info.filename.replace("\\", "/")
                if name.startswith("/") or re.match(r"^[A-Za-z]:", name):
                    continue
                parts = [part for part in name.split("/") if part not in {"", "."}]
                if any(part == ".." for part in parts) or not parts:
                    continue
                extension = os.path.splitext(parts[-1])[1].lower()
                if extension not in {".pdf", ".h5p"}:
                    continue
                unpacked_bytes += info.file_size
                if unpacked_bytes > MAX_ARCHIVE_UNPACKED_BYTES:
                    raise ValueError("El ZIP excede el límite de extracción permitido.")
                target = os.path.abspath(os.path.join(files_root, *parts))
                if os.path.commonpath((os.path.abspath(files_root), target)) != os.path.abspath(files_root):
                    continue
                os.makedirs(os.path.dirname(target), exist_ok=True)
                with archive.open(info, "r") as source, open(target, "wb") as output:
                    while True:
                        chunk = source.read(1024 * 1024)
                        if not chunk:
                            break
                        output.write(chunk)
                extracted += 1
    except zipfile.BadZipFile as exc:
        raise ValueError("El archivo de recursos no es un ZIP válido.") from exc
    if extracted == 0:
        raise ValueError("El ZIP no contiene archivos PDF o H5P.")
    return files_root


def validate_course_url(value):
    from urllib.parse import urlparse, parse_qs

    url = str(value or "").strip()
    parsed = urlparse(url)
    if (
        parsed.scheme != "https"
        or parsed.hostname != "prizma.americana.edu.co"
        or parsed.username
        or parsed.password
        or parsed.port not in (None, 443)
        or parsed.path != "/course/view.php"
        or not parse_qs(parsed.query).get("id", [""])[0]
    ):
        raise ValueError("Ingresa un enlace HTTPS válido de un curso de Prizma.")
    return url


def create_moodle_preview(owner, matrix_filename, matrix_bytes, zip_filename, zip_bytes, course_url):
    course_url = validate_course_url(course_url)
    matrix_filename = os.path.basename(str(matrix_filename or ""))
    zip_filename = os.path.basename(str(zip_filename or ""))
    if os.path.splitext(matrix_filename)[1].lower() not in {".csv", ".xlsx"}:
        raise ValueError("La matriz debe ser un archivo CSV o XLSX.")
    if not zip_filename.lower().endswith(".zip"):
        raise ValueError("Los recursos deben venir en un archivo ZIP.")

    rows, issues = _read_matrix(matrix_filename, matrix_bytes)
    job_id = uuid.uuid4().hex
    job_dir = os.path.join(UPLOADS_DIR, "moodle", job_id)
    os.makedirs(job_dir, exist_ok=False)
    try:
        matrix_path = os.path.join(job_dir, matrix_filename)
        zip_path = os.path.join(job_dir, "resources.zip")
        with open(matrix_path, "wb") as output:
            output.write(matrix_bytes)
        with open(zip_path, "wb") as output:
            output.write(zip_bytes)
        resources_root = _safe_extract_resources(zip_bytes, job_dir)
    except Exception:
        shutil.rmtree(job_dir, ignore_errors=True)
        raise
    file_index = index_resource_files(resources_root)
    entries = []
    for row in rows:
        destination = destination_for(row)
        match = resolve_resource(row, file_index)
        entry = {
            "key": (
                f"{row['sourceRow']}:{unicodedata.normalize('NFC', row['title'])}:"
                f"{destination['block'] if destination else 'unsupported'}"
            ),
            **row,
            "kind": resource_kind(row),
            "destination": destination,
            **match,
        }
        if entry["status"] == "matched" and str(entry.get("path") or "").lower().endswith(".h5p"):
            entry["h5p_preview"] = inspect_h5p_title(entry["path"], row["title"])
            if entry["h5p_preview"]["status"] != "match":
                entry["status"] = "title-review"
        entries.append(entry)
    summary = {}
    for entry in entries:
        summary[entry["status"]] = summary.get(entry["status"], 0) + 1

    job = {
        "id": job_id,
        "owner": str(owner or ""),
        "course_url": course_url,
        "matrix_filename": matrix_filename,
        "matrix_identity": hashlib.sha256(matrix_bytes).hexdigest(),
        "zip_filename": zip_filename,
        "job_dir": job_dir,
        "resources_root": resources_root,
        "entries": entries,
        "matrix_issues": issues,
        "label_plan": label_cleanup_plan(
            rows,
            [item for item in issues if item.get("reason") == "missing-required-field"],
        ),
        "summary": summary,
        "created_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        "status": "preview",
        "message": "Vista previa lista. Revisa los recursos antes de publicar.",
        "progress": {"index": 0, "total": 0, "week": "", "title": "", "phase": ""},
        "report": None,
        "cancel_requested": False,
        "pause_requested": False,
        "live_results": [],
        "started_at": "",
        "finished": False,
    }
    with MOODLE_JOBS_LOCK:
        MOODLE_JOBS[job_id] = job
    return job


def inspect_h5p_title(resource_path, expected_title):
    """Requiere un título del contenido igual al de matriz tras normalizar texto."""
    package_title, read_error = leer_titulo_h5p(resource_path)
    package_title = str(package_title or "").strip()
    primary_titles = titulos_principales_h5p(resource_path)
    visible_titles = titulos_visibles_h5p(resource_path)
    expected_key = h5p_title_key(expected_title)
    package_similarity = (
        SequenceMatcher(None, h5p_title_key(package_title), expected_key).ratio()
        if h5p_title_key(package_title) and expected_key else 0.0
    )
    cover_match = next(
        (title for title in primary_titles if h5p_title_key(title) == expected_key),
        "",
    )
    if read_error:
        status = "error"
        detail = "No se pudo leer el título interno del paquete: " + str(read_error) + "."
    elif cover_match:
        status = "match"
        detail = (
            "El título detectado dentro del contenido coincide con el de la matriz "
            "(ignorando mayúsculas, tildes y signos)."
        )
    elif not primary_titles:
        status = "no-cover"
        detail = (
            "No se detectó un título principal en el contenido que permita verificar la portada. "
            "El título del paquete por sí solo no confirma que la portada coincida."
        )
    else:
        status = "cover-mismatch"
        titulos = " · ".join(primary_titles[:3])
        detail = (
            "El título detectado dentro del contenido no coincide exactamente con el de la matriz «"
            + str(expected_title or "") + "». Título(s) encontrado(s): " + titulos + "."
        )
    return {
        "status": status,
        "package_title": package_title,
        "visible_titles": visible_titles,
        "primary_titles": primary_titles,
        "package_similarity": round(package_similarity, 4),
        "cover_title": cover_match,
        "detail": detail,
    }


def attach_manual_resource(job, entry_key, filename, content):
    """Asocia un PDF/H5P adjuntado manualmente a una fila de la vista previa."""
    entry = next((item for item in job.get("entries", []) if item.get("key") == str(entry_key)), None)
    if not entry:
        raise ValueError("No se encontró esa fila de la matriz.")
    destination = entry.get("destination") or {}
    expected_extension = str(destination.get("extension") or "").lower()
    safe_name = os.path.basename(str(filename or "").replace("\\", "/"))
    extension = os.path.splitext(safe_name)[1].lower()
    if expected_extension not in {".pdf", ".h5p"}:
        raise ValueError("Esta fila no admite un archivo PDF o H5P para publicación automática.")
    if extension != expected_extension:
        raise ValueError("Esta fila necesita un archivo " + expected_extension + ".")
    if not content:
        raise ValueError("El archivo seleccionado está vacío.")

    if extension == ".h5p":
        try:
            with zipfile.ZipFile(io.BytesIO(content)) as package:
                if "h5p.json" not in package.namelist():
                    raise ValueError("El H5P no contiene su archivo de configuración h5p.json.")
        except zipfile.BadZipFile as exc:
            raise ValueError("El archivo H5P seleccionado no es válido.") from exc

    manual_dir = os.path.join(job["job_dir"], "manuales")
    os.makedirs(manual_dir, exist_ok=True)
    stored_path = os.path.join(manual_dir, uuid.uuid4().hex + extension)
    with open(stored_path, "wb") as output:
        output.write(content)
    entry["path"] = stored_path
    entry["candidates"] = [stored_path]
    entry["status"] = "matched"
    entry["manual_filename"] = safe_name
    if extension == ".h5p":
        entry["h5p_preview"] = inspect_h5p_title(stored_path, entry.get("title"))
    else:
        entry.pop("h5p_preview", None)
    summary = {}
    for item in job.get("entries", []):
        summary[item["status"]] = summary.get(item["status"], 0) + 1
    job["summary"] = summary
    return entry


def get_moodle_job(job_id, owner):
    with MOODLE_JOBS_LOCK:
        job = MOODLE_JOBS.get(str(job_id))
        if not job or job.get("owner") != str(owner or ""):
            return None
        return job


def local_profile_path(owner):
    owner_hash = hashlib.sha256(str(owner or "").encode("utf-8")).hexdigest()[:24]
    return os.path.join(DATA_DIR, "moodle_profiles", owner_hash)


class _MoodleWorker:
    """Conserva una sesión de navegador y procesa sus trabajos en serie."""

    def __init__(self, owner):
        self.owner = str(owner or "")
        self.tasks = queue.Queue()
        self.client = None
        self.course_url = ""
        self.thread = threading.Thread(
            target=self._run,
            name="moodle-session-" + hashlib.sha256(self.owner.encode("utf-8")).hexdigest()[:8],
            daemon=True,
        )
        self.thread.start()

    def submit(self, job, mode, entries):
        self.tasks.put((job, mode, entries))

    def get_client(self, course_url, profile_path):
        if self.client:
            try:
                page_closed = self.client.page.is_closed()
            except Exception:
                page_closed = True
            if page_closed or self.course_url != course_url:
                self.close_client()
        if not self.client:
            from auto_prizma.automation.moodle.client import MoodleClient

            self.client = MoodleClient.launch(course_url, profile_path)
            self.course_url = course_url
        return self.client

    def close_client(self):
        if self.client:
            try:
                self.client.close()
            except Exception:
                pass
        self.client = None
        self.course_url = ""

    def _run(self):
        while True:
            job, mode, entries = self.tasks.get()
            try:
                _run_moodle_publish(job, mode, entries, self)
            finally:
                self.tasks.task_done()


def _worker_for(owner):
    owner = str(owner or "")
    with MOODLE_WORKERS_LOCK:
        worker = MOODLE_WORKERS.get(owner)
        if not worker or not worker.thread.is_alive():
            worker = _MoodleWorker(owner)
            MOODLE_WORKERS[owner] = worker
        return worker


def write_moodle_report(job, report):
    owner_hash = hashlib.sha256(str(job["owner"]).encode("utf-8")).hexdigest()[:24]
    report_dir = os.path.join(DATA_DIR, "reportes", "moodle", owner_hash)
    os.makedirs(report_dir, exist_ok=True)
    stamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    base = os.path.join(report_dir, "moodle-cargue-" + stamp + "-" + job["id"][:8])
    json_path, csv_path = base + ".json", base + ".csv"
    write_json_atomic(json_path, report)
    headers = ("sourceRow", "week", "title", "status", "file", "detail")
    with io.StringIO(newline="") as buffer:
        writer = csv.writer(buffer)
        writer.writerow(headers)
        for row in report.get("rows", []):
            writer.writerow([row.get(header, "") for header in headers])
        from auto_prizma.core.storage import write_text_atomic
        write_text_atomic(csv_path, buffer.getvalue())
    return {"jsonPath": json_path, "csvPath": csv_path}


def report_from_results(results, stopped_reason=""):
    rows = []
    for result in results:
        entry = result["entry"]
        status = result.get("status", "Revisar en Prizma")
        rows.append({
            "sourceRow": entry["sourceRow"],
            "week": entry.get("week", ""),
            "title": entry["title"],
            "file": entry.get("path", ""),
            "status": status,
            "detail": result.get("detail", ""),
            "kind": entry.get("kind", ""),
            "previewKey": entry.get("previewKey", ""),
        })
    rows.sort(key=lambda row: row.get("sourceRow") or 0)
    summary = {
        "uploaded": sum(row["status"] == "Cargado" for row in rows),
        "alreadyPresent": sum(row["status"] == "Ya estaba cargado" for row in rows),
        "notLoaded": sum(row["status"] == "No cargado" for row in rows),
        "needsReview": sum(row["status"] == "Revisar en Prizma" for row in rows),
        "notProcessed": sum(row["status"] == "No procesado" for row in rows),
    }
    return {
        "createdAt": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        "stoppedReason": stopped_reason,
        "summary": summary,
        "rows": rows,
    }


def publishable_entries(job, selected_keys):
    requested = set(str(value) for value in (selected_keys or []))
    allowed = {entry["key"]: entry for entry in job.get("entries", []) if entry.get("status") == "matched"}
    if not requested:
        raise ValueError("Selecciona al menos un recurso listo para publicar.")
    if not requested.issubset(allowed):
        raise ValueError("La selección contiene recursos que no están disponibles en esta vista previa.")
    return sorted((allowed[key] for key in requested), key=publish_order)


def start_moodle_publish(job, mode, selected_keys, clean_labels=False):
    if mode not in {"test", "batch"}:
        raise ValueError("El tipo de publicación Moodle no es válido.")
    selected = publishable_entries(job, selected_keys)
    if mode == "test":
        selected = selected[:1]
    matrix_path = job.get("matrix_identity") or job["matrix_filename"]
    selected = [
        {
            **entry,
            "previewKey": entry["key"],
            "key": f"{job['course_url']}:{matrix_path}:{entry.get('week', '')}:{entry['key']}",
        }
        for entry in selected
    ]

    with MOODLE_JOBS_LOCK:
        if job.get("status") in RUNNING_STATUSES:
            raise ValueError("Este cargue Moodle ya está en curso.")
        job.update({
            "status": "queued",
            "mode": mode,
            "message": "Abriendo el navegador de Moodle en este equipo…",
            "progress": {"index": 0, "total": len(selected), "week": "", "title": "", "phase": "En cola"},
            "report": None,
            "finished": False,
            "cancel_requested": False,
            "pause_requested": False,
            # Borrar etiquetas sin uso solo en la publicación completa, nunca en la prueba de 1 recurso.
            "clean_labels": bool(clean_labels) and mode == "batch",
            # Solo se limpian las semanas que se están publicando, no todo el curso.
            "label_weeks": sorted({entry.get("week", "") for entry in selected}),
            "label_results": [],
            "live_results": [],
            "queue": [
                {
                    "previewKey": entry["previewKey"],
                    "sourceRow": entry.get("sourceRow"),
                    "week": entry.get("week", ""),
                    "title": entry.get("title", ""),
                    "kind": entry.get("kind", ""),
                }
                for entry in selected
            ],
            "started_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        })

    _worker_for(job["owner"]).submit(job, mode, selected)
    return job


def pause_moodle_publish(job):
    """Pide pausar: el recurso en curso termina y la publicación espera antes del siguiente."""
    with MOODLE_JOBS_LOCK:
        if job.get("status") not in RUNNING_STATUSES or job.get("cancel_requested"):
            raise ValueError("No hay una publicación en curso para pausar.")
        job["pause_requested"] = True
        if job.get("status") == "publishing":
            job["message"] = "Pausando: se termina el recurso actual y luego se espera."
    return job


def resume_moodle_publish(job):
    with MOODLE_JOBS_LOCK:
        if job.get("status") not in RUNNING_STATUSES:
            raise ValueError("No hay una publicación en curso para reanudar.")
        job["pause_requested"] = False
    return job


def cancel_moodle_publish(job):
    """Pide cancelar: el recurso en curso termina y los demás quedan como «No procesado»."""
    with MOODLE_JOBS_LOCK:
        if job.get("status") not in RUNNING_STATUSES:
            raise ValueError("No hay una publicación en curso para cancelar.")
        job["cancel_requested"] = True
        job["pause_requested"] = False
        job["message"] = "Cancelando: se termina el recurso actual y no se publicarán los siguientes."
    return job


def public_moodle_status(job):
    """Estado que consulta la página cada pocos segundos (sin rutas locales de archivos)."""
    with MOODLE_JOBS_LOCK:
        return {
            "status": job.get("status"),
            "message": job.get("message"),
            "progress": dict(job.get("progress") or {}),
            "finished": bool(job.get("finished")),
            "paused": job.get("status") == "paused",
            "pause_requested": bool(job.get("pause_requested")),
            "cancel_requested": bool(job.get("cancel_requested")),
            "started_at": job.get("started_at", ""),
            "results": [dict(item) for item in job.get("live_results") or []],
            "queue": [dict(item) for item in job.get("queue") or []],
            "clean_labels": bool(job.get("clean_labels")),
            "label_results": [dict(item) for item in job.get("label_results") or []],
        }


def _wait_while_paused(job):
    while True:
        with MOODLE_JOBS_LOCK:
            if job.get("cancel_requested") or not job.get("pause_requested"):
                if job.get("status") == "paused":
                    job["status"] = "publishing"
                return
            if job.get("status") != "paused":
                job["status"] = "paused"
                job["message"] = "En pausa. El navegador sigue abierto; pulsa «Reanudar» para continuar."
                job["progress"]["phase"] = "En pausa"
        time.sleep(PAUSE_POLL_SECONDS)


def _record_result(job, results, result):
    results.append(result)
    entry = result["entry"]
    with MOODLE_JOBS_LOCK:
        job.setdefault("live_results", []).append({
            "previewKey": entry.get("previewKey", ""),
            "sourceRow": entry.get("sourceRow"),
            "week": entry.get("week", ""),
            "title": entry.get("title", ""),
            "kind": entry.get("kind", ""),
            "status": result.get("status", ""),
            "detail": result.get("detail", ""),
        })
        job["progress"]["done"] = len(results)


def _set_phase(job, phase):
    with MOODLE_JOBS_LOCK:
        job["progress"]["phase"] = phase


def _clean_unused_labels(job, client):
    """Borra, al final de la publicación, las etiquetas que cada semana no usa según la matriz.

    Va después de publicar porque las etiquetas marcan dónde se inserta cada recurso.
    """
    weeks = set(job.get("label_weeks") or [])
    plan = [item for item in job.get("label_plan") or [] if item.get("delete") and item.get("week") in weeks]
    for item in plan:
        _wait_while_paused(job)
        if job.get("cancel_requested"):
            return
        week = item["week"]
        with MOODLE_JOBS_LOCK:
            job["status"] = "publishing"
            job["message"] = f"Borrando etiquetas sin uso de {week}"
            job["progress"].update({"week": week, "title": "Etiquetas sin uso", "phase": "Borrando etiquetas"})
        try:
            client.open_course(job["course_url"])
            client.open_week(week)
            client.wait_for_week_content()
        except Exception as exc:
            for block in item["delete"]:
                _record_label(job, week, block, "No se pudo borrar", "No se abrió la semana: " + str(exc))
            continue
        for block in item["delete"]:
            try:
                outcome = client.delete_unused_label(block)
                if outcome == "deleted":
                    _record_label(job, week, block, "Borrada", "La semana no tiene recursos de este bloque en la matriz.")
                else:
                    _record_label(job, week, block, "No estaba", "La etiqueta no aparece en la semana.")
            except Exception as exc:
                _record_label(job, week, block, "No se pudo borrar", str(exc))


def _record_label(job, week, block, status, detail):
    with MOODLE_JOBS_LOCK:
        job.setdefault("label_results", []).append({
            "week": week,
            "block": block,
            "label": BLOCK_LABELS.get(block, block),
            "status": status,
            "detail": detail,
        })


def _published_before(state, entry, course_url):
    """True si el historial local ya registra este recurso como publicado en esa semana del curso.

    Cualquier registro cuenta (no solo los "verified-"): antes, con un id numérico se volvía a
    depender de verlo en la página, y si la semana aún no había cargado se duplicaba. También
    reconoce registros de otra versión de la matriz (otro hash o fila) por curso, semana, título y bloque.
    """
    completed = state.get("completed") or {}
    if entry["key"] in completed:
        return True
    title = unicodedata.normalize("NFC", str(entry.get("title") or ""))
    block = (entry.get("destination") or {}).get("block") or "unsupported"
    week_part = ":" + str(entry.get("week") or "") + ":"
    suffix = ":" + title + ":" + block
    prefix = str(course_url) + ":"
    return any(
        key.startswith(prefix) and week_part in key[len(prefix):] and key.endswith(suffix)
        for key in completed
    )


def _friendly_error(job, exc):
    """Mensaje legible para el usuario, con el detalle técnico al final."""
    detail = str(exc)
    if job.get("status") == "logging_in" and "Timeout" in detail:
        return (
            "No se detectó el inicio de sesión en Moodle (se espera hasta 2 minutos). "
            "Vuelve a publicar e inicia sesión en la ventana de Chromium. Detalle: " + detail
        )
    if "Target page, context or browser has been closed" in detail or "has been closed" in detail:
        return "Se cerró la ventana de Chromium antes de terminar. Vuelve a publicar y déjala abierta. Detalle: " + detail
    return detail


def _run_moodle_publish(job, mode, selected, worker):
    results = []
    try:
        profile = local_profile_path(job["owner"])
        state_path = os.path.join(profile, "completed-resources.json")
        state = read_json(state_path, {"completed": {}, "pending": {}}, dict)
        os.makedirs(profile, exist_ok=True)

        with MOODLE_JOBS_LOCK:
            job["status"] = "logging_in"
            job["message"] = "Inicia sesión manualmente en la ventana de Chromium; el navegador permanecerá abierto durante el cargue."
            job["progress"]["phase"] = "Esperando inicio de sesión"

        client = worker.get_client(job["course_url"], profile)
        for index, entry in enumerate(selected, start=1):
            _wait_while_paused(job)
            if job.get("cancel_requested"):
                _record_result(job, results, {"entry": entry, "status": "No procesado", "detail": "El usuario canceló la publicación."})
                continue
            with MOODLE_JOBS_LOCK:
                job["status"] = "publishing"
                job["message"] = f"Revisando {entry['week']}: {entry['title']}"
                job["progress"].update({
                    "index": index,
                    "total": len(selected),
                    "week": entry.get("week", ""),
                    "title": entry["title"],
                    "phase": "Revisando si ya está en el curso",
                })
            try:
                client.open_course(job["course_url"])
                client.open_week(entry.get("week", ""))
                completed = state.get("completed", {}).get(entry["key"], {})
                existing_id = client.find_existing_resource(
                    entry["title"],
                    completed.get("prizmaResourceId", ""),
                    entry.get("path", ""),
                )
                if existing_id:
                    state.setdefault("completed", {})[entry["key"]] = {
                        "prizmaResourceId": existing_id,
                        "completedAt": datetime.now(timezone.utc).isoformat(timespec="seconds"),
                    }
                    state.get("pending", {}).pop(entry["key"], None)
                    _record_result(job, results, {"entry": entry, "status": "Ya estaba cargado", "detail": "Encontrado en la semana del curso; no se duplicó."})
                    write_json_atomic(state_path, state)
                    continue
                if _published_before(state, entry, job["course_url"]):
                    state.get("pending", {}).pop(entry["key"], None)
                    write_json_atomic(state_path, state)
                    _record_result(job, results, {
                        "entry": entry,
                        "status": "Ya estaba cargado",
                        "detail": "El historial local indica que ya se cargó; se evitó duplicarlo.",
                    })
                    continue

                def mark_pending():
                    _set_phase(job, "Guardando en Moodle")
                    state.setdefault("pending", {})[entry["key"]] = {
                        "startedAt": datetime.now(timezone.utc).isoformat(timespec="seconds"),
                    }
                    write_json_atomic(state_path, state)

                _set_phase(job, "Subiendo el archivo")
                resource_id = client.publish(entry, on_before_save=mark_pending)
                state.setdefault("completed", {})[entry["key"]] = {
                    "prizmaResourceId": resource_id,
                    "completedAt": datetime.now(timezone.utc).isoformat(timespec="seconds"),
                }
                state.get("pending", {}).pop(entry["key"], None)
                write_json_atomic(state_path, state)
                _record_result(job, results, {"entry": entry, "status": "Cargado", "detail": "Carga completada y verificada."})
            except Exception as exc:
                detail = str(exc)
                h5p_rejected = entry.get("destination", {}).get("module") == "h5p" and (
                    "paquete h5p rechazado" in detail.lower()
                    or "editor h5p activo" in detail.lower()
                )
                status = "No cargado" if h5p_rejected or getattr(exc, "code", "") in {
                    "DESTINATION_BLOCK_UNAVAILABLE",
                    "H5P_PACKAGE_VALIDATION_FAILED",
                    "H5P_EDITOR_NOT_READY",
                } else "Revisar en Prizma"
                _record_result(job, results, {"entry": entry, "status": status, "detail": detail})
            if mode == "test":
                break

        if mode == "batch" and job.get("clean_labels"):
            _clean_unused_labels(job, client)

        cancelled = bool(job.get("cancel_requested"))
        report = report_from_results(results, "Cancelado por el usuario." if cancelled else "")
        report["labels"] = [dict(item) for item in job.get("label_results") or []]
        paths = write_moodle_report(job, report)
        with MOODLE_JOBS_LOCK:
            job.update({
                "status": "cancelled" if cancelled else "completed",
                "message": "Publicación cancelada por el usuario." if cancelled else "Publicación Moodle terminada.",
                "report": report,
                "report_paths": paths,
                "finished": True,
                "pause_requested": False,
                "progress": {
                    "index": len(selected),
                    "total": len(selected),
                    "done": len(results),
                    "week": "",
                    "title": "",
                    "phase": "Cancelado" if cancelled else "Completado",
                },
            })
    except Exception as exc:
        message = _friendly_error(job, exc)
        report = report_from_results(results, message) if results else None
        paths = write_moodle_report(job, report) if report else {}
        with MOODLE_JOBS_LOCK:
            job.update({
                "status": "error",
                "message": message,
                "report": report,
                "report_paths": paths,
                "finished": True,
                "pause_requested": False,
            })
