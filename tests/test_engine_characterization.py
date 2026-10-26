"""Pruebas de caracterizacion del motor de cargue de PRIZMA.

Fijan el comportamiento ACTUAL de las partes del motor que no tocan PRIZMA
(lectura de la matriz, tipo de archivo, indice del ZIP, resolucion de
recursos, reporte CSV y reglas de Retos) para que la futura division de
auto_prizma/automation/prizma/engine.py no cambie nada sin que se note.

Importante: al importar main, engine_overrides reemplaza en el modulo engine
leer_actividades_excel y resolver_recurso. Aqui se prueban las funciones
ORIGINALES, que engine_overrides conserva como _LEER_ACTIVIDADES_EXCEL_ORIGINAL
y _RESOLVER_RECURSO_ORIGINAL. Si se elimina el monkeypatching, apuntar estas
pruebas a las funciones del motor directamente.
"""

import csv
import io
import zipfile

import pytest
from openpyxl import Workbook

from auto_prizma.automation.prizma import engine
from auto_prizma.automation.prizma import engine_overrides

leer_actividades_excel = engine_overrides._LEER_ACTIVIDADES_EXCEL_ORIGINAL
resolver_recurso = engine_overrides._RESOLVER_RECURSO_ORIGINAL

DRIVE_URL = "https://drive.google.com/drive/folders/ABC123"


# ============================================================
# DATOS DE PRUEBA
# ============================================================

@pytest.fixture
def ruta_matriz(tmp_path):
    """Matriz con una hoja valida y una hoja sin cabecera (que se ignora)."""
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Curso A"
    hoja["A1"] = "Administración de Empresas"
    hoja["A2"] = "Contabilidad I"
    hoja["A4"] = "Semana correspondiente"

    # (semana, unidad, C, actividad, categoria, tipo de recurso, enlace)
    filas = [
        (1, "Unidad 1", None, "Infografía de costos", "OVI", "Infografía", "carpeta"),
        (1, "Unidad 1", None, "Lectura guía", "OVI", "PDF descargable", None),
        (2, "Unidad 2", None, "Video introducción", "OVI", "Video intro", None),
        (2, "Unidad 2", None, "Actividad interactiva", "OVA", "Cualquiera", None),
        (3, "Unidad 1", None, "Reto 1", "Challenge", "Reto evaluativo", None),
        (3, None, None, None, "OVI", "PDF", None),
        (3, None, None, "Foro de discusión", "Foro", "PDF", None),
        (4, None, None, "Recurso raro", "OVI", "Texto", None),
        (4, None, None, "Recurso por enlace", "ovi", None, "archivo.pdf"),
    ]
    for numero_fila, fila in enumerate(filas, start=5):
        for columna, valor in enumerate(fila, start=1):
            if valor is not None:
                hoja.cell(row=numero_fila, column=columna, value=valor)

    hoja.cell(row=5, column=7).hyperlink = DRIVE_URL
    libro.create_sheet("Notas")["A1"] = "sin cabecera"

    ruta = tmp_path / "matriz.xlsx"
    libro.save(ruta)
    return str(ruta)


@pytest.fixture
def ruta_zip(tmp_path):
    ruta = tmp_path / "recursos.zip"
    with zipfile.ZipFile(ruta, "w") as archivo:
        archivo.writestr("carpeta/", b"")
        archivo.writestr("carpeta/Infografia_de_costos.h5p", b"h5p")
        archivo.writestr("Lectura guía.pdf", b"%PDF")
        archivo.writestr("dup/Copia.pdf", b"1")
        archivo.writestr("otra/Copia.pdf", b"2")
        archivo.writestr("notas.txt", b"x")
    return str(ruta)


def _por_nombre(actividades):
    return {actividad["nombre"]: actividad for actividad in actividades}


# ============================================================
# LECTURA DE LA MATRIZ
# ============================================================

def test_leer_actividades_excel_extracts_only_supported_rows(ruta_matriz):
    actividades = leer_actividades_excel(ruta_matriz)

    resumen = [
        (a["fila_excel"], a["nombre"], a["categoria_prizma"], a["tipo_archivo"])
        for a in actividades
    ]
    assert resumen == [
        (5, "Infografía de costos", "OVI", "H5P"),
        (6, "Lectura guía", "OVI", "PDF"),
        (8, "Actividad interactiva", "OVA", "H5P"),
        (9, "Reto 1", "CHALLENGE", "PDF"),
        (13, "Recurso por enlace", "OVI", "PDF"),
    ]


def test_leer_actividades_excel_fills_course_metadata_and_hyperlink(ruta_matriz):
    primera = leer_actividades_excel(ruta_matriz)[0]

    assert primera == {
        "fila_excel": 5,
        "hoja": "Curso A",
        "programa": "Administración de Empresas",
        "curso": "Contabilidad I",
        "semana": "1",
        "unidad": "Unidad 1",
        "nombre": "Infografía de costos",
        "categoria_prizma": "OVI",
        "tipo_recurso": "Infografía",
        "tipo_archivo": "H5P",
        "enlace": "carpeta",
        "url_recurso": DRIVE_URL,
    }


def test_leer_actividades_excel_respects_category_flags(ruta_matriz):
    solo_ova = leer_actividades_excel(
        ruta_matriz,
        procesar_ovi=False,
        procesar_ova=True,
        procesar_retos=False,
    )

    assert [a["nombre"] for a in solo_ova] == ["Actividad interactiva"]


def test_main_replaces_leer_actividades_excel_with_target_sheet_filter():
    # Documenta el acople: tras importar main, el motor usa el override.
    assert engine.leer_actividades_excel is not leer_actividades_excel
    assert engine.resolver_recurso is not resolver_recurso


# ============================================================
# REGLAS DE TIPO Y CATEGORIA
# ============================================================

@pytest.mark.parametrize(
    ("categoria", "tipo_recurso", "enlace", "esperado"),
    [
        ("OVA", "x", "", "H5P"),
        ("OVI", "PDF", "", "PDF"),
        ("OVI", "H5P", "", "H5P"),
        ("OVI", "", ".../a.pdf", "PDF"),
        ("OVI", "Infografía", "", "H5P"),
        ("OVI", "Video cierre", "", None),
        ("Challenge", "Reto evaluativo", "", "PDF"),
        ("Challenge", "Otro", "", None),
        ("Foro", "PDF", "", None),
    ],
)
def test_determinar_tipo_archivo(categoria, tipo_recurso, enlace, esperado):
    assert engine.determinar_tipo_archivo(categoria, tipo_recurso, enlace) == esperado


@pytest.mark.parametrize(
    ("valor", "esperado"),
    [
        ("OVI", "OVI"),
        (" ova ", "OVA"),
        ("Challenge", "CHALLENGE"),
        ("CHALLENGE", "CHALLENGE"),
        ("Reto", None),
        ("", None),
        (None, None),
    ],
)
def test_normalizar_categoria(valor, esperado):
    assert engine.normalizar_categoria(valor) == esperado


@pytest.mark.parametrize(
    ("valor", "esperado"),
    [
        ("  Árbol  ÑANDÚ ", "arbol nandu"),
        (None, ""),
        (12, "12"),
        ("a\tb\n c", "a b c"),
    ],
)
def test_normalizar_texto(valor, esperado):
    assert engine.normalizar_texto(valor) == esperado


# ============================================================
# INDICE DEL ZIP Y RESOLUCION DE RECURSOS
# ============================================================

def test_crear_indice_recursos_keeps_only_h5p_and_pdf(ruta_zip):
    indice = engine.crear_indice_recursos(ruta_zip)

    assert [(r["nombre"], r["extension"], r["miembro"]) for r in indice] == [
        ("Infografia_de_costos.h5p", ".h5p", "carpeta/Infografia_de_costos.h5p"),
        ("Lectura guía.pdf", ".pdf", "Lectura guía.pdf"),
        ("Copia.pdf", ".pdf", "dup/Copia.pdf"),
        ("Copia.pdf", ".pdf", "otra/Copia.pdf"),
    ]
    assert all(r["zip"] == ruta_zip for r in indice)


def _resolver(actividad, ruta_zip, ruta_matriz, tmp_path, nombre_drive=None):
    actividad = dict(actividad)
    if nombre_drive:
        actividad["nombre_recurso_drive"] = nombre_drive
    carpeta = tmp_path / "extraccion"
    carpeta.mkdir(exist_ok=True)
    return resolver_recurso(
        actividad,
        engine.crear_indice_recursos(ruta_zip),
        str(carpeta),
        leer_actividades_excel(ruta_matriz),
    )


@pytest.mark.parametrize(
    ("actividad", "nombre_drive", "archivo", "puntuacion", "contenido"),
    [
        ("Infografía de costos", None, "Infografia_de_costos.h5p", 580, b"h5p"),
        ("Lectura guía", None, "Lectura guía.pdf", 580, b"%PDF"),
        ("Lectura guía", "Lectura guía.pdf", "Lectura guía.pdf", 10000, b"%PDF"),
    ],
)
def test_resolver_recurso_finds_and_extracts_file(
    ruta_matriz, ruta_zip, tmp_path, actividad, nombre_drive, archivo, puntuacion, contenido
):
    actividades = _por_nombre(leer_actividades_excel(ruta_matriz))

    resultado, error = _resolver(
        actividades[actividad], ruta_zip, ruta_matriz, tmp_path, nombre_drive
    )

    assert error is None
    assert resultado["nombre_original"] == archivo
    assert resultado["puntuacion"] == puntuacion
    with open(resultado["ruta"], "rb") as extraido:
        assert extraido.read() == contenido


@pytest.mark.parametrize(
    ("actividad", "cambios", "nombre_drive", "error_esperado"),
    [
        ("Lectura guía", {}, "No existe.pdf", "ERROR_RECURSO_DRIVE_NO_ESTA_EN_ZIP"),
        ("Lectura guía", {}, "Copia.pdf", "ERROR_RECURSO_DUPLICADO"),
        ("Actividad interactiva", {}, None, "ERROR_RECURSO_NO_ENCONTRADO"),
        ("Reto 1", {"tipo_archivo": "DOCX"}, None, "ERROR_TIPO_RECURSO_NO_SOPORTADO"),
    ],
)
def test_resolver_recurso_error_codes(
    ruta_matriz, ruta_zip, tmp_path, actividad, cambios, nombre_drive, error_esperado
):
    actividades = _por_nombre(leer_actividades_excel(ruta_matriz))
    objetivo = dict(actividades[actividad], **cambios)

    resultado, error = _resolver(objetivo, ruta_zip, ruta_matriz, tmp_path, nombre_drive)

    assert resultado is None
    assert error == error_esperado


# ============================================================
# REPORTE CSV
# ============================================================

def test_guardar_resultado_writes_utf8_bom_csv_with_header_once(ruta_matriz, tmp_path):
    actividades = leer_actividades_excel(ruta_matriz)
    ruta = tmp_path / "reporte.csv"

    engine.guardar_resultado(str(ruta), actividades[0], "OK", "obs", "Infografia_de_costos.h5p")
    engine.guardar_resultado(str(ruta), actividades[1], "ERROR", "ERROR_X")

    contenido = ruta.read_bytes()
    assert contenido.startswith(b"\xef\xbb\xbf")
    filas = list(csv.reader(io.StringIO(contenido.decode("utf-8-sig"))))
    assert filas == [
        ["Fila Excel", "Hoja", "Programa", "Curso", "Semana", "Unidad", "Actividad",
         "Categoría", "Tipo", "Recurso", "Resultado", "Observación", "Motivo"],
        ["5", "Curso A", "Administración de Empresas", "Contabilidad I", "1", "Unidad 1",
         "Infografía de costos", "OVI", "H5P", "Infografia_de_costos.h5p", "OK", "obs", ""],
        ["6", "Curso A", "Administración de Empresas", "Contabilidad I", "1", "Unidad 1",
         "Lectura guía", "OVI", "PDF", "", "ERROR", "ERROR_X", ""],
    ]


# ============================================================
# RETOS EVALUATIVOS Y ERRORES ESTRUCTURALES
# ============================================================

@pytest.mark.parametrize(
    ("programa", "esperado"),
    [
        ("Especialización en Gerencia", True),
        ("Maestría en X", True),
        ("Administración de Empresas", False),
    ],
)
def test_es_programa_posgrado(programa, esperado):
    assert engine.es_programa_posgrado(programa) is esperado


@pytest.mark.parametrize(
    ("unidad", "esperado"),
    [("Unidad 1", 1), ("unidad 2", 2), ("U3", None), ("Semana", None), ("", None), (None, None)],
)
def test_obtener_numero_unidad(unidad, esperado):
    assert engine.obtener_numero_unidad(unidad) == esperado


def test_obtener_configuracion_reto():
    assert engine.obtener_configuracion_reto(
        {"unidad": "Unidad 1", "programa": "Administración de Empresas"}
    ) == (
        {"corte": "Corte 1", "nivel": "Nivel básico", "posgrado": False, "numero_unidad": 1},
        None,
    )
    assert engine.obtener_configuracion_reto(
        {"unidad": "Unidad 2", "programa": "Especialización en Gerencia"}
    ) == (
        {"corte": "Corte 1", "nivel": "Nivel intermedio", "posgrado": True, "numero_unidad": 2},
        None,
    )
    assert engine.obtener_configuracion_reto(
        {"unidad": "Sin", "programa": "X"}
    ) == (None, "ERROR_UNIDAD_RETO_NO_RECONOCIDA")


@pytest.mark.parametrize(
    ("error", "esperado"),
    [
        ("Page crashed x", True),
        ("ERROR_LOGIN_PRIZMA", True),
        ("ERROR_RECURSO_NO_ENCONTRADO", False),
        ("", False),
        (None, False),
    ],
)
def test_es_error_estructural_prizma(error, esperado):
    assert engine._es_error_estructural_prizma(error) is esperado
