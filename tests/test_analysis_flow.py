"""Pruebas del flujo POST /analizar con matriz y ZIP locales.

Fijan el comportamiento del analisis previo al cargue: validaciones, archivos
guardados, trabajo registrado en memoria y datos que muestra la pagina.
"""

import io
import os
import shutil
import zipfile

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook

from main import app
from auto_prizma.core.config import RESULTADOS_DIR, TEMP_DIR, UPLOADS_DIR
from auto_prizma.services.auth import SESION_COOKIE, _firmar_sesion
from auto_prizma.services.upload_jobs import TRABAJOS


def _cliente():
    cliente = TestClient(app, client=("127.0.0.1", 50000))
    cliente.cookies.set(SESION_COOKIE, _firmar_sesion("adminlocal"))
    return cliente


def _matriz_xlsx():
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Curso A"
    hoja["A1"] = "Administración de Empresas"
    hoja["A2"] = "Contabilidad I"
    hoja["A4"] = "Semana correspondiente"
    filas = [
        (1, "Unidad 1", None, "Infografía de costos", "OVI", "Infografía", None),
        (1, "Unidad 1", None, "Lectura guía", "OVI", "PDF descargable", None),
        (3, "Unidad 1", None, "Reto 1", "Challenge", "Reto evaluativo", None),
    ]
    for numero_fila, fila in enumerate(filas, start=5):
        for columna, valor in enumerate(fila, start=1):
            if valor is not None:
                hoja.cell(row=numero_fila, column=columna, value=valor)
    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()


def _zip_recursos():
    salida = io.BytesIO()
    with zipfile.ZipFile(salida, "w") as archivo:
        archivo.writestr("Infografia_de_costos.h5p", b"h5p")
        archivo.writestr("Lectura guía.pdf", b"%PDF")
    return salida.getvalue()


def _analizar(datos=None, matriz=None, recursos=None, test_docx=None):
    archivos = {}
    if matriz is not None:
        archivos["matriz_archivo"] = ("matriz.xlsx", io.BytesIO(matriz), "application/octet-stream")
    if recursos is not None:
        archivos["recursos"] = ("recursos.zip", io.BytesIO(recursos), "application/zip")
    if test_docx is not None:
        archivos["test_archivo_docx"] = ("test.docx", io.BytesIO(test_docx), "application/octet-stream")
    previos = set(TRABAJOS)
    respuesta = _cliente().post(
        "/analizar",
        data=(datos or {}),
        files=archivos,
    )
    nuevos = [TRABAJOS[t] for t in TRABAJOS if t not in previos]
    return respuesta, nuevos


@pytest.fixture
def limpiar_trabajos():
    """Quita los trabajos creados y borra sus carpetas en uploads/ y temp/.

    /analizar escribe en las carpetas reales del proyecto (no dependen de
    AUTO_PRIZMA_DATA_DIR), asi que la prueba debe limpiar lo que crea.
    """
    previos = set(TRABAJOS)
    yield
    for trabajo_id in [t for t in TRABAJOS if t not in previos]:
        trabajo = TRABAJOS.pop(trabajo_id, None) or {}
        for carpeta in (os.path.dirname(trabajo.get("ruta_excel") or ""), trabajo.get("carpeta_temp")):
            if carpeta and os.path.basename(carpeta) == trabajo_id:
                shutil.rmtree(carpeta, ignore_errors=True)


def test_analyze_file_mode_registers_job_and_saves_files(limpiar_trabajos):
    respuesta, nuevos = _analizar(
        {"ovi": "1", "ova": "1", "retos": "1"},
        matriz=_matriz_xlsx(),
        recursos=_zip_recursos(),
    )

    assert respuesta.status_code == 200
    assert len(nuevos) == 1
    trabajo = nuevos[0]

    assert trabajo["etapa"] == "analizado"
    assert trabajo["usuario_app"] == "adminlocal"
    assert (trabajo["procesar_ovi"], trabajo["procesar_ova"], trabajo["procesar_retos"]) == (True, True, True)
    assert trabajo["cursos"] == [{"curso": "Contabilidad I", "programa": "Administración de Empresas"}]
    assert trabajo["ruta_excel"] == os.path.join(UPLOADS_DIR, trabajo["id"], "matriz.xlsx")
    assert trabajo["ruta_zip"] == os.path.join(UPLOADS_DIR, trabajo["id"], "recursos.zip")
    assert os.path.isfile(trabajo["ruta_excel"]) and os.path.isfile(trabajo["ruta_zip"])
    assert trabajo["carpeta_temp"] == os.path.join(TEMP_DIR, trabajo["id"])
    assert os.path.isdir(trabajo["carpeta_temp"])
    assert os.path.dirname(trabajo["ruta_reporte"]) == RESULTADOS_DIR
    assert trabajo["ruta_reporte"].endswith(" - " + trabajo["id"][:6] + ".csv")
    assert set(trabajo) == {
        "id", "ruta_excel", "ruta_zip", "carpeta_temp", "ruta_reporte",
        "procesar_ovi", "procesar_ova", "procesar_retos", "formato_recurso", "etapa", "mensaje",
        "total", "procesadas", "exitosas", "errores", "terminado",
        "detalle_actividades", "nombre_matriz", "usuario_app", "cursos",
        "historial_registrado", "creado_en", "cancelar_solicitado",
        "resultado_analisis",
    }

    resultado = trabajo["resultado_analisis"]
    assert set(resultado) == {"hojas", "zip", "nombre_matriz", "formato_recurso"}
    assert trabajo["formato_recurso"] == resultado["formato_recurso"] == "H5P"
    assert trabajo["id"] in respuesta.text
    assert "Previsualización y títulos H5P" in respuesta.text


@pytest.mark.parametrize(
    ("datos", "con_matriz", "recursos", "mensaje"),
    [
        ({}, True, b"zip", "Selecciona al menos una categoría"),
        ({"ovi": "1"}, False, b"zip", "Selecciona un archivo XLSX o CSV."),
        ({"ovi": "1"}, True, None, "Los recursos deben estar en un archivo .zip"),
        ({"ovi": "1"}, True, b"no es un zip", "El ZIP no es válido."),
    ],
)
def test_analyze_validation_errors_do_not_create_jobs(limpiar_trabajos, datos, con_matriz, recursos, mensaje):
    respuesta, nuevos = _analizar(
        datos,
        matriz=_matriz_xlsx() if con_matriz else None,
        recursos=recursos,
    )

    assert respuesta.status_code == 200
    assert mensaje in respuesta.text
    assert nuevos == []


def test_combined_upload_rejects_invalid_evaluative_test_before_creating_jobs(
    limpiar_trabajos, monkeypatch
):
    from auto_prizma.evaluative_tests.models import (
        DocumentTest,
        Question,
        Section,
        TestDocument,
    )
    from auto_prizma.evaluative_tests.state import _JOBS

    invalid_document = TestDocument(
        tests=[
            DocumentTest(
                title="",
                sections=[
                    Section(
                        code="multipleChoice",
                        name="Selección múltiple",
                        questions=[Question(prompt="")],
                    )
                ],
            )
        ]
    )
    monkeypatch.setattr(
        "auto_prizma.evaluative_tests.parser.parse_docx_bytes",
        lambda _content, source_name="": invalid_document,
    )
    trabajos_previos = set(TRABAJOS)
    tests_previos = set(_JOBS)

    respuesta, nuevos = _analizar(
        {"ovi": "1", "test_evaluativo": "1"},
        matriz=_matriz_xlsx(),
        recursos=_zip_recursos(),
        test_docx=b"contenido docx de prueba",
    )

    assert respuesta.status_code == 200
    assert "No se inició el cargue combinado" in respuesta.text
    assert "El test no tiene nombre o título" in respuesta.text
    assert "La pregunta está vacía" in respuesta.text
    assert "La pregunta debe tener al menos dos opciones" in respuesta.text
    assert "La pregunta de selección no tiene respuesta correcta marcada" in respuesta.text
    assert nuevos == []
    assert set(TRABAJOS) == trabajos_previos
    assert set(_JOBS) == tests_previos


# ============================================================
# FORMATO HTML
# ============================================================

def _zip_html():
    salida = io.BytesIO()
    with zipfile.ZipFile(salida, "w") as archivo:
        archivo.writestr("Infografía de costos.html", "<html><body>Costos</body></html>".encode("utf-8"))
        archivo.writestr("Lectura guía.pdf", b"%PDF")
    return salida.getvalue()


def test_analyze_html_format_marks_interactive_activities_as_html(limpiar_trabajos):
    respuesta, nuevos = _analizar(
        {"ovi": "1", "ova": "1", "retos": "1", "formato_recurso": "html"},
        matriz=_matriz_xlsx(),
        recursos=_zip_html(),
    )

    assert respuesta.status_code == 200
    trabajo = nuevos[0]
    resultado = trabajo["resultado_analisis"]
    assert trabajo["formato_recurso"] == resultado["formato_recurso"] == "HTML"
    tipos = {a["nombre"]: a["tipo"] for a in resultado["hojas"][0]["actividades"]}
    assert tipos == {"Infografía de costos": "HTML", "Lectura guía": "PDF", "Reto 1": "PDF"}
    assert resultado["hojas"][0]["html"] == 1 and resultado["hojas"][0]["h5p"] == 0
    assert resultado["zip"] == {"h5p": 0, "pdf": 1, "html": 1, "total": 2}
    assert '<span class="badge badge-html">HTML</span>' in respuesta.text
    assert "recurso(s) HTML/PDF detectado(s) en el ZIP" in respuesta.text


def test_analyze_rejects_unknown_resource_format(limpiar_trabajos):
    respuesta, nuevos = _analizar(
        {"ovi": "1", "formato_recurso": "SCORM"},
        matriz=_matriz_xlsx(),
        recursos=_zip_recursos(),
    )

    assert "Formato de recurso no válido" in respuesta.text
    assert nuevos == []


def test_html_file_upload_uses_zip_without_external_resource_validation(limpiar_trabajos):
    from auto_prizma.services import upload_jobs

    trabajo_id, resultado = upload_jobs.create_analyzed_job(
        _matriz_xlsx(), _zip_html(), [{"curso": "C", "programa": "P", "actividades": [{}]}],
        {"total": 2}, "Curso A", True, True, True, "adminlocal", formato_recurso="HTML",
    )

    assert resultado["formato_recurso"] == "HTML"
