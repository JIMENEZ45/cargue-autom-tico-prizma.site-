from io import BytesIO
from docx import Document
from fastapi.testclient import TestClient
from main import app
from auto_prizma.services.matrix_analysis import _parsear_datos_academicos
from auto_prizma.services.auth import _firmar_sesion
from auto_prizma.evaluative_tests.state import _AUTOMATION_JOBS, _AUTOMATION_SYNC, _JOBS
import threading


def _session_client() -> TestClient:
    client = TestClient(app, client=("127.0.0.1", 50000))
    client.cookies.set("ap_sesion", _firmar_sesion("adminlocal"))
    return client


def _crear_docx_prueba() -> bytes:
    doc = Document()
    doc.add_heading("EVALUACIÓN TEST", level=1)
    p = doc.add_paragraph()
    p.add_run("1. ¿Cuál es el objeto de la macroeconomía?\n")
    p.add_run("A) Los agregados económicos\n")
    p.add_run("B) La contabilidad individual\n")
    p.add_run("Respuesta correcta: A\n")
    b = BytesIO()
    doc.save(b)
    return b.getvalue()


def test_parsear_datos_academicos_espacios():
    res = _parsear_datos_academicos("ADM101 PEN2024 MAC001 Macroeconomía Aplicada")
    assert res is not None
    assert res["programa"] == "ADM101"
    assert res["pensum"] == "PEN2024"
    assert res["codigo_asignatura"] == "MAC001"
    assert res["asignatura"] == "Macroeconomía Aplicada"


def test_parsear_fila_academica_del_usuario():
    res = _parsear_datos_academicos("V1130\t11301\t30S01\tCOMPETENCIAS MATEMÁTICAS I")
    assert res == {
        "programa": "V1130",
        "pensum": "11301",
        "codigo_asignatura": "30S01",
        "asignatura": "COMPETENCIAS MATEMÁTICAS I",
    }


def test_parsear_datos_academicos_guiones():
    res = _parsear_datos_academicos("ADM101-PEN2024-MAC001-Macroeconomía Aplicada")
    assert res is not None
    assert res["programa"] == "ADM101"
    assert res["pensum"] == "PEN2024"
    assert res["codigo_asignatura"] == "MAC001"
    assert res["asignatura"] == "Macroeconomía Aplicada"


def test_parsear_datos_academicos_guiones_espacios():
    res = _parsear_datos_academicos("ADM101 - PEN2024 - MAC001 - Macroeconomía Aplicada")
    assert res is not None
    assert res["programa"] == "ADM101"
    assert res["pensum"] == "PEN2024"
    assert res["codigo_asignatura"] == "MAC001"
    assert res["asignatura"] == "Macroeconomía Aplicada"


def test_parsear_datos_academicos_invalido():
    assert _parsear_datos_academicos("ADM101 PEN2024") is None
    assert _parsear_datos_academicos("") is None
    assert _parsear_datos_academicos("   ") is None


def test_pagina_inicio_contiene_test_evaluativo():
    client = _session_client()
    response = client.get("/prizma")
    assert response.status_code == 200
    html = response.text
    # Debe contener la tarjeta de Test Evaluativo
    assert 'name="test_evaluativo"' in html
    # Debe estar desmarcado por defecto (no tener 'checked' en el input de test)
    assert '<input type="checkbox" name="test_evaluativo" value="1" id="chk-test">' in html
    # Debe contener el bloque de test
    assert 'id="bloque-test"' in html
    assert 'id="test-datos-academicos"' not in html
    assert 'name="test_semestre"' not in html
    assert 'id="archivo-test-docx"' in html
    assert "Google Docs" not in html
    assert 'name="test_archivo_docx"' in html


def test_analizar_solo_test_flujo_exclusivo():
    client = _session_client()
    doc_bytes = _crear_docx_prueba()

    # Enviar solo test sin OVI, OVA ni Retos, y sin matriz ni zip
    response = client.post(
        "/analizar",
        data={
            "test_evaluativo": "1",
        },
        files={
            "test_archivo_docx": ("test_eval.docx", BytesIO(doc_bytes), "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
        },
        follow_redirects=False,
    )

    # Debe redirigir a /tests/{job_id}/revision
    assert response.status_code == 303
    location = response.headers.get("location", "")
    assert location.startswith("/tests/")
    assert location.endswith("/revision")

    # La línea académica se ingresa en la pantalla de revisión, junto al nivel.
    job_id = location.split("/")[2]
    assert job_id in _JOBS
    revision = client.get(location)
    assert 'name="linea_academica"' in revision.text
    assert 'placeholder="V1130 11301 30S01 COMPETENCIAS MATEMÁTICAS I"' in revision.text
    assert 'name="cascada_facultad"' in revision.text
    assert 'name="cascada_nivel"' in revision.text
    assert 'name="cascada_programa"' not in revision.text
    assert 'name="cascada_pensum"' not in revision.text
    assert 'name="cascada_asignatura"' not in revision.text
    assert 'name="cascada_semana"' not in revision.text
    assert 'name="cascada_corte"' not in revision.text


def test_seleccionar_asignatura_endpoint():
    client = _session_client()
    job_id = "test_job_sync_123"
    _JOBS[job_id] = {"owner": "adminlocal", "document": None}
    sync_obj = {"event": threading.Event(), "seleccion": {"asignatura": None}}
    _AUTOMATION_SYNC[job_id] = sync_obj
    _AUTOMATION_JOBS[job_id] = {
        "owner": "adminlocal",
        "estado": "esperando_asignatura",
        "esperando_asignatura": True,
        "opciones_asignatura": ["Macroeconomía I (MAC101)", "Macroeconomía II (MAC102)"],
        "asignatura_buscada": "Macroeconomía",
    }

    resp = client.post(
        f"/tests/{job_id}/seleccionar-asignatura",
        data={"asignatura": "Macroeconomía I (MAC101)", "selector": "Programa"},
    )
    assert resp.status_code == 200
    data = resp.json()
    assert data["ok"] is True
    assert data["asignatura"] == "Macroeconomía I (MAC101)"
    assert sync_obj["seleccion"]["asignatura"] == "Macroeconomía I (MAC101)"
    assert sync_obj["seleccion"]["campo"] == "Programa"
    assert sync_obj["event"].is_set()


def test_facultad_opcional_en_automatizacion_y_revision():
    from unittest.mock import MagicMock
    from auto_prizma.evaluative_tests.automation import PrizmaTestAutomator

    # 1. Verificar que seleccionar_mui con Facultad retorna True si el selector no existe
    automator = PrizmaTestAutomator(base_url="https://admin.prizma.solutions", headless=True)
    mock_page = MagicMock()
    # Simular que no se localiza ningún campo
    mock_page.locator.return_value.count.return_value = 0
    mock_page.get_by_placeholder.return_value.count.return_value = 0
    mock_page.get_by_label.return_value.count.return_value = 0

    # Para cualquier campo obligatorio, debe lanzar ValueError
    try:
        automator.seleccionar_mui(mock_page, "Programa", "ADM101")
        assert False, "Debió fallar para Programa no encontrado"
    except ValueError:
        pass

    # Para Facultad, al ser opcional, debe retornar True sin lanzar excepción
    resultado_fac = automator.seleccionar_mui(mock_page, "Facultad", "Virtual - Derecho")
    assert resultado_fac is True

    # 2. Verificar que en la página de revisión el campo Facultad esté etiquetado como Opcional
    client = _session_client()
    doc_bytes = _crear_docx_prueba()
    resp_upload = client.post(
        "/tests/analizar",
        files={"archivo_docx": ("test.docx", BytesIO(doc_bytes), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        data={"modo_entrada": "archivo"},
        follow_redirects=False,
    )
    job_id = resp_upload.headers.get("location", "").split("/")[-2]
    resp_rev = client.get(f"/tests/{job_id}/revision")
    assert resp_rev.status_code == 200
    assert "Facultad" in resp_rev.text
    assert "Opcional" in resp_rev.text
    assert 'name="cascada_nivel"' in resp_rev.text
    assert 'name="cascada_semana"' not in resp_rev.text
    assert 'name="cascada_corte"' not in resp_rev.text
